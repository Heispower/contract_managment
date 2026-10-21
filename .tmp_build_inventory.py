import shutil
import sqlite3
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(r"C:\Users\Administrator\Desktop\Contract_System")
SRC = Path(r"C:\Users\Administrator\Desktop\博赛矿业\22-安装包-工具VB\工具VB\传输.xlsm")
OUT = ROOT / "data" / "传输_API梳理.xlsm"
BACKUP = ROOT / "data" / "传输_梳理前备份.xlsm"

# ---------- 当前数据库 ----------
conn = sqlite3.connect(f"file:{ROOT / 'data' / 'contract.db'}?mode=ro", uri=True)
cur = conn.cursor()


def count(table):
    return cur.execute(f"select count(*) from {table}").fetchone()[0]


def columns(table):
    return [r[1] for r in cur.execute(f"pragma table_info({table})")]


def db_type(table, field):
    for row in cur.execute(f"pragma table_info({table})"):
        if row[1] == field:
            raw = (row[2] or "").upper()
            if "INT" in raw:
                return "整数"
            if any(x in raw for x in ("DEC", "NUM", "REAL", "FLOAT", "DOUBLE")):
                return "数值"
            if "DATE" in raw and "TIME" in raw:
                return "时间"
            if "DATE" in raw:
                return "日期"
            if "BOOL" in raw:
                return "是否"
            return "文本"
    return "文本"


# ---------- 字段与接口 ----------
# 字段格式：(api/数据库字段, 中文名, 必填, 说明/枚举)
F = {}


def form(module, name, table, segments):
    F[(module, name, table)] = segments


form("商务信息", "客商", "party", [
    ("主表", [
        ("code", "编码", "是", "公司内唯一；导入按编码更新"),
        ("name", "名称", "是", ""),
        ("type", "客商类型", "是", "supplier 供应商 / customer 客户；保存后不可改"),
        ("short_name", "简称", "否", ""),
        ("tax_no", "纳税人识别号", "否", ""),
        ("default_tax_rate", "默认税率", "否", "新建合同时带出，可改"),
        ("status", "状态", "否", "启用 / 停用；停用不出现在新单据下拉"),
        ("supplier_category", "供应商类别", "否", "仅供应商"),
        ("salesman_id", "业务员", "否", "员工编号；仅客户常用"),
        ("invoice_title", "开票名称", "否", ""),
        ("bank_name", "开户银行", "否", ""),
        ("account_name", "账户名称", "否", ""),
        ("bank_account", "银行账号", "否", ""),
        ("bank_branch", "开户支行", "否", ""),
        ("invoice_phone", "开票电话", "否", ""),
        ("invoice_address", "开票地址", "否", ""),
        ("address", "通讯地址", "否", ""),
        ("email", "邮箱", "否", ""),
        ("mobile", "手机", "否", ""),
        ("legal_name", "法人代表", "否", ""),
        ("agent_name", "委托代理人", "否", ""),
        ("remark", "备注", "否", ""),
    ]),
    ("联系人", [
        ("name", "联系人", "否", "party_contact 子表"),
        ("phone", "联系电话", "否", "party_contact 子表"),
        ("is_default", "默认联系人", "否", "同一客商至多一个默认"),
    ]),
])

form("商务信息", "项目", "project", [("主表", [
    ("project_code", "项目编号", "是", "公司内唯一"),
    ("name", "项目名称", "是", ""),
    ("year", "年度", "否", ""),
    ("project_type", "项目类型", "否", ""),
    ("source", "来源", "否", ""),
    ("customer_name", "客户名称", "否", "自由文本，不挂客商ID"),
    ("start_date", "开始日期", "否", ""),
    ("end_date", "结束日期", "否", ""),
    ("status", "状态", "否", "进行中 / 已完成 / 已终止"),
    ("remark", "备注", "否", ""),
])])

form("采购计划", "采购计划", "purchase_plan", [("主表", [
    ("plan_month", "计划月份", "否", ""),
    ("purchase_seq", "采购序号", "否", ""),
    ("material_code", "物料编号", "是", "公司内唯一；编辑时不可改"),
    ("plan_no", "采购计划单号", "否", ""),
    ("category", "使用设备", "否", ""),
    ("name_cn", "中文名称", "是", ""),
    ("name_en", "英文名称", "否", ""),
    ("spec", "规格型号", "否", ""),
    ("unit", "单位", "否", ""),
    ("qty", "采购数量", "是", "修改后不能小于已采购数量"),
    ("purpose", "用途", "否", ""),
    ("use_dept", "使用部门", "否", ""),
    ("stock_qty", "库存数量", "否", ""),
    ("transit_qty", "在途数量", "否", ""),
    ("consume_4m", "4个月消耗量", "否", ""),
    ("remark", "备注", "否", ""),
])])

form("合同", "采购合同", "contract", [
    ("主表", [
        ("contract_code", "合同编号", "否", "留空按编号规则生成；保存后不改"),
        ("name", "合同名称", "是", ""),
        ("nature", "合同性质", "是", "main 主合同 / supplement 补充协议"),
        ("direction", "方向", "是", "purchase"),
        ("type_name", "合同类型", "否", "基础资料合同类型"),
        ("party_id", "供应商ID", "是", "客商 party_id"),
        ("party_name", "供应商名称", "否", "保存时快照"),
        ("party_tax_no", "供应商税号", "否", "快照"),
        ("owner_id", "业务员", "是", "员工编号"),
        ("status", "合同状态", "是", "未签订 / 已签订 / 已终止"),
        ("archived", "归档状态", "否", "只读，恒为未归档"),
        ("project_id", "项目ID", "否", ""),
        ("sign_date", "签订日期", "否", ""),
        ("start_date", "开始日期", "否", ""),
        ("end_date", "结束日期", "否", "到期提醒取此日期"),
        ("pricing", "计价方式", "是", "fixed 固定总价 / actual 按实结算"),
        ("tax_rate", "税率", "是", "百分比数值，13 表示 13%"),
        ("amount_manual", "手填合同金额", "否", "无明细时使用"),
        ("amount", "合同金额", "否", "服务端计算，含税"),
        ("amount_ex_tax", "不含税金额", "否", "服务端计算"),
        ("settled_sum", "已收/已付", "否", "收付款回写，只读"),
        ("content", "内容概述", "否", ""),
        ("terms", "条款/付款条款", "否", "Word 模板使用"),
        ("subject_name", "标的名称", "否", ""),
        ("ship_address", "收货地址", "否", ""),
        ("contact_name", "联系人", "否", ""),
        ("contact_phone", "联系电话", "否", ""),
        ("party_bank_name", "对方开户银行", "否", "对方快照"),
        ("party_account_name", "对方账户名称", "否", "对方快照"),
        ("party_bank_account", "对方银行账号", "否", "对方快照"),
        ("party_bank_branch", "对方开户支行", "否", "对方快照"),
        ("party_invoice_title", "对方开票名称", "否", "对方快照"),
        ("party_invoice_phone", "对方开票电话", "否", "对方快照"),
        ("party_invoice_address", "对方开票地址", "否", "对方快照"),
        ("party_address", "对方地址", "否", "对方快照"),
        ("party_legal_name", "对方法人", "否", "对方快照"),
        ("party_agent", "对方委托代理人", "否", "对方快照"),
        ("transport_mode", "运输方式", "否", ""),
        ("delivery_term", "交货期限", "否", ""),
        ("freight_bearer", "运费承担", "否", ""),
        ("receiver_name", "收货人", "否", ""),
        ("acceptance_text", "验收标准", "否", ""),
        ("annex_note", "附件说明", "否", ""),
        ("invoice_locked", "发票锁定", "否", "只读，超开时置位"),
        ("remark", "备注", "否", ""),
    ]),
    ("物料明细", [
        ("material_code", "物料编号", "是", "同一合同内唯一"),
        ("name_cn", "中文名称", "否", ""),
        ("name_en", "英文名称", "否", ""),
        ("spec", "规格", "否", ""),
        ("unit", "单位", "否", ""),
        ("qty", "数量", "是", "必须大于 0"),
        ("price", "含税单价", "是", "不能小于 0"),
        ("plan_no", "计划单号", "否", ""),
        ("category", "使用设备", "否", ""),
        ("remark", "备注", "否", ""),
    ]),
    ("阶段付款", [
        ("seq", "阶段序号", "否", "服务端排序"),
        ("name", "阶段名称", "是", ""),
        ("ratio", "比例", "否", "百分比数值"),
        ("plan_amount", "计划金额", "否", "合计必须等于合同金额"),
        ("plan_date", "计划日期", "否", "逾期判断依据"),
        ("settled_amount", "已收/已付", "否", "收付款回写，只读"),
        ("arrival_triggered", "到货触发", "否", "圆钮确认/取消后置位"),
        ("ship_triggered", "发货触发", "否", "圆钮确认/取消后置位"),
        ("remark", "备注", "否", ""),
    ]),
])

form("仓储", "采购入库", "wms_bill", [
    ("主表", [
        ("bill_code", "单据编号", "否", "留空按 RK-日期-流水生成"),
        ("bill_date", "单据日期", "是", ""),
        ("type", "类型", "是", "inbound"),
        ("contract_id", "合同ID", "是", "只能选已签订采购合同"),
        ("party_id", "供应商ID", "否", "随合同带出"),
        ("party_name", "供应商名称", "否", "快照"),
        ("warehouse_id", "仓库ID", "是", ""),
        ("remark", "备注", "否", ""),
    ]),
    ("物料明细", [
        ("material_code", "物料编号", "是", "必须属于合同"),
        ("name_cn", "中文名称", "否", ""),
        ("spec", "规格", "否", ""),
        ("unit", "单位", "否", ""),
        ("qty", "数量", "是", "累计不能超过合同签约数量"),
        ("price", "含税单价", "否", "库存成本取最近入库价"),
        ("amount", "金额", "否", "服务端计算"),
        ("category", "使用设备", "否", ""),
        ("remark", "备注", "否", ""),
    ]),
])

form("仓储", "库存余额", "stock", [("主表", [
    ("warehouse_id", "仓库ID", "是", "仓库 + 物料唯一"),
    ("material_code", "物料编号", "是", ""),
    ("goods_code", "商品编码", "否", "服务端回写"),
    ("goods_name", "商品名称", "否", "服务端回写"),
    ("unit", "单位", "否", "服务端回写"),
    ("qty", "数量", "否", "入库 − 出库重算，只读"),
    ("cost", "成本", "否", "最近入库含税单价，只读"),
])])

form("结算", "付款单", "settle_bill", [("主表", [
    ("bill_code", "单据编号", "否", "留空按 FK-日期-流水生成"),
    ("bill_date", "单据日期", "是", ""),
    ("direction", "方向", "是", "payment"),
    ("party_id", "供应商ID", "是", ""),
    ("party_name", "供应商名称", "否", "快照"),
    ("contract_id", "合同ID", "是", ""),
    ("stage_id", "阶段ID", "是", ""),
    ("amount", "金额", "是", "大于 0，不超过阶段剩余"),
    ("account_name", "账户", "否", "资金账户名称"),
    ("remark", "备注", "否", ""),
])])

form("发票", "进项发票", "invoice", [
    ("主表", [
        ("bill_code", "单据编号", "否", "留空按 JX-日期-流水生成"),
        ("invoice_no", "发票号码", "是", "公司内唯一"),
        ("invoice_date", "发票日期", "是", ""),
        ("invoice_type", "发票类型", "是", "增值税专用发票/普通发票/电子专票/电子普票"),
        ("direction", "方向", "是", "input"),
        ("party_id", "供应商ID", "是", ""),
        ("party_name", "供应商名称", "否", "快照"),
        ("party_tax_no", "税号", "否", "快照"),
        ("contract_id", "合同ID", "是", "只能选已签订采购合同"),
        ("attachments", "附件", "否", "JSON/文件标识"),
        ("amount_ex_tax", "不含税金额", "否", "服务端按明细汇总"),
        ("tax", "税额", "否", "服务端按明细汇总"),
        ("total", "价税合计", "否", "服务端按明细汇总"),
        ("remark", "备注", "否", ""),
    ]),
    ("物料明细", [
        ("material_code", "物料编号", "是", "必须在合同物料中"),
        ("name_cn", "中文名称", "否", ""),
        ("name_en", "英文名称", "否", ""),
        ("spec", "规格", "否", ""),
        ("unit", "单位", "否", ""),
        ("qty", "数量", "是", "累计不超合同签约数量"),
        ("price_ex_tax", "不含税单价", "否", ""),
        ("price", "含税单价", "否", ""),
        ("amount_ex_tax", "不含税金额", "否", ""),
        ("tax_rate", "税率", "否", ""),
        ("tax", "税额", "否", ""),
        ("total", "价税合计", "否", "行主值"),
        ("manual_tax", "手工改税", "否", "是否直接改过不含税/税额"),
        ("remark", "备注", "否", ""),
    ]),
])

form("报关", "报关物料", "customs_material", [("主表", [
    ("material_code", "物料编号", "是", "唯一主键"),
    ("receipt_no", "入仓单号", "否", ""),
    ("contract_no", "合同号", "否", ""),
    ("packing_list_no", "箱单号", "否", ""),
    ("material_category", "物料大类", "否", ""),
    ("equipment", "使用设备", "否", ""),
    ("origin_country", "产地", "否", ""),
    ("quantity", "数量", "否", "大于 0"),
    ("unit", "单位", "否", ""),
    ("price", "含税单价", "否", ""),
    ("gross_weight", "毛重", "否", ""),
    ("net_weight", "净重", "否", ""),
    ("volume", "体积", "否", ""),
    ("length", "长", "否", ""),
    ("width", "宽", "否", ""),
    ("height", "高", "否", ""),
    ("box_no", "箱号", "否", ""),
    ("container_no", "集装箱号", "否", ""),
    ("bill_no", "提单号", "否", ""),
    ("vessel_name", "船名", "否", ""),
    ("voyage_no", "航号", "否", ""),
    ("po_no", "PO号", "否", ""),
    ("port_of_shipment", "出发港", "否", ""),
    ("port_of_destination", "目的港", "否", ""),
    ("consignee", "收货人", "否", ""),
    ("payment_terms", "付款条款", "否", ""),
    ("hs_code", "HS Code", "否", ""),
    ("declaration_elements", "申报要素", "否", ""),
    ("special_declare", "特殊申报", "否", "是/否"),
    ("package_count", "包件件数", "否", ""),
    ("status", "状态", "否", ""),
    ("remark", "备注", "否", ""),
])])

form("编号", "合同编号", "number_contract", [("主表", [
    ("contract_no", "合同编号", "是", "编号规则生成或手工登记"),
    ("business_type", "业务类型", "是", "锰矿 / 铝矿"),
    ("number_year", "编号年份", "是", "2 位或 4 位按规则"),
    ("num_pool", "编号池", "否", "服务端记录"),
    ("num_seq", "顺序号", "否", "服务端记录；改名后继续递增"),
    ("product_name", "产品/合同名称", "否", ""),
    ("party_b", "乙方", "否", ""),
    ("party_a", "甲方", "否", ""),
    ("amount_cny", "合同金额CNY", "否", ""),
    ("sign_date", "签订时间", "否", ""),
    ("owner_id", "归属人", "否", "员工编号"),
    ("purchase_contract_id", "关联采购合同ID", "否", "系统回写"),
    ("manual_edited", "手工改号", "否", "只读"),
])])

form("编号", "发票编号", "number_invoice", [("主表", [
    ("invoice_no", "发票号", "是", "编号规则生成或手工登记"),
    ("number_rule", "编号规则", "是", "锰矿 / 铝矿海运 / 铝矿空运"),
    ("business_type", "业务类型", "否", ""),
    ("number_year", "编号年份", "是", ""),
    ("num_pool", "编号池", "否", "服务端记录"),
    ("num_seq", "顺序号", "否", "服务端记录；改名后继续递增"),
    ("contract_no", "关联合同", "否", ""),
    ("number_contract_id", "关联合同编号ID", "否", ""),
    ("content", "所含内容", "是", ""),
    ("bl_date", "提单日期", "否", ""),
    ("price_term", "价格条款", "否", "FOB / CIF"),
    ("currency", "结算币种", "否", "USD / CNY"),
    ("amount_usd", "美元金额", "否", ""),
    ("amount_cny", "人民币金额", "否", ""),
    ("exchange_rate", "汇率", "否", ""),
    ("exchange_month", "汇率月份", "否", "YYYY-MM"),
    ("po_no", "PO号", "否", ""),
    ("operator_name", "登记人", "否", "留空默认当前用户"),
    ("transport_type", "运输方式", "否", "海运 / 空运"),
    ("email_sent_date", "发邮件时间", "否", ""),
    ("bl_no", "提单号", "否", ""),
    ("is_batch", "批次", "否", "是否批次"),
    ("is_batch_master", "批次主发票", "否", "主发票金额由小批汇总"),
    ("batch_prefix", "批次前缀", "否", ""),
    ("batch_sub_no", "分单号", "否", ""),
    ("manual_edited", "手工改号", "否", "只读"),
    ("remark", "备注", "否", ""),
])])

form("基础资料", "公司", "company", [("主表", [
    ("code", "公司编码", "否", ""),
    ("name", "公司名称", "是", ""),
    ("short_name", "简称", "否", "顶栏显示"),
    ("status", "状态", "否", "启用 / 停用"),
    ("default_tax_rate", "默认税率", "否", ""),
    ("salesman_id", "业务员", "否", ""),
    ("tax_no", "纳税人识别号", "否", ""),
    ("invoice_title", "开票名称", "否", ""),
    ("address", "通讯地址", "否", ""),
    ("email", "邮箱", "否", ""),
    ("mobile", "手机", "否", ""),
    ("bank_name", "开户银行", "否", ""),
    ("account_name", "账户名称", "否", ""),
    ("bank_account", "银行账号", "否", ""),
    ("bank_branch", "开户支行", "否", ""),
    ("invoice_phone", "开票电话", "否", ""),
    ("invoice_address", "开票地址", "否", ""),
    ("legal_name", "法人代表", "否", ""),
    ("agent_name", "委托代理人", "否", ""),
    ("remark", "备注", "否", ""),
])])

form("基础资料", "仓库", "warehouse", [("主表", [
    ("code", "编码", "是", "公司内唯一"),
    ("name", "名称", "是", ""),
    ("status", "状态", "否", "启用 / 停用"),
    ("remark", "备注", "否", ""),
])])

form("基础资料", "HS Code库", "hs_code_library", [("主表", [
    ("hs_code", "HS Code", "是", "公司内唯一"),
    ("declaration_name", "申报名称", "是", ""),
    ("product_name", "产品名称", "否", ""),
    ("specification", "申报规格", "否", ""),
    ("declaration_elements", "报关要素", "否", ""),
    ("description", "描述", "否", ""),
    ("status", "状态", "否", "启用 / 停用"),
])])

form("基础资料", "汇率库", "exchange_rate", [("主表", [
    ("year", "年", "是", ""),
    ("month", "月", "是", ""),
    ("currency", "币种", "否", "USD/CNY 等"),
    ("rate", "汇率", "是", "大于 0"),
    ("effective_date", "生效日期", "否", ""),
])])

form("基础资料", "编号规则", "number_rule", [("主表", [
    ("pool_key", "编号池", "是", "contract_mn/contract_al/invoice_*"),
    ("label", "单据名称", "是", ""),
    ("prefix", "前缀", "否", "改名只改这里"),
    ("year_digits", "年份位数", "是", "2 或 4"),
    ("seq_width", "序号位数", "是", "1–8"),
    ("suffix", "后缀", "否", "如 S/A"),
    ("enabled", "启用", "否", "是/否"),
    ("remark", "备注", "否", ""),
])])

form("基础资料", "导出模板", "export_profile", [("主表", [
    ("doc_code", "单据类型", "是", "purchase_contract/sale_contract 等"),
    ("doc_name", "单据名称", "是", ""),
    ("title", "模板名称", "否", ""),
    ("party_id", "适用客商", "否", "空为不限"),
    ("template_path", "模板绝对地址", "是", ""),
    ("name_rule", "导出名称规则", "否", "支持 {{字段}}"),
    ("sort_no", "排序", "否", ""),
])])

form("组织管理", "用户", "sys_user", [("主表", [
    ("employee_id", "员工编号", "是", "全局唯一"),
    ("login_name", "登录名", "是", "唯一"),
    ("name", "中文姓名", "是", ""),
    ("password", "初始密码", "否", "仅新增；空则 123456"),
    ("role_text", "角色文本", "否", "不控制权限"),
    ("status", "状态", "否", "启用 / 停用"),
    ("is_super_admin", "超级管理员", "否", "仅 admin"),
    ("can_import_overwrite", "导入覆盖", "否", "是否允许覆盖他人记录"),
])])

# ---------- API 行 ----------
API = [
    ("商务信息", "客商", "party", count("party"),
     "GET /api/parties\n参数 type=supplier|customer、status、keyword、page、pageSize",
     "GET /api/parties/{party_id}", "POST /api/parties", "PUT /api/parties/{party_id}",
     "DELETE /api/parties/{party_id}?version=",
     "POST /api/import/party/parse|execute\n模板 GET /api/import/party/template",
     "GET/POST /api/export/party?type=supplier|customer",
     "type 保存后不可改；contacts 子表随主表提交；改名可调 sync-name(s)"),
    ("商务信息", "项目", "project", count("project"),
     "GET /api/projects\n参数 keyword、project_type、source、year、status、customer_name、start_date、end_date",
     "GET 列表即可（无独立详情接口）", "POST /api/projects", "PUT /api/projects/{row_id}",
     "DELETE /api/projects/{row_id}?version=", "无导入",
     "GET/POST /api/export/project", "项目客户名称是自由文本"),
    ("采购计划", "采购计划", "purchase_plan", count("purchase_plan"),
     "GET /api/purchase-plans\n参数 keyword、plan_month、name_cn、purchase_state",
     "无独立详情接口", "无新增接口（Excel 导入）", "PUT /api/purchase-plans/{row_id}",
     "DELETE /api/purchase-plans/{row_id}?version=",
     "POST /api/import/purchase-plan/parse|execute\n模式 update/replace/append",
     "GET/POST /api/export/purchase-plan",
     "物料编号不可改；数量不能小于已采购数量"),
    ("合同", "采购合同", "contract", count("contract"),
     "GET /api/contracts?direction=purchase\n参数 status、party_id、project_id、pay_status 等",
     "GET /api/contracts/{contract_id}", "POST /api/contracts", "PUT /api/contracts/{contract_id}",
     "DELETE /api/contracts/{contract_id}?version=",
     "POST /api/import/contract/parse|execute?direction=purchase",
     "GET/POST /api/export/contract?direction=purchase\nWord GET /api/contracts/{id}/word",
     "items/stages 内嵌在合同 payload；金额、已收付由服务端计算"),
    ("仓储", "采购入库", "wms_bill", count("wms_bill"),
     "GET /api/wms-bills?type=inbound\n参数 contract_id、keyword",
     "GET /api/wms-bills/{bill_id}", "POST /api/wms-bills", "PUT /api/wms-bills/{bill_id}",
     "DELETE /api/wms-bills/{bill_id}?version=", "无通用导入（可走入仓单）",
     "GET/POST /api/export/wms?type=inbound",
     "items 内嵌；保存后整单重算库存"),
    ("仓储", "库存余额", "stock", count("stock"),
     "GET /api/stocks\n参数 keyword、warehouse_id",
     "无", "无（由出入库重算）", "无",
     "DELETE /api/stocks/{row_id}?version=（仅数量为 0）",
     "无", "GET/POST /api/export/stock", "只读派生表，不允许直接改数量"),
    ("结算", "付款单", "settle_bill", count("settle_bill"),
     "GET /api/settle-bills?direction=payment\n参数 party_id、contract_id",
     "无独立详情接口", "POST /api/settle-bills/multi-stage", "PUT /api/settle-bills/{bill_id}",
     "DELETE /api/settle-bills/{bill_id}?version=",
     "POST /api/import/settle/parse|execute?direction=payment",
     "GET/POST /api/export/settle?direction=payment",
     "按合同阶段付款；保存后回写阶段与合同已付"),
    ("发票", "进项发票", "invoice", count("invoice"),
     "GET /api/invoices?direction=input\n参数 keyword、party_id、project_id、invoice_type",
     "GET /api/invoices/{invoice_id}", "POST /api/invoices", "PUT /api/invoices/{invoice_id}",
     "DELETE /api/invoices/{invoice_id}?version=",
     "POST /api/import/invoice/parse|execute?direction=input",
     "GET/POST /api/export/invoice?direction=input\nWord GET /api/forms/input_invoice/{id}/word",
     "items 内嵌；数量与金额双层校验；发票号唯一"),
    ("报关", "报关物料", "customs_material", count("customs_material"),
     "GET /api/customs/materials\n参数 contract_no、packing_list_no、material_code、receipt_no、supplier_name、equipment",
     "GET /api/customs/materials/{material_id}", "POST /api/customs/materials", "PUT /api/customs/materials/{material_id}",
     "DELETE /api/customs/materials/{material_id}?version=",
     "POST /api/import/customs-material-fill/parse|execute",
     "GET/POST /api/export/customs-material",
     "批量赋值 POST /api/customs/materials/batch-assign；被发票/要素引用时不可删"),
    ("编号", "合同编号", "number_contract", count("number_contract"),
     "GET /api/numbering/contracts\n参数 keyword、business_type、number_year、order",
     "GET /api/numbering/contracts/{row_id}", "POST /api/numbering/contracts", "PUT /api/numbering/contracts/{row_id}",
     "DELETE /api/numbering/contracts/{row_id}?version=",
     "POST /api/import/number-contract/parse|execute",
     "GET/POST /api/export/contract（编号页面用编号导出/导入）",
     "手工改号 POST /{id}/manual；顺序号下限 PUT /api/numbering/floors"),
    ("编号", "发票编号", "number_invoice", count("number_invoice"),
     "GET /api/numbering/invoices\n参数 keyword、number_rule、number_year、order",
     "GET /api/numbering/invoices/{row_id}", "POST /api/numbering/invoices", "PUT /api/numbering/invoices/{row_id}",
     "DELETE /api/numbering/invoices/{row_id}?version=",
     "POST /api/import/number-invoice/parse|execute",
     "编号导出/导入",
     "登记人字段 operator_name；手工改号 POST /{id}/manual"),
    ("基础资料", "公司", "company", count("company"),
     "GET /api/companies", "无", "POST /api/companies", "PUT /api/companies/{row_id}",
     "DELETE /api/companies/{row_id}?version=", "无", "无", "改简称即改顶栏"),
    ("基础资料", "仓库", "warehouse", count("warehouse"),
     "GET /api/warehouses", "无", "POST /api/warehouses", "PUT /api/warehouses/{row_id}",
     "DELETE /api/warehouses/{row_id}?version=", "无", "GET/POST /api/export/warehouses",
     "库存不为零或存在有效单据时不能停用/删除"),
    ("基础资料", "HS Code库", "hs_code_library", count("hs_code_library"),
     "GET /api/customs/hs", "无", "POST /api/customs/hs", "PUT /api/customs/hs/{hs_id}",
     "DELETE /api/customs/hs/{hs_id}?version=", "POST /api/import/hs-code/parse|execute",
     "GET/POST /api/export/hs-code", "报关要素匹配用"),
    ("基础资料", "汇率库", "exchange_rate", count("exchange_rate"),
     "GET /api/customs/rates", "无", "POST /api/customs/rates", "PUT /api/customs/rates/{rate_id}",
     "DELETE /api/customs/rates/{rate_id}?version=", "无", "GET/POST /api/export/exchange-rate",
     "同月同币种唯一；报关换算用"),
    ("基础资料", "编号规则", "number_rule", count("number_rule"),
     "GET /api/numbering/rules", "无（列表已含预览）", "无", "PUT /api/numbering/rules/{rule_id}",
     "无", "无", "无", "管理员维护；只影响以后新编号"),
    ("基础资料", "导出模板", "export_profile", count("export_profile"),
     "GET /api/export-profiles?doc_code=", "无", "POST /api/export-profiles", "PUT /api/export-profiles/{row_id}",
     "DELETE /api/export-profiles/{row_id}?version=", "POST /api/export-profiles/upload",
     "GET /api/export-profiles/fields", "模板绝对地址；按客商匹配"),
    ("组织管理", "用户", "sys_user", count("sys_user"),
     "GET /api/users", "无", "POST /api/users", "PUT /api/users/{row_id}",
     "DELETE /api/users/{row_id}?version=", "无", "GET/POST /api/export/users",
     "重置密码 PUT /{id}/reset-password；强制退出 POST /{id}/force-logout"),
    ("视图", "待付款", "contract + settle_bill", "派生",
     "GET /api/pending?direction=purchase&months=&party_id=", "无", "只读", "只读", "只读",
     "无", "GET/POST /api/export/pending", "按阶段汇总，不是独立表"),
    ("视图", "到期提醒", "reminder_record", count("reminder_record"),
     "GET /api/reminders?tab=sale|purchase|output|input", "无", "POST /api/reminders/check",
     "POST /api/reminders/mark", "无", "无", "GET/POST /api/export/reminders",
     "只记录处理状态，不改合同/发票"),
    ("视图", "首页/统计", "聚合", "派生",
     "GET /api/stats/home、/overview、/customer、/supplier、/supplier-detail",
     "无", "只读", "只读", "只读", "无", "GET/POST /api/export/stats-*", "金额口径只计已签订合同"),
]

# ---------- 写工作簿 ----------
if not SRC.exists():
    raise SystemExit(f"源文件不存在：{SRC}")
OUT.parent.mkdir(parents=True, exist_ok=True)
shutil.copy2(SRC, BACKUP)
wb = load_workbook(SRC, keep_vba=True)

for name in ("API接口清单", "字段清单", "梳理说明"):
    if name in wb.sheetnames:
        del wb[name]

HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
HEAD_FONT = Font(color="FFFFFF", bold=True, size=11)
THIN = Side(style="thin", color="D9E2F3")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_sheet(ws, widths, row_height=24):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 30
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    for cell in ws[1]:
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = BORDER
        ws.row_dimensions[row[0].row].height = row_height
    ws.sheet_view.showGridLines = False


# API 清单
ws = wb.create_sheet("API接口清单")
api_headers = ["序号", "模块", "表单/页面", "主表", "数据量", "查询接口", "详情接口", "新增接口",
               "修改接口", "删除接口", "导入接口", "导出/打印", "关键字段/说明"]
ws.append(api_headers)
for idx, row in enumerate(API, start=1):
    ws.append([idx, *row])
style_sheet(ws, [6, 10, 18, 18, 9, 42, 32, 34, 38, 34, 40, 38, 52], 78)
ws.sheet_properties.tabColor = "1F4E78"

# 字段清单
ws2 = wb.create_sheet("字段清单")
ws2.append(["序号", "模块", "表单/页面", "分段", "字段名（API/数据库）", "中文名", "类型", "必填", "说明/枚举"])
idx = 1
for (module, name, table), segments in F.items():
    for segment, fields in segments:
        for field, label, required, note in fields:
            ws2.append([idx, module, name, segment, field, label, db_type(table, field), required, note])
            idx += 1
style_sheet(ws2, [6, 10, 18, 14, 26, 20, 8, 8, 52], 22)
ws2.sheet_properties.tabColor = "2E75B6"

# 说明
ws3 = wb.create_sheet("梳理说明")
ws3.append(["项目", "说明"])
notes = [
    ("梳理对象", "合同管理系统中当前常用、且有数据的表单/台账"),
    ("数据源", "data/contract.db 当前库（只读统计）"),
    ("梳理时间", "2026-10-21"),
    ("接口风格", "REST + JSON，统一前缀 /api；登录使用 Authorization: Bearer {token}"),
    ("响应结构", "成功 {code:0,message:'ok',data:...}；列表 data={items,total,page,pageSize}"),
    ("分页", "page 从 1 开始，pageSize 默认 50，最大 500"),
    ("并发", "修改/删除带 version；单条删除用 ?version=；批量删除 items=[{id,version}]"),
    ("公司隔离", "服务端从 token 取当前公司；company_id 只做一致性校验"),
    ("公共字段", "除系统表外业务表均带 id、company_id、created_by/at、updated_by/at、version、deleted_at/by、reserve1-3；字段清单不重复列出"),
    ("列表展示字段", "列表常额外返回 party_code/party_name、created_by_name、pay_status、stage_text 等展示字段，不落库"),
    ("金额", "接口金额一般是 2 位小数字符串；计算以服务端 Decimal 为准，行价税合计为主值"),
    ("当前有数据的表单", "客商 88、项目 1、采购计划 131、采购合同 11、采购入库 4、库存 4、付款单 11、进项发票 11、报关物料 4、合同编号 112、发票编号 9、HS Code 63、汇率 1、编号规则 5、导出模板 13、用户 5"),
    ("当前无数据/未展开", "销售合同 0、销售出库 0、收款单 0、销项发票 0、采购订单 0、销售订单 0、其他收入/支出 0、报关发票/清关发票/报关要素 0、调拨/盘点 0、收货入仓 0；接口存在但本次未列字段明细"),
    ("导入导出", "列表导出统一 GET/POST /api/export/{module}；Excel 导入统一 /api/import/{module}/template、/parse、/execute"),
    ("字段说明", "字段名以 API/数据库 snake_case 为准；类型来自 SQLite 列类型，日期通过接口传 YYYY-MM-DD"),
    ("维护建议", "新增业务字段时，先更新模型/接口，再在“字段清单”追加；如该表参与 Excel 导入，同步更新“字段库”和 VB 模板"),
]
for row in notes:
    ws3.append(row)
style_sheet(ws3, [24, 130], 30)
ws3.sheet_properties.tabColor = "548235"

# 使用说明补一行
use = wb["使用说明"]
note_row = 14
if not use.cell(note_row, 1).value:
    use.cell(note_row, 1, "10. 「API接口清单」「字段清单」「梳理说明」是系统常用有数据表单的接口/字段梳理，只读，不参与导入。")

wb.save(OUT)
print(f"written: {OUT}")
print(f"backup:  {BACKUP}")
print(f"sheets:  {wb.sheetnames}")
print(f"api rows={len(API)}, field rows={idx - 1}")
wb.close()
conn.close()
