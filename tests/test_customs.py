"""报关数据链：采购合同 → 采购入库 → 报关物料 → 发票 / 要素，以及特殊申报与导出校验。"""
from decimal import Decimal
from io import BytesIO
from zipfile import ZipFile

import openpyxl
import pytest
from fastapi.testclient import TestClient

from tests.test_rules import login
from backend.app.main import app


def _write_docx(path, text):
    document = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f'<w:body><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:body></w:document>'
    )
    types = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
        '</Types>'
    )
    rels = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
        '</Relationships>'
    )
    with ZipFile(path, "w") as zf:
        zf.writestr("[Content_Types].xml", types)
        zf.writestr("_rels/.rels", rels)
        zf.writestr("word/document.xml", document)


@pytest.fixture(scope="module")
def client():
    with TestClient(app) as c:
        yield c


def _item(code, name, category="轴承", special=False, qty="10", extra=None):
    row = {
        "material_code": code,
        "material_name": name,
        "material_name_en": "Bearing",
        "material_category": category,
        "specification": "6205",
        "unit": "套",
        "quantity": qty,
        "origin_country": "中国",
        "equipment": "包装机",
        "box_no": "BOX001",
        "container_no": "CONT001",
        "bill_no": "BL001",
        "vessel_name": "SEA STAR",
        "voyage_no": "V01",
        "gross_weight": "12",
        "net_weight": "10",
        "length": "30",
        "width": "20",
        "height": "15",
        "volume": "0.009",
        "po_no": "PO-1",
        "price": "100",
        "special_declare": special,
    }
    if extra:
        row.update(extra)
    return row


def _plans(items):
    from backend.app.database import SessionLocal
    from backend.app.models import Company, PurchasePlan, now

    db = SessionLocal()
    try:
        company = db.query(Company).filter(Company.short_name == "星河科技").one()
        for item in items:
            plan = db.query(PurchasePlan).filter(
                PurchasePlan.company_id == company.id,
                PurchasePlan.material_code == item["material_code"],
                PurchasePlan.deleted_at.is_(None),
            ).first()
            need = Decimal(str(item["quantity"]))
            if plan is None:
                db.add(PurchasePlan(
                    company_id=company.id,
                    plan_month="2026-10",
                    material_code=item["material_code"],
                    name_cn=item["material_name"],
                    spec=item.get("specification") or "",
                    unit=item.get("unit") or "",
                    qty=need,
                    purpose="测试",
                    version=1,
                    created_by="admin",
                    updated_by="admin",
                    created_at=now(),
                    updated_at=now(),
                ))
            elif plan.qty < need:
                plan.qty = need
        db.commit()
    finally:
        db.close()


def _chain(client, headers, receipt_no, contract_no, packing, items):
    _plans(items)
    party = client.get("/api/parties", params={"type": "supplier"}, headers=headers).json()["data"]["items"][0]
    warehouse = client.get("/api/warehouses", headers=headers).json()["data"]["items"][0]
    created = client.post("/api/contracts", json={
        "contract_code": contract_no,
        "name": f"{contract_no}采购合同",
        "direction": "purchase",
        "nature": "main",
        "type_name": "材料采购",
        "party_id": party["id"],
        "owner_id": "admin",
        "status": "已签订",
        "pricing": "fixed",
        "tax_rate": "13",
        "sign_date": "2026-10-01",
        "items": [{
            "material_code": item["material_code"],
            "name_cn": item["material_name"],
            "spec": item.get("specification") or "",
            "unit": item.get("unit") or "",
            "qty": item["quantity"],
            "price": item.get("price") or "100",
            "category": item.get("equipment") or "",
        } for item in items],
    }, headers=headers)
    assert created.status_code == 200, created.text
    inbound = client.post("/api/wms-bills", json={
        "type": "inbound",
        "bill_code": receipt_no,
        "bill_date": "2026-10-01",
        "contract_id": created.json()["data"]["id"],
        "warehouse_id": warehouse["id"],
        "items": [{
            "material_code": item["material_code"],
            "name_cn": item["material_name"],
            "spec": item.get("specification") or "",
            "unit": item.get("unit") or "",
            "qty": item["quantity"],
            "price": item.get("price") or "100",
            "category": item.get("equipment") or "",
        } for item in items],
    }, headers=headers)
    assert inbound.status_code == 200, inbound.text
    mats = client.get("/api/customs/materials", params={"contract_no": contract_no, "pageSize": 50}, headers=headers).json()["data"]["items"]
    by_code = {row["material_code"]: row for row in mats}
    for item in items:
        synced = by_code[item["material_code"]]
        assert synced["receipt_no"] in ("", None)
        assert synced["material_category"] == ""
        assert synced["equipment"] == (item.get("equipment") or "")
    for item in items:
        row = by_code[item["material_code"]]
        payload = dict(item)
        payload.update({"id": row["id"], "version": row["version"], "packing_list_no": packing})
        payload["vessel_voyage"] = "/".join(part for part in (item.get("vessel_name"), item.get("voyage_no")) if part)
        saved = client.put(f"/api/customs/materials/{row['id']}", json=payload, headers=headers)
        assert saved.status_code == 200, saved.text
    return client.get("/api/customs/materials", params={"contract_no": contract_no, "pageSize": 50}, headers=headers).json()["data"]["items"]


def test_receipt_syncs_material(client):
    headers, _ = login(client)
    blocked = client.post("/api/warehouse/receipts", json={
        "receipt_no": "RK-TEST-001",
        "receipt_date": "2026-10-01",
        "contract_no": "HT-TEST",
        "packing_list_no": "PL-TEST",
        "items": [_item("M-RK-001", "深沟球轴承")],
    }, headers=headers)
    assert blocked.status_code == 400
    assert "报关物料只能来自采购入库" in blocked.json()["message"]
    mats = _chain(client, headers, "RK-TEST-001", "HT-TEST", "PL-TEST", [_item("M-RK-001", "深沟球轴承")])
    assert mats[0]["material_code"] == "M-RK-001"
    assert mats[0]["receipt_no"] in ("", None)
    assert mats[0]["wms_bill_id"]
    assert mats[0]["vessel_voyage"] == "SEA STAR/V01"
    detail = client.get(f"/api/customs/materials/{mats[0]['id']}", headers=headers).json()["data"]
    assert detail["contract_code"] == "HT-TEST"
    assert detail["inbound_code"] == "RK-TEST-001"
    assert detail["customs_invoice_no"] == ""


def test_materials_create_from_inbound_form(client):
    headers, _ = login(client)
    blocked = client.post("/api/customs/materials", json={
        "receipt_no": "RK-UI-001",
        "receipt_date": "2026-10-01",
        "contract_no": "HT-UI",
        "packing_list_no": "PL-UI",
        "items": [_item("UI1", "清单物料")],
    }, headers=headers)
    assert blocked.status_code == 400
    assert "报关物料只能来自采购入库" in blocked.json()["message"]
    mats = _chain(client, headers, "RK-UI-001", "HT-UI", "PL-UI", [_item("UI1", "清单物料")])
    assert mats[0]["material_code"] == "UI1"
    assert mats[0]["receipt_no"] in ("", None)
    assert mats[0]["receipt_date"] == "2026-10-01"
    opts = client.get("/api/customs/materials/options", headers=headers).json()["data"]
    assert "RK-UI-001" not in opts["receipt_no"]
    assert "UI1" in opts["material_code"]


def test_inbound_syncs_customs_material(client):
    headers, _ = login(client)
    bills = client.get("/api/wms-bills", params={"type": "inbound", "pageSize": 200}, headers=headers).json()["data"]["items"]
    inbound = next(row for row in bills if row["contract_code"] == "CG-2026-001")
    mats = client.get("/api/customs/materials", params={"material_code": "M001", "wms_bill_id": inbound["id"]}, headers=headers).json()["data"]["items"]
    assert mats, "采购入库应写入报关物料清单"
    assert mats[0]["receipt_no"] in ("", None)
    assert mats[0]["wms_bill_id"] == inbound["id"]
    assert mats[0]["material_code"] == "M001"
    assert mats[0]["receipt_date"]
    assert mats[0]["supplier_name"] == inbound["party_name"]


def test_material_doc_status_and_hide_filters(client):
    headers, _ = login(client)
    pending = _chain(client, headers, "RK-HIDE-001", "HT-HIDE-001", "PL-HIDE-1", [_item("HIDE1", "未报关物料")])
    done = _chain(client, headers, "RK-HIDE-002", "HT-HIDE-002", "PL-HIDE-2", [_item("HIDE2", "已报关物料")])
    fresh = client.get(f"/api/customs/materials/{pending[0]['id']}", headers=headers).json()["data"]
    assert fresh["invoice_generated"] == "未生成"
    assert fresh["clearance_generated"] == "未生成"
    assert fresh["element_generated"] == "未生成"
    assert fresh["customs_done"] == "未报关"
    ids = [done[0]["id"]]
    invoice = client.post("/api/customs/invoices", json={"ids": ids}, headers=headers)
    assert invoice.status_code == 200, invoice.text
    clearance = client.post("/api/customs/clearances", json={"ids": ids}, headers=headers)
    assert clearance.status_code == 200, clearance.text
    element = client.post("/api/customs/elements", json={"invoice_id": invoice.json()["data"]["id"], "ids": ids}, headers=headers)
    assert element.status_code == 200, element.text
    finished = client.get(f"/api/customs/materials/{done[0]['id']}", headers=headers).json()["data"]
    assert finished["invoice_generated"] == "已生成"
    assert finished["clearance_generated"] == "已生成"
    assert finished["element_generated"] == "已生成"
    assert finished["customs_done"] == "已报关"
    listed = client.get("/api/customs/materials", params={"contract_no": "HT-HIDE-002", "hide_invoice": "1"}, headers=headers).json()["data"]
    assert listed["total"] == 0
    kept = client.get("/api/customs/materials", params={"contract_no": "HT-HIDE-001", "hide_invoice": "1", "hide_clearance": "1", "hide_element": "1"}, headers=headers).json()["data"]["items"]
    assert [row["material_code"] for row in kept] == ["HIDE1"]
    options = client.get("/api/customs/materials/options", params={"hide_invoice": "1"}, headers=headers).json()["data"]
    assert "HIDE1" in options["material_code"]
    assert "HIDE2" not in options["material_code"]


def test_invoice_merge_and_special_declare(client):
    headers, _ = login(client)
    mats = _chain(client, headers, "RK-TEST-002", "HT-TEST-002", "PL-TEST", [
        _item("A1", "物料A", extra={"box_no": "B1"}),
        _item("A2", "物料A", extra={"box_no": "B2"}),
        _item("A3", "物料A", special=True, extra={"box_no": "B3"}),
    ])
    ids = [row["id"] for row in mats]
    assert len(ids) == 3
    inv = client.post("/api/customs/invoices", json={"ids": ids}, headers=headers).json()["data"]
    names = [item["name_cn"] for item in inv["items"]]
    assert names.count("轴承") == 2
    specials = [item for item in inv["items"] if item["special_declare"]]
    assert len(specials) == 1
    merged = [item for item in inv["items"] if not item["special_declare"]][0]
    assert float(merged["quantity"]) == 20


def test_export_requires_element(client, tmp_path):
    headers, _ = login(client)
    mats = _chain(client, headers, "RK-TEST-003", "HT-TEST-003", "PL-TEST", [_item("C1", "密封圈", category="密封件")])
    inv = client.post("/api/customs/invoices", json={"ids": [mats[0]["id"]]}, headers=headers).json()["data"]
    blocked = client.get(f"/api/customs/invoices/{inv['id']}/export", headers=headers)
    assert blocked.status_code == 400
    assert "尚未关联报关要素" in blocked.json()["message"]
    elem = client.post("/api/customs/elements", json={"invoice_id": inv["id"]}, headers=headers).json()["data"]
    assert elem["invoice_id"] == inv["id"]
    missing = client.get(f"/api/customs/invoices/{inv['id']}/export", headers=headers)
    assert missing.status_code == 400
    assert "导出模板" in missing.json()["message"]
    kinds = client.get("/api/export-profiles", headers=headers).json()["data"]["kinds"]
    assert {"customs_invoice", "customs_clearance", "customs_element"} <= {item["code"] for item in kinds}
    docx = tmp_path / "customs-invoice.docx"
    _write_docx(docx, "{{发票编号}} {{申报大类}}")
    saved = client.post("/api/export-profiles", json={
        "doc_code": "customs_invoice",
        "title": "测试报关发票",
        "template_path": str(docx.resolve()),
        "name_rule": "{{发票编号}}.docx",
    }, headers=headers)
    assert saved.status_code == 200, saved.text
    ok = client.get(f"/api/customs/invoices/{inv['id']}/export", headers=headers)
    assert ok.status_code == 200
    assert ok.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.wordprocessingml")


def test_export_template_accepts_xlsx_and_doc(client, tmp_path):
    from backend.app.services.export_profile import require_template_path
    from backend.app.services.form_docx import render_office

    book = openpyxl.Workbook()
    sheet = book.active
    sheet["A1"] = "发票 {{发票编号}}"
    sheet["A2"] = "{{申报大类}}"
    sheet["B2"] = "{{数量}}"
    xlsx = tmp_path / "customs.xlsx"
    book.save(xlsx)
    assert require_template_path(str(xlsx.resolve()), for_save=True).suffix == ".xlsx"
    content = render_office(
        xlsx.resolve(),
        {"发票编号": "BG-1"},
        [(("申报大类", "数量"), [{"申报大类": "钢材", "数量": "6"}, {"申报大类": "密封件", "数量": "2"}])],
    )
    filled = openpyxl.load_workbook(BytesIO(content))
    assert filled.active["A1"].value == "发票 BG-1"
    assert filled.active["A2"].value == "钢材"
    assert filled.active["A3"].value == "密封件"
    assert filled.active["B3"].value == "2"
    merged = openpyxl.Workbook()
    merged_sheet = merged.active
    merged_sheet.merge_cells("A1:C1")
    merged_sheet["A1"] = "抬头 {{发票编号}}"
    merged_sheet.merge_cells("A2:B2")
    merged_sheet["A2"] = "{{申报大类}}"
    merged_sheet["C2"] = "{{数量}}"
    merged_path = tmp_path / "merged.xlsx"
    merged.save(merged_path)
    merged_content = render_office(
        merged_path.resolve(),
        {"发票编号": "BG-9"},
        [(("申报大类", "数量"), [{"申报大类": "钢材", "数量": "1"}, {"申报大类": "密封件", "数量": "3"}])],
    )
    merged_book = openpyxl.load_workbook(BytesIO(merged_content))
    assert merged_book.active["A1"].value == "抬头 BG-9"
    assert merged_book.active["A2"].value == "钢材"
    assert merged_book.active["A3"].value == "密封件"
    assert merged_book.active["C3"].value == "3"
    for suffix in (".doc", ".xls"):
        sample = tmp_path / f"template{suffix}"
        sample.write_bytes(b"placeholder")
        assert require_template_path(str(sample.resolve())).suffix == suffix
    headers, _ = login(client)
    saved = client.post("/api/export-profiles", json={
        "doc_code": "customs_invoice",
        "title": "表格模板",
        "template_path": str(xlsx.resolve()),
        "name_rule": "{{发票编号}}.docx",
    }, headers=headers)
    assert saved.status_code == 200, saved.text
    listed = client.get("/api/export-profiles", params={"doc_code": "customs_invoice"}, headers=headers)
    paths = [item["template_path"].lower() for item in listed.json()["data"]["items"]]
    assert any(path.endswith(".xlsx") for path in paths)


def test_element_consistency_issues(client):
    headers, _ = login(client)
    client.post("/api/customs/hs", json={"hs_code": "84821000", "declaration_name": "轴承", "declaration_elements": "品牌;型号"}, headers=headers)
    mats = _chain(client, headers, "RK-TEST-004", "HT-TEST-004", "PL-TEST", [_item("D1", "轴承A"), _item("D2", "轴承B")])
    inv = client.post("/api/customs/invoices", json={"ids": [row["id"] for row in mats]}, headers=headers).json()["data"]
    elem = client.post("/api/customs/elements", json={"invoice_id": inv["id"]}, headers=headers).json()["data"]
    items = elem["items"]
    assert sorted(line["product_name"] for line in items[0]["lines"]) == ["轴承A", "轴承B"]
    assert {line["model"] for line in items[0]["lines"]} == {"6205"}
    assert {line["origin_place"] for line in items[0]["lines"]} == {"中国"}
    assert items[0]["lines"][0]["declaration_elements"]
    assert len({line["declaration_elements"] for line in items[0]["lines"]}) == 1
    items[0]["declaration_name"] = "轴承座"
    failed = client.put(f"/api/customs/elements/{elem['id']}", json={"version": elem["version"], "items": items}, headers=headers)
    assert failed.status_code == 400
    assert any("申报名称不一致" in text for text in failed.json()["data"]["issues"])
    confirmed = client.put(f"/api/customs/elements/{elem['id']}", json={"version": elem["version"], "items": items, "confirm": True}, headers=headers)
    assert confirmed.status_code == 200


def test_excel_import_all_or_nothing(client):
    headers, _ = login(client)
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(["入仓单号", "入仓日期", "合同号", "箱单号", "物料大类", "物料编号", "中文名称", "数量", "产地", "使用设备", "毛重", "净重", "长", "宽", "高", "体积"])
    sheet.append(["RK-XLS-1", "2026-10-01", "HT-X", "PL-X", "轴承", "X1", "好物料", "2", "中国", "包装机", "1", "1", "1", "1", "1", "1"])
    sheet.append(["RK-XLS-1", "2026-10-01", "HT-X", "PL-X", "轴承", "", "空编号", "2", "中国", "包装机", "1", "1", "1", "1", "1", "1"])
    buf = BytesIO()
    book.save(buf)
    parsed = client.post(
        "/api/import/warehouse-receipt/parse",
        files={"file": ("bad.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=headers,
    )
    assert parsed.status_code == 200
    data = parsed.json()["data"]
    assert data["batch_error"]
    executed = client.post("/api/import/warehouse-receipt/execute", json=data, headers=headers)
    assert executed.status_code == 400
    rows = client.get("/api/warehouse/receipts", params={"keyword": "RK-XLS-1"}, headers=headers).json()["data"]["items"]
    assert rows == []


def test_batch_assign_and_api_import(client):
    headers, _ = login(client)
    imported = client.post("/api/warehouse/receipts/import", json={
        "receiptNo": "RK-API-001",
        "contractNo": "HT-API",
        "packingListNo": "PL-API",
        "receiptDate": "2026-10-01",
        "items": [_item("API1", "进口物料")],
    }, headers=headers)
    assert imported.status_code == 400
    assert "报关物料只能来自采购入库" in imported.json()["message"]
    mats = _chain(client, headers, "RK-API-001", "HT-API", "PL-API", [_item("API1", "进口物料")])
    assert mats
    assigned = client.post("/api/customs/materials/batch-assign", json={"ids": [mats[0]["id"]], "field": "containerNo", "value": "CONT999"}, headers=headers)
    assert assigned.status_code == 200
    assert assigned.json()["data"]["count"] == 1
    again = client.get(f"/api/customs/materials/{mats[0]['id']}", headers=headers).json()["data"]
    assert again["container_no"] == "CONT999"
    special = client.post("/api/customs/materials/batch-assign", json={"ids": [mats[0]["id"]], "field": "special_declare", "value": "特殊申报"}, headers=headers)
    assert special.status_code == 200
    marked = client.get(f"/api/customs/materials/{mats[0]['id']}", headers=headers).json()["data"]
    assert marked["special_declare"] is True
    plain = client.post("/api/customs/materials/batch-assign", json={"ids": [mats[0]["id"]], "field": "special_declare", "value": "普通"}, headers=headers)
    assert plain.status_code == 200
    cleared = client.get(f"/api/customs/materials/{mats[0]['id']}", headers=headers).json()["data"]
    assert cleared["special_declare"] is False


def test_demo_materials_run_full_chain(client):
    headers, _ = login(client)
    r1 = client.get("/api/customs/materials", params={"contract_no": "HT-BG-2026-001", "pageSize": 50}, headers=headers).json()["data"]["items"]
    r2 = client.get("/api/customs/materials", params={"contract_no": "HT-BG-2026-002", "pageSize": 50}, headers=headers).json()["data"]["items"]
    r3 = client.get("/api/customs/materials", params={"contract_no": "HT-BG-2026-003", "pageSize": 50}, headers=headers).json()["data"]["items"]
    assert len(r1) == 3 and len(r2) == 3 and len(r3) == 2
    assert all(row.get("wms_bill_id") for row in r1 + r2 + r3)
    contracts = client.get("/api/contracts", params={"direction": "purchase", "keyword": "HT-BG-2026-002"}, headers=headers).json()["data"]["items"]
    assert contracts and contracts[0]["contract_code"] == "HT-BG-2026-002"
    bills = client.get("/api/wms-bills", params={"type": "inbound", "pageSize": 200}, headers=headers).json()["data"]["items"]
    bill = next(row for row in bills if row["bill_code"] == "RK-BG-002")
    assert bill["contract_code"] == "HT-BG-2026-002"
    detail = client.get(f"/api/customs/materials/{r2[0]['id']}", headers=headers).json()["data"]
    assert detail["contract_code"] == "HT-BG-2026-002"
    assert detail["inbound_code"] == "RK-BG-002"
    assert detail["customs_invoice_no"] == "BG-DEMO-001"
    assert detail["clearance_invoice_no"] == "QG-DEMO-001"
    assert all(row.get("invoice_id") for row in r2)
    invoices = client.get("/api/customs/invoices", headers=headers).json()["data"]["items"]
    demo_inv = next(row for row in invoices if row["invoice_no"] == "BG-DEMO-001")
    assert demo_inv["supplier_name"]
    assert demo_inv["main_materials"] == "钢材、密封件"
    clearances = client.get("/api/customs/clearances", headers=headers).json()["data"]["items"]
    demo_clr = next(row for row in clearances if row.get("invoice_no") == "QG-DEMO-001")
    assert demo_clr["supplier_name"] == demo_inv["supplier_name"]
    assert demo_clr["main_materials"] == "钢材、密封件"
    assert demo_clr["contract_no"]
    elements = client.get("/api/customs/elements", headers=headers).json()["data"]["items"]
    demo_el = next(row for row in elements if row["invoice_no"] == "BG-DEMO-001")
    assert demo_el["supplier_name"]
    assert demo_el["main_materials"] == "钢材、密封件"
    assert demo_el["contract_no"]
    inv = client.post("/api/customs/invoices", json={"ids": [row["id"] for row in r1]}, headers=headers)
    assert inv.status_code == 200, inv.text
    names = [item["name_cn"] for item in inv.json()["data"]["items"]]
    assert names.count("轴承") == 2
    elem = client.post("/api/customs/elements", json={"invoice_id": inv.json()["data"]["id"]}, headers=headers)
    assert elem.status_code == 200, elem.text
    codes = {item["hs_code"] for item in elem.json()["data"]["items"]}
    assert "84821000" in codes
    clr = client.post("/api/customs/clearances", json={"ids": [row["id"] for row in r1], "invoice_date": "2026-10-01"}, headers=headers)
    assert clr.status_code == 200, clr.text
    assert float(clr.json()["data"]["total_usd"] or 0) > 0


def test_invoice_pages_and_excel(client):
    headers, _ = login(client)
    mats = _chain(client, headers, "RK-PAGE-001", "HT-PAGE", "PL-PAGE", [
        _item("P1", "托辊", category="托辊", extra={"container_no": "BOX-A", "box_no": "1#", "length": "1000", "width": "800", "height": "600", "material_name_en": "ROLL", "specification": "Φ127", "port_of_shipment": "NINGBO, CHINA", "hs_code": "84821000", "declaration_elements": "品牌;型号", "remark": "加急"}),
        _item("P2", "螺母", category="螺母", extra={"container_no": "BOX-B", "box_no": "2#", "length": "10", "width": "10", "height": "10", "material_name_en": "nut"}),
    ])
    inv = client.post("/api/customs/invoices", json={"ids": [row["id"] for row in mats]}, headers=headers).json()["data"]
    assert inv["invoice_date_2"] == inv["invoice_date_1"]
    p1 = next(row for row in mats if row["material_code"] == "P1")
    assert p1["hs_code"] == "84821000"
    assert p1["remark"] == "加急"
    assert inv["port_of_shipment"] == "NINGBO, CHINA"
    assert inv["customs_elements"] == "品牌;型号"
    assert [row["size_display"] for row in inv["categories"]] == ["1000*800*600", "10*10*10"]
    assert [item["label"] for item in inv["specification_options"]] == ["Specification # 1", "Specification # 2"]
    assert inv["detail_groups"][0]["items"][0]["seq"] == 1
    assert inv["detail_groups"][1]["items"][0]["seq"] == 1
    kept = inv["invoice_date_2"]
    saved = client.put(f"/api/customs/invoices/{inv['id']}", json={
        "version": inv["version"],
        "customs_elements": "品牌",
        "invoice_date_1": "2026-10-02",
    }, headers=headers).json()["data"]
    assert saved["invoice_date_2"] == kept
    assert saved["invoice_date_1"] == "2026-10-02"
    assert saved["customs_elements"] == "品牌"
    from backend.app.database import SessionLocal
    from backend.app.models import CustomsInvoice
    db = SessionLocal()
    try:
        row = db.get(CustomsInvoice, inv["id"])
        row.status = "已导出"
        db.commit()
        version = row.version
    finally:
        db.close()
    exported = client.put(f"/api/customs/invoices/{inv['id']}", json={
        "version": version,
        "customs_elements": "仍可改",
    }, headers=headers)
    assert exported.status_code == 200, exported.text
    assert exported.json()["data"]["status"] == "已导出"
    assert exported.json()["data"]["customs_elements"] == "仍可改"
    excel = client.get(f"/api/customs/invoices/{inv['id']}/excel", headers=headers)
    assert excel.status_code == 200, excel.text
    book = openpyxl.load_workbook(BytesIO(excel.content))
    assert book.sheetnames == ["PL", "Commercial Invoice", "明细表"]
    assert book["PL"]["H8"].value == "PL-PAGE"
    assert book["PL"]["H9"].value == kept
    assert book["PL"]["I17"].value == "1000*800*600"
    assert "{{" not in str(book["Commercial Invoice"]["F4"].value)
    assert book["Commercial Invoice"]["D10"].value == "NINGBO, CHINA"
    titles = [book["明细表"].cell(row, 1).value for row in range(1, book["明细表"].max_row + 1)]
    assert "Specification # 1" in titles and "Specification # 2" in titles


def test_clearance_pages_stay_editable(client):
    headers, _ = login(client)
    mats = _chain(client, headers, "RK-CLR-001", "HT-CLR", "PL-CLR", [
        _item("Q1", "托辊", category="托辊", extra={"remark": "先发"}),
        _item("Q2", "螺母", category="螺母"),
    ])
    for row in mats:
        detail = client.get(f"/api/customs/materials/{row['id']}", headers=headers).json()["data"]
        detail["receipt_no"] = "RC-CLR-9"
        saved = client.put(f"/api/customs/materials/{row['id']}", json=detail, headers=headers)
        assert saved.status_code == 200, saved.text
    created = client.post("/api/customs/clearances", json={"ids": [row["id"] for row in mats]}, headers=headers)
    assert created.status_code == 200, created.text
    data = created.json()["data"]
    assert data["port_of_shipment"] == "SHANGHAI, CHINA"
    assert data["packing_list_no"] == "PL-CLR"
    assert data["items"][0]["unit"] == "套"
    assert any(row["remark"] == "先发" for row in data["items"])
    assert data["items"][0]["price"]
    assert sorted(row["category"] for row in data["packing"]) == ["托辊", "螺母"]
    assert {row["receipt_no"] for row in data["packing"]} == {"RC-CLR-9"}
    from backend.app.database import SessionLocal
    from backend.app.models import CustomsClearance
    db = SessionLocal()
    try:
        row = db.get(CustomsClearance, data["id"])
        row.status = "已导出"
        db.commit()
        version = row.version
    finally:
        db.close()
    saved = client.put(f"/api/customs/clearances/{data['id']}", json={
        "version": version,
        "invoice_no": "",
        "port_of_shipment": "NINGBO, CHINA",
        "status": "已导出",
        "items": data["items"],
        "packing": data["packing"],
    }, headers=headers)
    assert saved.status_code == 200, saved.text
    body = saved.json()["data"]
    assert body["status"] == "已导出"
    assert body["invoice_no"] == ""
    assert body["invoice_missing"] is True
    assert body["port_of_shipment"] == "NINGBO, CHINA"


def _xlsx(rows):
    book = openpyxl.Workbook()
    sheet = book.active
    for row in rows:
        sheet.append(row)
    buf = BytesIO()
    book.save(buf)
    return buf.getvalue()


def _fill_file(headers, rows):
    return {"file": ("fill.xlsx", _xlsx([headers, *rows]), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}


FILL_HEAD = ["物料编号", "入仓单号", "产地", "箱号", "尺寸", "净重", "毛重", "体积", "出发港"]


def test_material_fields_are_in_catalog():
    from openpyxl import load_workbook
    from backend.app.services.export_profile import build_field_workbook

    book = load_workbook(BytesIO(build_field_workbook()))
    assert "报关物料清单" in book.sheetnames
    sheet = book["报关物料清单"]
    names = [row[2] for row in sheet.iter_rows(min_row=2, values_only=True)]
    for label in ("英文名称", "规格", "毛重", "净重", "长", "宽", "高", "体积", "出发港", "收货人", "询价名称", "申报要素", "报关发票", "清关发票", "报关要素", "报关完毕", "入仓单号"):
        assert label in names
    assert "入库单号" not in names


def test_material_filter_options_follow_selection(client):
    headers, _ = login(client)
    scoped = client.get("/api/customs/materials/options", params={"material_category": "钢材"}, headers=headers).json()["data"]
    assert {"BG-S01", "BG-S02"} <= set(scoped["material_code"])
    for unrelated in ("BG-M01", "BG-F01", "BG-B01"):
        assert unrelated not in scoped["material_code"]
    assert "HT-BG-2026-002" in scoped["contract_no"]
    assert "HT-BG-2026-001" not in scoped["contract_no"]
    assert "HT-BG-2026-003" not in scoped["contract_no"]
    assert "RK-BG-003" not in scoped["receipt_no"]
    assert "产线A" in scoped["equipment"]
    assert "输送线" not in scoped["equipment"]
    assert "包装机" not in scoped["equipment"]
    assert "电机" in scoped["material_category"]
    both = client.get("/api/customs/materials/options", params={"contract_no": "HT-BG-2026-002", "material_category": "钢材"}, headers=headers).json()["data"]
    assert "BG-F01" not in both["material_code"]
    assert "BG-S01" in both["material_code"]
    assert "密封件" in both["material_category"]
    assert "电机" not in both["material_category"]
    one = client.get("/api/customs/materials/options", params={"material_code": "BG-M01"}, headers=headers).json()["data"]
    assert one["contract_no"] == ["HT-BG-2026-003"]
    assert one["receipt_no"] == []
    assert one["material_category"] == ["电机"]
    assert "钢材" not in one["material_name"]
    assert "三相异步电机" in one["material_name"]


def test_material_fill_import(client):
    headers, _ = login(client)
    template = client.get("/api/import/customs-material-fill/template", headers=headers)
    assert template.status_code == 200
    assert template.content[:2] == b"PK"
    template_book = openpyxl.load_workbook(BytesIO(template.content))
    assert template_book.active.cell(1, 2).value == "入仓单号"
    mats = _chain(client, headers, "RK-FILL-001", "HT-FILL-001", "PL-FILL", [_item("FILL1", "补录物料")])
    first_no = client.post("/api/import/customs-material-fill/parse", files=_fill_file(FILL_HEAD, [
        ["FILL1", "RC-NEW", "美国", "BX-1", "", "", "", "", ""],
    ]), headers=headers).json()["data"]
    assert first_no["rows"][0]["ok"] is True
    assert client.post("/api/import/customs-material-fill/execute", json=first_no, headers=headers).status_code == 200
    named_blank = client.get(f"/api/customs/materials/{mats[0]['id']}", headers=headers).json()["data"]
    assert named_blank["receipt_no"] == "RC-NEW"
    assert named_blank["origin_country"] == "美国"
    named_blank["receipt_no"] = ""
    named_blank["origin_country"] = "中国"
    cleared = client.put(f"/api/customs/materials/{mats[0]['id']}", json=named_blank, headers=headers)
    assert cleared.status_code == 200, cleared.text
    before = client.get(f"/api/customs/materials/{mats[0]['id']}", headers=headers).json()["data"]
    assert before["origin_country"] == "中国"
    parsed = client.post("/api/import/customs-material-fill/parse", files=_fill_file(FILL_HEAD, [
        ["FILL1", "", "德国", "BX-9", "40*30*20", "8", "9", "0.024", "NINGBO, CHINA"],
    ]), headers=headers)
    assert parsed.status_code == 200, parsed.text
    data = parsed.json()["data"]
    assert data["rows"][0]["ok"] is True
    assert "将更新 1 条" in data["rows"][0]["message"]
    done = client.post("/api/import/customs-material-fill/execute", json=data, headers=headers)
    assert done.status_code == 200, done.text
    assert done.json()["data"]["success"] == 1
    after = client.get(f"/api/customs/materials/{mats[0]['id']}", headers=headers).json()["data"]
    assert after["origin_country"] == "德国"
    assert after["box_no"] == "BX-9"
    assert after["length"] == "40"
    assert after["width"] == "30"
    assert after["height"] == "20"
    assert after["net_weight"] == "8"
    assert after["gross_weight"] == "9"
    assert after["volume"] == "0.024"
    assert after["port_of_shipment"] == "NINGBO, CHINA"
    assert after["material_name"] == "补录物料"
    after["receipt_no"] = "RC-FILL-9"
    named = client.put(f"/api/customs/materials/{mats[0]['id']}", json=after, headers=headers)
    assert named.status_code == 200, named.text
    missed = client.post("/api/import/customs-material-fill/parse", files=_fill_file(FILL_HEAD, [
        ["FILL1", "RK-FILL-001", "", "BX-10", "", "", "", "", ""],
    ]), headers=headers).json()["data"]
    assert missed["rows"][0]["ok"] is False
    assert "入仓单" in missed["rows"][0]["message"]
    kept = client.post("/api/import/customs-material-fill/parse", files=_fill_file(FILL_HEAD, [
        ["FILL1", "RC-FILL-9", "", "BX-10", "", "", "", "", ""],
    ]), headers=headers).json()["data"]
    assert kept["rows"][0]["ok"] is True
    applied = client.post("/api/import/customs-material-fill/execute", json=kept, headers=headers)
    assert applied.status_code == 200, applied.text
    latest = client.get(f"/api/customs/materials/{mats[0]['id']}", headers=headers).json()["data"]
    assert latest["origin_country"] == "德国"
    assert latest["box_no"] == "BX-10"
    assert latest["port_of_shipment"] == "NINGBO, CHINA"
    blocked = client.post("/api/import/customs-material-fill/parse", files=_fill_file(FILL_HEAD, [
        ["FILL1", "", "日本", "BX-11", "1*1*1", "1", "1", "1", "SHANGHAI"],
        ["NO-SUCH", "", "日本", "BX-12", "1*1*1", "1", "1", "1", "SHANGHAI"],
    ]), headers=headers).json()["data"]
    assert blocked["batch_error"]
    assert blocked["rows"][1]["ok"] is False
    refused = client.post("/api/import/customs-material-fill/execute", json=blocked, headers=headers)
    assert refused.status_code == 400
    unchanged = client.get(f"/api/customs/materials/{mats[0]['id']}", headers=headers).json()["data"]
    assert unchanged["origin_country"] == "德国"
    assert unchanged["box_no"] == "BX-10"
    bad_size = client.post("/api/import/customs-material-fill/parse", files=_fill_file(FILL_HEAD, [
        ["FILL1", "", "", "", "30*20", "", "", "", ""],
    ]), headers=headers).json()["data"]
    assert bad_size["rows"][0]["ok"] is False
    assert "尺寸" in bad_size["rows"][0]["message"]
    duplicated = client.post("/api/import/customs-material-fill/parse", files=_fill_file(FILL_HEAD, [
        ["FILL1", "", "日本", "", "", "", "", "", ""],
        ["FILL1", "", "美国", "", "", "", "", "", ""],
    ]), headers=headers).json()["data"]
    assert all(not row["ok"] for row in duplicated["rows"])
    assert "重复" in duplicated["rows"][0]["message"]


def test_form_preview_waits_for_review(client):
    headers, _ = login(client)
    mats = _chain(client, headers, "RK-DRAFT-001", "HT-DRAFT-001", "PL-DRAFT", [
        _item("DRAFT1", "待审物料"),
    ])
    ids = [row["id"] for row in mats]
    preview = client.post("/api/customs/invoices/preview", json={"ids": ids}, headers=headers)
    assert preview.status_code == 200, preview.text
    draft = preview.json()["data"]
    assert draft["draft"] is True
    assert draft["id"] is None
    assert draft["invoice_no"]
    assert draft["items"]
    material = client.get(f"/api/customs/materials/{ids[0]}", headers=headers).json()["data"]
    assert material["customs_invoice_no"] == ""
    created = client.post("/api/customs/invoices", json={"ids": ids, "invoice_no": draft["invoice_no"]}, headers=headers)
    assert created.status_code == 200, created.text
    assert created.json()["data"]["invoice_no"] == draft["invoice_no"]
    clearance = client.post("/api/customs/clearances/preview", json={"ids": ids}, headers=headers)
    assert clearance.status_code == 200, clearance.text
    assert clearance.json()["data"]["draft"] is True
    assert clearance.json()["data"]["items"]
    still = client.get(f"/api/customs/materials/{ids[0]}", headers=headers).json()["data"]
    assert not still.get("clearance_invoice_no")
    element = client.post("/api/customs/elements/preview", json={"invoice_id": created.json()["data"]["id"], "ids": ids}, headers=headers)
    assert element.status_code == 200, element.text
    assert element.json()["data"]["draft"] is True
    assert element.json()["data"]["items"]
    invoice = client.get(f"/api/customs/invoices/{created.json()['data']['id']}", headers=headers).json()["data"]
    assert invoice["element_status"] == "未关联"


def test_unused_invoices_follow_selected_materials(client):
    headers, _ = login(client)
    left = _chain(client, headers, "RK-REL-001", "HT-REL-001", "PL-REL-1", [_item("REL1", "关联物料甲")])
    right = _chain(client, headers, "RK-REL-002", "HT-REL-002", "PL-REL-2", [_item("REL2", "关联物料乙")])
    inv_left = client.post("/api/customs/invoices", json={"ids": [left[0]["id"]]}, headers=headers).json()["data"]
    inv_right = client.post("/api/customs/invoices", json={"ids": [right[0]["id"]]}, headers=headers).json()["data"]
    scoped = client.get("/api/customs/invoices/unused", params={"material_ids": str(left[0]["id"])}, headers=headers)
    assert scoped.status_code == 200, scoped.text
    scoped_ids = {item["id"] for item in scoped.json()["data"]}
    assert inv_left["id"] in scoped_ids
    assert inv_right["id"] not in scoped_ids
    whole = {item["id"] for item in client.get("/api/customs/invoices/unused", headers=headers).json()["data"]}
    assert inv_left["id"] in whole and inv_right["id"] in whole


def test_hs_product_name_can_be_empty(client):
    headers, _ = login(client)
    created = client.post("/api/customs/hs", json={"hs_code": "99990001", "declaration_name": "空产品名"}, headers=headers)
    assert created.status_code == 200, created.text
    listed = client.get("/api/customs/hs", params={"keyword": "空产品名", "pageSize": 20}, headers=headers)
    assert listed.status_code == 200, listed.text
    item = listed.json()["data"]["items"][0]
    assert item["product_name"] == ""
    saved = client.put(
        f"/api/customs/hs/{item['id']}",
        json={**item, "product_name": "深沟球轴承"},
        headers=headers,
    )
    assert saved.status_code == 200, saved.text
    again = client.get("/api/customs/hs", params={"keyword": "深沟球轴承", "pageSize": 20}, headers=headers)
    assert again.json()["data"]["items"][0]["product_name"] == "深沟球轴承"
    blank = client.put(
        f"/api/customs/hs/{item['id']}",
        json={**again.json()["data"]["items"][0], "product_name": ""},
        headers=headers,
    )
    assert blank.status_code == 200, blank.text


def test_deepseek_setting_drives_check(client):
    headers, _ = login(client)
    got = client.get("/api/settings/deepseek", headers=headers)
    assert got.status_code == 200, got.text
    data = got.json()["data"]
    assert data["api_base"] == "https://api.deepseek.com"
    assert data["model"] == "deepseek-v4-flash"
    assert data["enabled"] is True
    assert data["api_key"].startswith("sk-")
    assert data["thinking"] is False
    saved = client.put("/api/settings/deepseek", json={**data, "timeout_seconds": 25, "model": "deepseek-flash"}, headers=headers)
    assert saved.status_code == 200, saved.text
    body = saved.json()["data"]
    assert body["model"] == "deepseek-flash"
    assert body["timeout_seconds"] == 25
    off = client.put("/api/settings/deepseek", json={**body, "enabled": False}, headers=headers)
    assert off.status_code == 200, off.text
    check = client.post("/api/ai/customs/check", json={"declarationName": "", "hsCode": "", "elements": ""}, headers=headers)
    assert check.status_code == 200, check.text
    assert check.json()["data"]["source"] == "local"
    restored = client.put(
        "/api/settings/deepseek",
        json={**off.json()["data"], "enabled": True, "model": "deepseek-v4-flash", "timeout_seconds": 30},
        headers=headers,
    )
    assert restored.status_code == 200, restored.text
    assert restored.json()["data"]["model"] == "deepseek-v4-flash"


def test_customs_templates_merge_category_and_keep_specials():
    from pathlib import Path

    from backend.app.services.customs_xlsx import export_element_xlsx, export_invoice_xlsx

    invoice_template = Path(r"d:\Python\Cursor\需求的模板\报关发票.xlsx")
    element_template = Path(r"d:\Python\Cursor\需求的模板\报关要素.xlsx")
    if not invoice_template.is_file() or not element_template.is_file():
        pytest.skip("本机模板不在预期目录")
    categories = [
        {"material_category": "非公路用自卸车零配件", "specification": "", "quantity": "598", "unit": "件", "package_count": 1, "gross_weight": "10", "net_weight": "9", "size_display": "600*200*10", "weight": "10", "name_en": "", "price": "12.31", "amount": "7358.68"},
        {"material_category": "对开法兰", "specification": "9409984", "quantity": "16", "unit": "件", "package_count": 1, "gross_weight": "1", "net_weight": "1", "size_display": "1*1*1", "weight": "1", "name_en": "split flange", "price": "27.03", "amount": "432.48"},
        {"material_category": "对开法兰", "specification": "09409985", "quantity": "32", "unit": "件", "package_count": 1, "gross_weight": "1", "net_weight": "1", "size_display": "1*1*1", "weight": "1", "name_en": "split flange", "price": "27.03", "amount": "864.96"},
        {"material_category": "对开法兰", "specification": "09167602", "quantity": "40", "unit": "件", "package_count": 1, "gross_weight": "1", "net_weight": "1", "size_display": "1*1*1", "weight": "1", "name_en": "split flange", "price": "41.58", "amount": "1663.20"},
    ]
    content, _name = export_invoice_xlsx({
        "invoice_no": "BG-20261002-003",
        "invoice_date_1": "2026-10-02",
        "invoice_date_2": "2026-10-02",
        "packing_list_no": "PL-1",
        "vessel_voyage": "SHIP/V1",
        "container_no": "CONT1",
        "categories": categories,
        "detail_groups": [],
    }, template_path=invoice_template)
    book = openpyxl.load_workbook(BytesIO(content))
    pl = book["PL"]
    consignee_rows = [cell.row for row in pl.iter_rows() for cell in row if isinstance(cell.value, str) and cell.value.startswith("收货人")]
    assert len(consignee_rows) == 1
    names = [pl.cell(row, 2).value for row in range(16, 20)]
    assert names == ["非公路用自卸车零配件", "对开法兰", "对开法兰", "对开法兰"]
    assert pl.cell(16, 3).value in ("", None)
    assert pl.cell(17, 3).value == "9409984"
    assert "{{" not in "".join(str(pl.cell(row, col).value or "") for row in range(7, 21) for col in range(1, 11))
    invoice = book["Commercial Invoice"]
    assert [invoice.cell(row, 2).value for row in range(15, 19)] == names
    assert not any(invoice.row_dimensions[row].hidden for row in range(15, 19))
    assert invoice.cell(15, 3).value in ("", None)
    assert invoice.cell(16, 3).value == "split flange"
    element = export_element_xlsx({
        "items": [
            {
                "hs_code": "8708993900",
                "declaration_name": "非公路用自卸车零配件",
                "declaration_elements": "同一套申报要素",
                "special_declare": False,
                "lines": [
                    {"product_name": "螺栓", "model": "00179887", "origin_place": "中国"},
                    {"product_name": "垫圈", "model": "15235606", "origin_place": "中国"},
                ],
            },
            {
                "hs_code": "7307910000",
                "declaration_name": "对开法兰",
                "declaration_elements": "法兰申报要素",
                "special_declare": True,
                "lines": [{"product_name": "对开法兰", "model": "9409984", "origin_place": "中国"}],
            },
        ],
    }, element_template)
    sheet = openpyxl.load_workbook(BytesIO(element)).active
    assert sheet.cell(2, 2).value == "非公路用自卸车零配件"
    assert sheet.cell(2, 3).value == "螺栓"
    assert sheet.cell(3, 3).value == "垫圈"
    assert sheet.cell(3, 2).value in ("", None)
    assert sheet.cell(4, 2).value == "对开法兰"
    assert sheet.cell(4, 3).value == "对开法兰"
    assert sheet.cell(2, 5).value == "同一套申报要素"
    assert "明细物料" not in "".join(str(sheet.cell(row, col).value or "") for row in range(2, 8) for col in range(1, 7))
    merged = {str(item) for item in sheet.merged_cells.ranges}
    assert "B2:B3" in merged
    assert "E2:E3" in merged


def test_clearance_template_numbers_rows_and_writes_totals():
    from pathlib import Path

    from backend.app.services.customs_xlsx import export_clearance_xlsx, usd_amount_words

    template = Path(r"d:\Python\Cursor\需求的模板\清关发票.xlsx")
    if not template.is_file():
        pytest.skip("本机清关发票模板不在预期目录")
    assert usd_amount_words("3673.85") == "USD THREE THOUSAND SIX HUNDRED SEVENTY-THREE AND EIGHTY-FIVE CENTS ONLY"
    content = export_clearance_xlsx({
        "invoice_no": "QGFP-1",
        "invoice_date": "2026-10-02",
        "vessel_voyage": "YM WISH/V03",
        "container_no": "TCLU0001001",
        "bill_no": "COSU0001001",
        "port_of_shipment": "SHANGHAI, CHINA",
        "packing_list_no": "PL-1",
        "exchange_rate": "7.12",
        "items": [
            {"name_cn": "螺栓", "name_en": "bolt", "specification": "00179887", "quantity": "100", "unit": "件", "price": "2.08", "amount": "208.00", "amount_usd": "0", "po_no": "PO-1", "remark": ""},
            {"name_cn": "对开法兰", "name_en": "split flange", "specification": "9409984", "quantity": "16", "unit": "件", "price": "27.03", "amount": "432.48", "amount_usd": "0", "po_no": "PO-1", "remark": ""},
        ],
        "packing": [
            {"category": "非公路用自卸车零配件", "specification": "", "quantity": "598", "unit": "件", "packages": 1, "gross_weight": "10", "net_weight": "9", "size_text": "600*200*10", "receipt_no": "RK-1"},
            {"category": "对开法兰", "specification": "9409984", "quantity": "16", "unit": "件", "packages": 1, "gross_weight": "2", "net_weight": "1", "size_text": "1*1*1", "receipt_no": "RK-1"},
        ],
    }, template)
    book = openpyxl.load_workbook(BytesIO(content))
    ci = book["CI"]
    assert [ci.cell(row, 1).value for row in range(16, 18)] == [1, 2]
    assert ci.cell(16, 2).value == "螺栓"
    assert ci.cell(17, 2).value == "对开法兰"
    assert ci.cell(18, 5).value == 116
    assert ci.cell(18, 8).value == 640.48
    assert ci.cell(19, 1).value.strip().startswith("TOTAL AMOUNT")
    assert ci.cell(19, 4).value == usd_amount_words(round(round(208 / 7.12, 2) + round(432.48 / 7.12, 2), 2))
    assert not any(merged.min_row <= row <= merged.max_row and merged.min_col <= 2 <= merged.max_col for row in (16, 17) for merged in ci.merged_cells.ranges)
    pl = book["PL"]
    assert [pl.cell(row, 1).value for row in range(11, 13)] == [1, 2]
    assert pl.cell(11, 2).value == "非公路用自卸车零配件"
    assert pl.cell(12, 2).value == "对开法兰"
    assert pl.cell(13, 4).value == 614
    assert pl.cell(13, 6).value == 2
