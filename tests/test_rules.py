import os
from decimal import Decimal
from io import BytesIO
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[1]
os.environ["CONTRACT_DB"] = str(ROOT / "data" / "test_rules.db")
os.environ["CONTRACT_DEPLOY_MODE"] = "demo"

if Path(os.environ["CONTRACT_DB"]).exists():
    Path(os.environ["CONTRACT_DB"]).unlink()

from backend.app import config
from backend.app.money import line_ex_tax, line_tax, line_total, split_line

config.DB_PATH = os.environ["CONTRACT_DB"]

from fastapi.testclient import TestClient
from backend.app.main import app


@pytest.fixture(scope="module")
def client():
    with TestClient(app) as c:
        yield c


def login(client, name="admin", password="123456"):
    res = client.post("/api/auth/login", json={"login_name": name, "password": password})
    assert res.status_code == 200, res.text
    data = res.json()["data"]
    return {"Authorization": f"Bearer {data['token']}"}, data


def test_move_label_states():
    from backend.app.services.calc import move_label

    assert move_label("purchase", []) == "未到货"
    assert move_label("purchase", [(0, 0)]) == "未到货"
    assert move_label("purchase", [(10, 0), (5, 0)]) == "未到货"
    assert move_label("purchase", [(10, 4), (5, 0)]) == "部分到货"
    assert move_label("purchase", [(10, 10), (5, 1)]) == "部分到货"
    assert move_label("purchase", [(10, 10), (5, 5)]) == "已到货"
    assert move_label("purchase", [(10, 12)]) == "已到货"
    assert move_label("sale", [(1, 0)]) == "未发货"
    assert move_label("sale", [(2, 1)]) == "部分发货"
    assert move_label("sale", [(1, 1)]) == "已发货"


def test_money_chain_no_cent():
    samples = [(1, "0.88", 13), (3, "1", 13), (10, "5000", 13), ("0.3", "0.1", 6), (7, "123.456", 9)]
    for qty, price, rate in samples:
        part = split_line(qty, price, rate)
        assert part["ex_tax"] + part["tax"] == part["total"]
        assert part["total"] == line_total(qty, price)
        assert part["ex_tax"] == line_ex_tax(part["total"], rate)
        assert part["tax"] == line_tax(part["total"], part["ex_tax"])


def test_header_ex_tax_uses_whole_amount():
    from backend.app.services.calc import compute_amount_from_payload

    class Fake:
        pricing = "fixed"
        tax_rate = Decimal("13")

    amount, ex_tax = compute_amount_from_payload(Fake(), [
        {"material_code": "A", "qty": "2", "price": "62"},
        {"material_code": "B", "qty": "1", "price": "3596"},
    ], {})
    assert ex_tax == line_ex_tax(amount, 13)
    assert ex_tax == (amount / Decimal("1.13")).quantize(Decimal("0.01"))


def test_login_and_company_isolation(client):
    headers, data = login(client)
    assert data["need_choose_company"] is True
    companies = data["user"]["companies"]
    other = [c for c in companies if c["id"] != data["user"]["company_id"]][0]
    star = client.get("/api/contracts", params={"direction": "purchase"}, headers=headers).json()["data"]["items"]
    switched = client.post("/api/companies/switch", json={"company_id": other["id"]}, headers=headers).json()["data"]
    headers2 = {"Authorization": f"Bearer {switched['token']}"}
    other_rows = client.get("/api/contracts", params={"direction": "purchase"}, headers=headers2).json()["data"]["items"]
    assert star and other_rows
    assert {r["name"] for r in star} != {r["name"] for r in other_rows} or star[0]["party_name"] != other_rows[0]["party_name"]


def test_creator_and_version(client):
    admin, _ = login(client)
    zhang, _ = login(client, "zhang", "123456")
    rows = client.get("/api/contracts", params={"direction": "purchase"}, headers=admin).json()["data"]["items"]
    row = rows[0]
    detail = client.get(f"/api/contracts/{row['id']}", headers=admin).json()["data"]
    detail["name"] = detail["name"] + "改"
    denied = client.put(f"/api/contracts/{row['id']}", json=detail, headers=zhang)
    assert denied.status_code == 403
    assert denied.json()["message"] == "无权修改，仅创建者可修改"
    detail["version"] = 999
    conflict = client.put(f"/api/contracts/{row['id']}", json=detail, headers=admin)
    assert conflict.status_code == 409
    assert conflict.json()["message"] == "数据已被他人修改，请刷新"


def test_plan_and_order_and_stock(client):
    admin, _ = login(client)
    party = client.get("/api/parties", params={"type": "supplier"}, headers=admin).json()["data"]["items"][0]
    users = client.get("/api/users", headers=admin).json()["data"]["items"]
    owner = next(u["employee_id"] for u in users if u["employee_id"] == "E001")
    missing = client.post("/api/contracts", json={
        "direction": "purchase", "party_id": party["id"], "owner_id": owner, "name": "无计划",
        "status": "已签订", "pricing": "fixed", "tax_rate": "13",
        "items": [{"material_code": "NO-PLAN", "name_cn": "无计划物料", "unit": "个", "qty": "1", "price": "1"}],
    }, headers=admin)
    assert missing.status_code == 400
    assert missing.json()["message"] == "物料 NO-PLAN 没有采购计划"
    over = client.post("/api/contracts", json={
        "direction": "purchase", "party_id": party["id"], "owner_id": owner, "name": "超计划",
        "status": "已签订", "pricing": "fixed", "tax_rate": "13",
        "items": [{"material_code": "M001", "name_cn": "热轧钢板", "unit": "吨", "qty": "95", "price": "1"}],
    }, headers=admin)
    assert over.status_code == 400
    assert "超出" in over.json()["message"]
    orders = client.get("/api/orders", params={"direction": "purchase"}, headers=admin).json()["data"]["items"]
    assert orders[0]["order_code"].startswith("orderCG-2026-001")
    order_id = orders[0]["order_id"]
    version = orders[0]["version"]
    deleted = client.delete(f"/api/orders/{order_id}", params={"version": version}, headers=admin)
    assert deleted.status_code == 200, deleted.text
    contracts = client.get("/api/contracts", params={"direction": "purchase", "orderable": "1"}, headers=admin).json()["data"]["items"]
    target = next(c for c in contracts if c["contract_code"] == "CG-2026-001")
    detail = client.get(f"/api/contracts/{target['id']}", headers=admin).json()["data"]
    again = client.post("/api/orders", json={
        "contract_id": target["id"], "direction": "purchase",
        "items": [{"material_code": "M001", "name_cn": "热轧钢板", "unit": "吨", "qty": "10", "price": "5000"}],
    }, headers=admin)
    assert again.status_code == 200, again.text
    assert again.json()["data"]["order_code"] == "orderCG-2026-001"
    wh = client.get("/api/warehouses", headers=admin).json()["data"]["items"][0]
    outbound = client.post("/api/wms-bills", json={
        "type": "outbound", "contract_id": detail["id"], "warehouse_id": wh["id"], "bill_date": "2026-03-20",
        "items": [{"material_code": "M001", "qty": "8", "price": "5000"}],
    }, headers=admin)
    assert outbound.status_code == 400
    sale = client.get("/api/contracts", params={"direction": "sale"}, headers=admin).json()["data"]["items"][0]
    ship = client.post("/api/wms-bills", json={
        "type": "outbound", "contract_id": sale["id"], "warehouse_id": wh["id"], "bill_date": "2026-03-21",
        "items": [{"material_code": "M001", "name_cn": "热轧钢板", "unit": "吨", "qty": "3", "price": "8000"}],
    }, headers=admin)
    assert ship.status_code == 200, ship.text
    inbound = client.get("/api/wms-bills", params={"type": "inbound"}, headers=admin).json()["data"]["items"][0]
    blocked = client.delete(f"/api/wms-bills/{inbound['id']}", params={"version": inbound["version"]}, headers=admin)
    assert blocked.status_code == 400
    assert "库存" in blocked.json()["message"]


def test_contract_move_state(client):
    admin, _ = login(client)
    rows = client.get("/api/contracts", params={"direction": "purchase"}, headers=admin).json()["data"]["items"]
    main = next(item for item in rows if item["contract_code"] == "CG-2026-001")
    assert main["move_state"] == "部分到货"
    full = next(item for item in rows if item["contract_code"] == "HT-BG-2026-001")
    assert full["move_state"] == "已到货"
    partial = client.get("/api/contracts", params={"direction": "purchase", "move_state": "部分到货"}, headers=admin).json()["data"]["items"]
    assert {item["contract_code"] for item in partial} >= {"CG-2026-001"}
    assert all(item["move_state"] == "部分到货" for item in partial)
    missed = client.get("/api/contracts", params={"direction": "purchase", "move_state": "未到货"}, headers=admin).json()["data"]["items"]
    assert all(item["contract_code"] != "CG-2026-001" for item in missed)
    sales = client.get("/api/contracts", params={"direction": "sale"}, headers=admin).json()["data"]["items"]
    sale = next(item for item in sales if item["contract_code"] == "XS-2026-001")
    assert sale["move_state"] in {"未发货", "部分发货", "已发货"}


def test_stage_delete_and_unsigned(client):
    admin, _ = login(client)
    row = client.get("/api/contracts", params={"direction": "purchase"}, headers=admin).json()["data"]["items"]
    main = next(item for item in row if item["contract_code"] == "CG-2026-001")
    detail = client.get(f"/api/contracts/{main['id']}", headers=admin).json()["data"]
    detail["stages"] = [stage for stage in detail["stages"] if stage["seq"] != 1]
    removed = client.put(f"/api/contracts/{main['id']}", json=detail, headers=admin)
    assert removed.status_code == 400
    assert removed.json()["message"] == "后续已生成单据，无法删除。"
    party = client.get("/api/parties", params={"type": "customer"}, headers=admin).json()["data"]["items"][0]
    draft = client.post("/api/contracts", json={
        "direction": "sale", "party_id": party["id"], "owner_id": "E001", "name": "未签订合同",
        "status": "未签订", "pricing": "fixed", "tax_rate": "13", "amount_manual": "100",
        "items": [{"material_code": "M002", "name_cn": "高强螺栓", "unit": "件", "qty": "1", "price": "2"}],
    }, headers=admin)
    assert draft.status_code == 200, draft.text
    cid = draft.json()["data"]["id"]
    order = client.post("/api/orders", json={
        "contract_id": cid, "direction": "sale",
        "items": [{"material_code": "M002", "name_cn": "高强螺栓", "unit": "件", "qty": "1", "price": "2"}],
    }, headers=admin)
    assert order.status_code == 400
    detail = client.get(f"/api/contracts/{cid}", headers=admin).json()["data"]
    code = detail["contract_code"]
    removed = client.delete(f"/api/contracts/{cid}", params={"version": detail["version"]}, headers=admin)
    assert removed.status_code == 200, removed.text
    reused = client.post("/api/contracts", json={
        "contract_code": code, "direction": "sale", "party_id": party["id"], "owner_id": "E001", "name": "未签订合同",
        "status": "未签订", "pricing": "fixed", "tax_rate": "13", "amount_manual": "100",
        "items": [{"material_code": "M002", "name_cn": "高强螺栓", "unit": "件", "qty": "1", "price": "2"}],
    }, headers=admin)
    assert reused.status_code == 200, reused.text
    assert reused.json()["data"]["contract_code"] == code


def test_password_invalidates_token(client):
    headers, data = login(client, "zhang", "123456")
    changed = client.post("/api/auth/change-password", json={"old_password": "123456", "new_password": "123456"}, headers=headers)
    assert changed.status_code == 200, changed.text
    new_token = changed.json()["data"]["token"]
    old = client.get("/api/auth/me", headers=headers)
    assert old.status_code == 401
    me = client.get("/api/auth/me", headers={"Authorization": f"Bearer {new_token}"})
    assert me.status_code == 200


def _book(rows):
    book = openpyxl.Workbook()
    sheet = book.active
    for row in rows:
        sheet.append(row)
    buf = BytesIO()
    book.save(buf)
    return buf.getvalue()


def _import(client, headers, module, rows, filename, **fields):
    content = rows if isinstance(rows, bytes) else _book(rows)
    parsed = client.post(
        f"/api/import/{module}/parse",
        headers=headers,
        data=fields,
        files={"file": (filename, content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert parsed.status_code == 200, parsed.text
    body = parsed.json()["data"]
    executed = client.post(
        f"/api/import/{module}/execute",
        headers=headers,
        json={
            "file_hash": body["file_hash"],
            "mode": fields.get("mode", ""),
            "direction": fields.get("direction", ""),
            "party_type": fields.get("party_type", ""),
            "rows": body["rows"],
            "locks": body.get("locks") or [],
        },
    )
    assert executed.status_code == 200, executed.text
    return body, executed.json()["data"]


def test_batch_import(client):
    admin, _ = login(client)
    bad = client.post(
        "/api/import/party/parse",
        headers=admin,
        data={"mode": "update", "party_type": "supplier"},
        files={"file": ("a.txt", b"hello", "text/plain")},
    )
    assert bad.status_code == 400
    assert bad.json()["message"] == "文件格式不正确"
    bad_xlsx = client.post(
        "/api/import/party/parse",
        headers=admin,
        data={"mode": "update", "party_type": "supplier"},
        files={"file": ("a.xlsx", b"not-a-zip", "application/octet-stream")},
    )
    assert bad_xlsx.status_code == 400
    assert bad_xlsx.json()["message"] == "文件格式不正确"
    template = client.get("/api/import/contract/template", params={"direction": "purchase"}, headers=admin)
    assert template.status_code == 200
    assert template.content.startswith(b"PK")

    first = [
        ["编码", "名称", "类型", "税号", "银行账号", "默认税率", "开票名称", "供应商类别", "业务员"],
        ["S901", "=南方导入", "供应商", "", "", 13, "南方导入", "钢材", ""],
    ]
    _, created = _import(client, admin, "party", first, "party.xlsx", mode="update", party_type="supplier")
    assert created["success"] == 1 and created["fail"] == 0
    again = client.post(
        "/api/import/party/execute",
        headers=admin,
        json={"file_hash": _import_hash(client, admin, first), "mode": "update", "party_type": "supplier", "direction": "", "rows": [{"ok": True, "row_no": 2, "data": {"code": "S901", "name": "不该写入"}}]},
    )
    assert again.status_code == 200
    party = client.get("/api/parties", params={"type": "supplier", "keyword": "S901"}, headers=admin).json()["data"]["items"][0]
    assert party["name"] == "南方导入"
    assert Decimal(party["default_tax_rate"]) == Decimal("13")

    second = [
        ["编码", "名称", "类型", "税号", "银行账号", "默认税率", "开票名称", "供应商类别", "业务员"],
        ["S901", "南方导入改名", "供应商", "", "", 6, "南方导入改名", "钢材", ""],
    ]
    _, renamed = _import(client, admin, "party", second, "party2.xlsx", mode="update", party_type="supplier")
    assert renamed["need_name_sync"] is True
    party = client.get("/api/parties", params={"type": "supplier", "keyword": "S901"}, headers=admin).json()["data"]["items"][0]
    assert party["name"] == "南方导入改名"
    assert Decimal(party["default_tax_rate"]) == Decimal("13")

    zhang, _ = login(client, "zhang", "123456")
    third = [
        ["编码", "名称", "类型", "税号", "银行账号", "默认税率", "开票名称", "供应商类别", "业务员"],
        ["S901", "张三越权", "供应商", "", "", 6, "张三越权", "钢材", ""],
    ]
    preview, denied = _import(client, zhang, "party", third, "party3.xlsx", mode="update", party_type="supplier")
    assert preview["rows"][0]["ok"] is False
    assert denied["success"] == 0
    assert denied["errors"][0]["message"] == "无权修改，仅创建者可修改"
    party = client.get("/api/parties", params={"type": "supplier", "keyword": "S901"}, headers=admin).json()["data"]["items"][0]
    assert party["name"] == "南方导入改名"

    contract_rows = [
        ["合同编号", "合同名称", "合同类型", "对方编码", "业务员员工编号", "合同状态", "计价方式", "税率", "签订日期", "开始日期", "结束日期", "内容概述", "备注"],
        ["CG-IMP-900", "导入采购合同", "", "S001", "E001", "已签订", "固定总价", 13, "2026-03-01", "", "", "导入", ""],
        [],
        ["对应合同编号", "物料编号", "中文名称", "规格", "单位", "数量", "含税单价", "备注", "计划单号", "使用设备"],
        ["CG-IMP-900", "M001", "热轧钢板", "10mm", "吨", 1, 100, "", ""],
    ]
    _, contract = _import(client, admin, "contract", contract_rows, "contract.xlsx", mode="overwrite", direction="purchase")
    assert contract["success"] == 1, contract
    saved = client.get("/api/contracts", params={"direction": "purchase", "keyword": "CG-IMP-900"}, headers=admin).json()["data"]["items"]
    assert saved and Decimal(saved[0]["amount"]) == Decimal("100")

    order_rows = [
        ["订单编号", "源合同编号", "单据日期", "备注"],
        ["", "CG-IMP-900", "2026-03-02", ""],
        [],
        ["对应订单编号", "物料编号", "数量", "含税单价", "备注"],
        ["", "M001", 1, 100, ""],
    ]
    _, order = _import(client, admin, "order", order_rows, "order.xlsx", mode="create", direction="purchase")
    assert order["success"] == 1, order
    _, again_order = _import(client, admin, "order", order_rows, "order.xlsx", mode="create", direction="purchase")
    assert again_order["success"] == 1 and again_order["fail"] == 0

    fresh_order = [
        ["订单编号", "源合同编号", "单据日期", "备注"],
        ["PO-IMP-901", "CG-IMP-900", "2026-03-03", ""],
        [],
        ["对应订单编号", "物料编号", "数量", "含税单价", "备注"],
        ["PO-IMP-901", "M001", 1, 100, ""],
    ]
    parsed_order, blocked = _import(client, admin, "order", fresh_order, "order2.xlsx", mode="create", direction="purchase")
    assert parsed_order["rows"][0]["ok"] is False
    assert blocked["success"] == 0
    assert "该合同已有订单" in blocked["errors"][0]["message"]

    invoice_rows = [
        ["单据编号", "发票号码", "发票日期", "发票类型", "客商编码", "合同编号", "不含税金额", "税额", "价税合计", "备注"],
        ["", "INV-IMP-900", "2026-03-04", "增值税专用发票", "S001", "CG-IMP-900", 1, 1, 1, ""],
        [],
        ["对应发票号码", "物料编号", "数量", "含税单价", "税率", "备注"],
        ["INV-IMP-900", "M001", 1, 100, 13, ""],
    ]
    preview_inv, invoice = _import(client, admin, "invoice", invoice_rows, "invoice.xlsx", mode="create", direction="input")
    assert "将以明细为准" in preview_inv["rows"][0]["message"]
    assert invoice["success"] == 1, invoice

    settle_rows = [
        ["单据编号", "单据日期", "客商编码", "合同编号", "阶段序号", "阶段名称", "金额", "账户", "备注"],
        ["FK-IMP-900", "2026-03-05", "S001", "CG-IMP-900", 1, "", 10, "基本户", ""],
    ]
    _, settle = _import(client, admin, "settle", settle_rows, "settle.xlsx", mode="create", direction="payment")
    assert settle["success"] == 1, settle

    goods_rows = [
        ["物料编号", "名称", "规格", "单位", "分类编码", "含税单价"],
        ["M901", "导入商品", "10mm", "吨", "", 88],
    ]
    _, goods = _import(client, admin, "goods", goods_rows, "goods.xlsx", mode="update")
    assert goods["success"] == 1, goods

    plan_rows = [
        ["计划月份", "采购序号", "物料编号", "采购计划单号", "使用设备", "中文名称", "英文名称", "规格型号", "单位", "采购数量", "用途", "使用部门", "备注"],
        ["2026-03", "9", "M901", "PL-901", "钢材", "导入商品", "", "10mm", "吨", 20, "生产", "采购部", ""],
        ["", "", "", "", "", "合计", "", "", "", 20, "", "", ""],
    ]
    _, plan = _import(client, admin, "purchase-plan", plan_rows, "plan.xlsx", mode="append")
    assert plan["success"] == 1, plan
    _, duplicated = _import(client, admin, "purchase-plan", plan_rows, "plan.xlsx", mode="append")
    assert duplicated["success"] == 1 and duplicated["fail"] == 0
    plan_again = [
        ["计划月份", "采购序号", "物料编号", "采购计划单号", "使用设备", "中文名称", "英文名称", "规格型号", "单位", "采购数量", "用途", "使用部门", "备注"],
        ["2026-03", "9", "M901", "PL-901", "钢材", "导入商品", "", "10mm", "吨", 20, "生产", "采购部", ""],
    ]
    parsed_plan, blocked_plan = _import(client, admin, "purchase-plan", plan_again, "plan2.xlsx", mode="append")
    assert parsed_plan["rows"][0]["ok"] is False
    assert blocked_plan["success"] == 0

    replace_rows = [
        ["计划月份", "采购序号", "物料编号", "采购计划单号", "使用设备", "中文名称", "英文名称", "规格型号", "单位", "采购数量", "用途", "使用部门", "备注"],
        ["2026-03", "1", "M901", "PL-901", "钢材", "导入商品", "", "10mm", "吨", 20, "生产", "采购部", ""],
    ]
    _, replaced = _import(client, admin, "purchase-plan", replace_rows, "plan3.xlsx", mode="replace")
    assert replaced["success"] == 0
    assert any("超出" in item["message"] for item in replaced["errors"])
    still = client.get("/api/purchase-plans", params={"keyword": "M001"}, headers=admin).json()["data"]["items"]
    assert any(item["material_code"] == "M001" for item in still)


def test_contract_sections_and_added_lines(client):
    admin, _ = login(client)
    party = [
        ["供应商编码", "供应商名称", "类型"],
        ["S902", "导入供应商", "供应商"],
    ]
    _, created = _import(client, admin, "party", party, "s902.xlsx", mode="update", party_type="supplier")
    assert created["success"] == 1
    blocked_rows = [
        ["合同编号", "合同名称", "对方编码", "业务员员工编号", "合同状态", "计价方式", "税率"],
        ["CG-NOPLAN", "无计划合同", "S902", "admin", "已签订", "固定总价", 13],
        ["对应合同编号", "物料编号", "中文名称", "规格", "单位", "数量", "含税单价", "备注"],
        ["CG-NOPLAN", "NO-PLAN-IMP", "无计划", "X", "个", 1, 10, ""],
    ]
    blocked_preview, blocked = _import(client, admin, "contract", blocked_rows, "noplan.xlsx", mode="append", direction="purchase")
    assert blocked_preview["rows"][0]["ok"] is False
    assert "没有采购计划" in blocked_preview["rows"][0]["message"]
    assert blocked["success"] == 0
    plans = [
        ["计划月份", "采购序号", "物料编号", "采购计划单号", "使用设备", "中文名称", "规格型号", "单位", "采购数量"],
        ["2026-09", "1", "SEC-M901", "PL-901", "钢材", "钢板", "SEC-SPEC-1", "吨", 10],
        ["2026-09", "2", "SEC-M902", "PL-902", "五金", "螺栓", "SEC-SPEC-2", "套", 10],
        ["2026-09", "3", "SEC-M903", "PL-903", "五金", "垫片", "SEC-SPEC-3", "个", 10],
    ]
    _, planned = _import(client, admin, "purchase-plan", plans, "sec-plan.xlsx", mode="append")
    assert planned["success"] == 3, planned
    rows = [
        ["合同编号", "合同名称", "对方编码", "业务员员工编号", "合同状态", "计价方式", "税率", "签订日期", "开始日期", "结束日期", "标的名称", "收货地址", "收货人", "运输方式", "交货期限", "运费承担", "付款条款", "验收标准", "附件说明"],
        ["CG-SEC-900", "分段导入合同", "S902", "admin", "已签订", "固定总价", 13, "2026-09-29", "2026-09-29", "2026-12-31", "分段标的", "示例地址", "张三", "汽运", "签订后30天", "供方", "按阶段付款", "按样品验收", "见附件"],
        ["对应合同编号", "物料编号", "中文名称", "规格", "单位", "数量", "含税单价", "备注"],
        ["CG-SEC-900", "SEC-M901", "钢板", "SEC-SPEC-1", "吨", 2, 100, ""],
        ["CG-SEC-900", "SEC-M902", "螺栓", "SEC-SPEC-2", "套", 4, 5, ""],
        ["阶段名称", "比例", "计划日期", "备注"],
        ["预付款", 30, "2026-09-29", ""],
        ["进度款", 40, "", ""],
        ["验收款", 25, "", ""],
        ["质保金", 5, "2026-12-31", ""],
    ]
    preview, result = _import(client, admin, "contract", rows, "sec.xlsx", mode="append", direction="purchase")
    assert preview["rows"][0]["ok"] is True, preview["rows"][0]
    assert result["success"] == 1, result
    listed = client.get("/api/contracts", params={"direction": "purchase", "keyword": "CG-SEC-900"}, headers=admin).json()["data"]["items"]
    detail = client.get(f"/api/contracts/{listed[0]['id']}", headers=admin).json()["data"]
    assert Decimal(detail["amount"]) == Decimal("220")
    assert detail["subject_name"] == "分段标的"
    assert detail["ship_address"] == "示例地址"
    assert detail["receiver_name"] == "张三"
    assert detail["transport_mode"] == "汽运"
    assert detail["delivery_term"] == "签订后30天"
    assert detail["freight_bearer"] == "供方"
    assert detail["terms"] == "按阶段付款"
    assert detail["acceptance_text"] == "按样品验收"
    assert detail["annex_note"] == "见附件"
    assert [stage["name"] for stage in detail["stages"]] == ["预付款", "进度款", "验收款", "质保金"]
    assert sum(Decimal(stage["plan_amount"]) for stage in detail["stages"]) == Decimal("220")
    extra = [
        ["合同编号", "合同名称", "对方编码", "业务员员工编号", "合同状态", "计价方式", "税率"],
        ["CG-SEC-900", "分段导入合同", "S902", "admin", "已签订", "固定总价", 13],
        ["对应合同编号", "物料编号", "中文名称", "规格", "单位", "数量", "含税单价", "备注"],
        ["CG-SEC-900", "SEC-M903", "垫片", "SEC-SPEC-3", "个", 1, 10, ""],
    ]
    _, added = _import(client, admin, "contract", extra, "sec2.xlsx", mode="append", direction="purchase")
    assert added["success"] == 1, added
    detail = client.get(f"/api/contracts/{listed[0]['id']}", headers=admin).json()["data"]
    assert {item["material_code"] for item in detail["items"]} == {"SEC-M901", "SEC-M902", "SEC-M903"}
    assert [stage["name"] for stage in detail["stages"]] == ["预付款", "进度款", "验收款", "质保金"]
    assert Decimal(detail["amount"]) == Decimal("230")
    assert sum(Decimal(stage["plan_amount"]) for stage in detail["stages"]) == Decimal("230")
    assert detail["subject_name"] == "分段标的"
    assert detail["terms"] == "按阶段付款"


def test_contract_import_by_party_name(client):
    admin, _ = login(client)
    party = [
        ["供应商编码", "供应商名称", "类型"],
        ["S903", "名称供应商", "供应商"],
    ]
    _, created = _import(client, admin, "party", party, "s903.xlsx", mode="update", party_type="supplier")
    assert created["success"] == 1
    plans = [
        ["计划月份", "采购序号", "物料编号", "采购计划单号", "使用设备", "中文名称", "规格型号", "单位", "采购数量"],
        ["2026-09", "1", "NAME-M1", "PL-NAME", "钢材", "钢板", "NAME-SPEC", "吨", 10],
    ]
    _, planned = _import(client, admin, "purchase-plan", plans, "name-plan.xlsx", mode="append")
    assert planned["success"] == 1, planned
    rows = [
        ["合同编号", "合同名称", "对方名称", "业务员员工编号", "税率", "签订日期", "标的名称", "付款条款"],
        ["CG-NAME-1", "按名称导入", "名称供应商", "admin", 13, "2026-09-29", "名称标的", "按阶段付款"],
        ["对应合同编号", "物料编号", "中文名称", "规格", "单位", "数量", "含税单价", "备注"],
        ["CG-NAME-1", "NAME-M1", "钢板", "NAME-SPEC", "吨", 1, 10, ""],
    ]
    preview, result = _import(client, admin, "contract", rows, "by-name.xlsx", mode="append", direction="purchase")
    assert preview["rows"][0]["ok"] is True, preview["rows"][0]
    assert result["success"] == 1, result
    listed = client.get("/api/contracts", params={"direction": "purchase", "keyword": "CG-NAME-1"}, headers=admin).json()["data"]["items"]
    detail = client.get(f"/api/contracts/{listed[0]['id']}", headers=admin).json()["data"]
    assert detail["party_name"] == "名称供应商"
    assert detail["status"] == "已签订"
    assert detail["pricing"] == "fixed"
    assert detail["subject_name"] == "名称标的"
    assert detail["terms"] == "按阶段付款"
    assert detail["content"] == ""
    assert detail["remark"] == ""
    again = [
        ["供应商编码", "供应商名称", "类型"],
        ["S904", "名称供应商", "供应商"],
    ]
    _, duplicated = _import(client, admin, "party", again, "s904.xlsx", mode="update", party_type="supplier")
    assert duplicated["success"] == 1
    clash = [
        ["合同编号", "合同名称", "对方名称", "业务员员工编号", "税率"],
        ["CG-NAME-2", "重名供应商", "名称供应商", "admin", 13],
        ["对应合同编号", "物料编号", "中文名称", "规格", "单位", "数量", "含税单价", "备注"],
        ["CG-NAME-2", "NAME-M1", "钢板", "NAME-SPEC", "吨", 1, 10, ""],
    ]
    blocked, failed = _import(client, admin, "contract", clash, "clash.xlsx", mode="append", direction="purchase")
    assert blocked["rows"][0]["ok"] is False
    assert "对方名称不唯一" in blocked["rows"][0]["message"]
    assert failed["success"] == 0


def test_stock_delete_only_when_zero(client):
    admin, _ = login(client)
    created = client.post("/api/warehouses", headers=admin, json={"code": "WH-DEL", "name": "删除仓"}).json()
    assert created["code"] == 0, created
    from backend.app.database import SessionLocal
    from backend.app.models import Stock, Warehouse
    from backend.app.services.util import now

    db = SessionLocal()
    try:
        warehouse = db.get(Warehouse, created["data"]["id"])
        moment = now()
        zero = Stock(company_id=warehouse.company_id, warehouse_id=warehouse.id, material_code="ZERO-1", qty=Decimal("0"), version=1, created_by="admin", updated_by="admin", created_at=moment, updated_at=moment)
        busy = Stock(company_id=warehouse.company_id, warehouse_id=warehouse.id, material_code="BUSY-1", qty=Decimal("3"), version=1, created_by="admin", updated_by="admin", created_at=moment, updated_at=moment)
        db.add_all([zero, busy])
        db.commit()
        zero_id, zero_version = zero.id, zero.version
        busy_id, busy_version = busy.id, busy.version
    finally:
        db.close()
    denied = client.delete(f"/api/stocks/{busy_id}", headers=admin, params={"version": busy_version})
    assert denied.status_code == 400
    assert "不为零" in denied.json()["message"]
    removed = client.delete(f"/api/stocks/{zero_id}", headers=admin, params={"version": zero_version})
    assert removed.status_code == 200, removed.text
    listed = client.get("/api/stocks", headers=admin, params={"material_code": "ZERO-1"}).json()["data"]["items"]
    assert all(item["material_code"] != "ZERO-1" for item in listed)
    still = client.get("/api/stocks", headers=admin, params={"material_code": "BUSY-1"}).json()["data"]["items"]
    assert any(item["material_code"] == "BUSY-1" for item in still)


def test_export_contract_excel(client):
    from backend.app.services.list_export import _safe

    assert _safe("=1+1") == "'=1+1"
    headers, _ = login(client)
    columns = [{"key": "contract_code", "label": "合同编号"}, {"key": "name", "label": "合同名称"}, {"key": "owner_name", "label": "业务员"}]
    got = client.get("/api/export/contract", params={"direction": "purchase"}, headers=headers)
    assert got.status_code == 200, got.text
    assert "spreadsheetml" in got.headers["content-type"]
    listed = client.get("/api/contracts", params={"direction": "purchase", "pageSize": 5}, headers=headers).json()["data"]["items"]
    assert listed
    one = listed[0]
    posted = client.post(
        "/api/export/contract",
        headers=headers,
        json={"direction": "purchase", "ids": [str(one["id"])], "title": "采购合同", "columns": columns},
    )
    assert posted.status_code == 200, posted.text
    book = openpyxl.load_workbook(BytesIO(posted.content))
    sheet = book.active
    assert [cell.value for cell in sheet[1]] == ["合同编号", "合同名称", "业务员"]
    assert sheet.max_row == 2
    assert sheet["A2"].value == one["contract_code"]


def test_export_pending_and_csv(client):
    headers, _ = login(client)
    pending = client.post(
        "/api/export/pending",
        headers=headers,
        json={"direction": "purchase", "title": "待付款", "columns": [{"key": "contract_code", "label": "合同编号"}, {"key": "detail", "label": "付款详情"}, {"key": "remark", "label": "备注"}]},
    )
    assert pending.status_code == 200, pending.text
    book = openpyxl.load_workbook(BytesIO(pending.content))
    assert "付款详情" in [cell.value for cell in book.active[1]]
    csv_res = client.post(
        "/api/export/supplier-detail",
        headers=headers,
        json={"format": "csv", "title": "供应商往来统计", "columns": [{"key": "name", "label": "供应商名称"}, {"key": "unpaid", "label": "未付款"}]},
    )
    assert csv_res.status_code == 200, csv_res.text
    assert "text/csv" in csv_res.headers["content-type"]
    assert "供应商名称" in csv_res.content.decode("utf-8-sig")


def _import_hash(client, headers, rows):
    content = _book(rows)
    parsed = client.post(
        "/api/import/party/parse",
        headers=headers,
        data={"mode": "update", "party_type": "supplier"},
        files={"file": ("party.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert parsed.status_code == 200, parsed.text
    return parsed.json()["data"]["file_hash"]


def test_generate_input_invoice_prefill(client):
    admin, _ = login(client)
    purchase = client.get("/api/contracts", params={"direction": "purchase", "status": "已签订"}, headers=admin).json()["data"]["items"][0]
    res = client.post(f"/api/contracts/{purchase['id']}/generate-input-invoice", headers=admin)
    assert res.status_code == 200, res.text
    data = res.json()["data"]
    assert data["saved"] is False
    assert data["direction"] == "input"
    assert data["contract_id"] == purchase["id"]
    assert data["party_id"] == purchase["party_id"]
    assert data["items"]
    assert data["items"][0]["material_code"]
    assert float(data["items"][0]["qty"]) > 0
    sale = client.get("/api/contracts", params={"direction": "sale", "status": "已签订"}, headers=admin).json()["data"]["items"][0]
    refused = client.post(f"/api/contracts/{sale['id']}/generate-input-invoice", headers=admin)
    assert refused.status_code == 400


def test_plan_short_field_names(client):
    admin, _ = login(client)
    rows = [
        ["序号", "月份", "计划单号", "物料编号", "中文名称", "采购数量"],
        ["3", "202609", "MCS20260901-1", "SHORT-M1", "短字段物料", 6],
    ]
    _, plan = _import(client, admin, "purchase-plan", rows, "short-plan.xlsx", mode="append")
    assert plan["success"] == 1, plan
    items = client.get("/api/purchase-plans", params={"keyword": "SHORT-M1"}, headers=admin).json()["data"]["items"]
    hit = next(item for item in items if item["material_code"] == "SHORT-M1")
    assert hit["purchase_seq"] == "3"
    assert hit["plan_month"] == "202609"
    assert hit["plan_no"] == "MCS20260901-1"
