"""报关导入模板与一键导出压缩包。正式报关/清关文件改由基础资料导出模板生成。"""
from copy import copy
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..models import BizFile, Company
from .util import q

THIN = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)
HEAD = Font(name="Microsoft YaHei", bold=True, color="FFFFFF")
BODY = Font(name="Microsoft YaHei", size=10)
FILL = PatternFill("solid", fgColor="1677FF")


def _cell(sheet, row, col, value, header=False):
    cell = sheet.cell(row, col, value if value not in (None, "") else "")
    cell.font = HEAD if header else BODY
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    cell.border = THIN
    if header:
        cell.fill = FILL
    return cell


def _width(sheet, widths):
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def zip_files(entries: list[tuple[str, bytes]]) -> bytes:
    buf = BytesIO()
    with ZipFile(buf, "w", ZIP_DEFLATED) as zf:
        for name, content in entries:
            zf.writestr(name, content)
    return buf.getvalue()


def _attach(db, company_id, business_type, business_id):
    rows = q(db, BizFile, company_id).filter(
        BizFile.business_type == business_type,
        BizFile.business_id == int(business_id),
        BizFile.status == "启用",
    ).all()
    entries = []
    used = set()
    for row in rows:
        path = Path(row.file_path)
        if not path.is_file():
            continue
        name = row.file_name
        if name in used:
            name = f"{row.id}_{name}"
        used.add(name)
        entries.append((f"附件/{name}", path.read_bytes()))
    return entries


def _form_zip(db, company_id, doc_code, record_id, template_id, extra_attach):
    from .form_docx import build_form

    company = db.get(Company, company_id)
    content, filename = build_form(db, company, doc_code, record_id, template_id)
    entries = [(filename, content)]
    for business_type, business_id in extra_attach:
        if business_id:
            entries += _attach(db, company_id, business_type, business_id)
    return zip_files(entries)


def invoice_zip(db, invoice, template_id=None) -> bytes:
    extra = [("customs_invoice", invoice.id)]
    if invoice.customs_element_id:
        extra.append(("customs_element", invoice.customs_element_id))
    return _form_zip(db, invoice.company_id, "customs_invoice", invoice.id, template_id, extra)


def clearance_zip(db, row, template_id=None) -> bytes:
    return _form_zip(db, row.company_id, "customs_clearance", row.id, template_id, [("customs_clearance", row.id)])


def receipt_template() -> tuple[bytes, str]:
    book = Workbook()
    sheet = book.active
    sheet.title = "入仓明细"
    headers = ["入仓单号", "入仓日期", "合同号", "箱单号", "物料大类", "物料编号", "中文名称", "英文名称", "规格", "单位", "数量", "产地", "使用设备", "箱号", "集装箱号", "提单号", "船名/航号", "毛重", "净重", "长", "宽", "高", "体积", "PO单号", "含税单价", "特殊申报", "包件件数", "询价名称", "询价型号", "询价数量", "出发港", "目的港", "收货人", "付款条款", "HS Code", "申报要素", "备注"]
    for col, text in enumerate(headers, 1):
        _cell(sheet, 1, col, text, True)
    _width(sheet, [16] * len(headers))
    output = BytesIO()
    book.save(output)
    return output.getvalue(), "入仓单导入模板.xlsx"


def material_fill_template() -> tuple[bytes, str]:
    book = Workbook()
    sheet = book.active
    sheet.title = "报关物料补录"
    headers = [
        "物料编号", "入仓单号", "产地", "箱号", "尺寸", "净重", "毛重", "体积",
        "出发港", "目的港", "集装箱号", "提单号", "船名/航号", "箱单号",
        "物料大类", "使用设备", "HS Code", "收货人", "包件件数", "备注",
    ]
    for col, text in enumerate(headers, 1):
        _cell(sheet, 1, col, text, True)
    _width(sheet, [18] * len(headers))
    output = BytesIO()
    book.save(output)
    return output.getvalue(), "报关物料批量导入模板.xlsx"


def _num(value):
    if value in (None, ""):
        return None
    return float(str(value).replace(",", ""))


def _find_row(sheet, token: str):
    needle = "{{" + token + "}}"
    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and needle in cell.value:
                return cell.row
    return None


def _token_value(raw, mapping: dict):
    if not isinstance(raw, str) or "{{" not in raw:
        return raw
    text = raw
    numeric = {
        "大类数量", "数量", "大类毛重", "大类净重", "大类重量", "大类包件件数",
        "人民币单价", "人民币总价", "装箱数", "换算美元金额", "毛重", "净重", "包件数", "包装件数", "汇率",
    }
    while "{{" in text:
        start = text.find("{{")
        end = text.find("}}", start)
        if end < 0:
            break
        key = text[start + 2:end].strip()
        if key not in mapping:
            break
        value = mapping.get(key)
        if text.strip() == "{{" + key + "}}" and key in numeric:
            return _num(value)
        replacement = "" if value is None else str(value)
        text = text[:start] + replacement + text[end + 2:]
    return text


def _detach_merges_below(sheet, row: int):
    from openpyxl.worksheet.cell_range import CellRange

    moved = []
    for merged in list(sheet.merged_cells.ranges):
        if merged.min_row >= row:
            moved.append(CellRange(str(merged)))
            sheet.unmerge_cells(str(merged))
    return moved


def _repeat_placeholder_row(sheet, token: str, records: list[dict]):
    src = _find_row(sheet, token)
    if src is None:
        return None
    max_col = max(sheet.max_column or 1, 1)
    original = [sheet.cell(src, col).value for col in range(1, max_col + 1)]
    count = len(records)
    detached = []
    if count > 1:
        detached = _detach_merges_below(sheet, src + 1)
        sheet.insert_rows(src + 1, count - 1)
        for offset in range(1, count):
            _clone_row(sheet, src, src + offset, max_col)
    usable = records or [{}]
    for index, record in enumerate(usable):
        line = src + index
        for col, raw in enumerate(original, start=1):
            if col == 1 and raw in (1, "1"):
                value = index + 1
            else:
                value = _token_value(raw, record)
            sheet.cell(line, col, value if value is not None else "")
    for offset in range(max(count, 1)):
        sheet.row_dimensions[src + offset].hidden = False
    if count:
        total_row = src + count
        first, last = src, src + count - 1
        for cell in sheet[total_row]:
            formula = cell.value
            if isinstance(formula, str) and "SUM(" in formula.upper():
                letter = get_column_letter(cell.column)
                cell.value = f"=SUM({letter}{first}:{letter}{last})"
    for rng in detached:
        rng.shift(0, count - 1)
        sheet.merge_cells(str(rng))
    return src


def _fill_remaining_tokens(sheet, mapping: dict):
    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and "{{" in cell.value:
                cell.value = _token_value(cell.value, mapping)


def _invoice_template() -> Path:
    from ..config import ROOT
    from ..errors import BizError

    name = "报关发票.xlsx"
    for folder in (ROOT.parent / "需求的模板", ROOT / "需求的模板"):
        path = folder / name
        if path.is_file():
            return path
    raise BizError("找不到报关发票.xlsx，请放在「需求的模板」目录")


def _clone_row(sheet, src, dst, max_col):
    for col in range(1, max_col + 1):
        source = sheet.cell(src, col)
        target = sheet.cell(dst, col)
        target.value = None
        if source.has_style:
            target.font = copy(source.font)
            target.border = copy(source.border)
            target.fill = copy(source.fill)
            target.number_format = source.number_format
            target.alignment = copy(source.alignment)
    height = sheet.row_dimensions[src].height
    if height:
        sheet.row_dimensions[dst].height = height


def _expand_rows(sheet, template_row, count, max_col):
    extra = max(count - 1, 0)
    if extra:
        sheet.insert_rows(template_row + 1, extra)
        for offset in range(1, count):
            _clone_row(sheet, template_row, template_row + offset, max_col)
    return template_row + max(count, 1)


def _invoice_header(detail) -> dict:
    return {
        "箱单号": detail.get("packing_list_no") or "",
        "发票日期2": detail.get("invoice_date_2") or "",
        "发票编号": detail.get("invoice_no") or "",
        "发票日期": detail.get("invoice_date_1") or detail.get("invoice_date") or "",
        "船名/航号": detail.get("vessel_voyage") or "",
        "集装箱号": detail.get("container_no") or "",
        "发票号": detail.get("invoice_no") or "",
    }


def _pl_record(row) -> dict:
    return {
        "物料大类": row.get("material_category") or "",
        "大类规格": row.get("specification") or "",
        "大类数量": row.get("quantity") or "",
        "大类单位": row.get("unit") or "",
        "大类包件件数": row.get("package_count") if row.get("package_count") not in (None, "") else 0,
        "大类毛重": row.get("gross_weight") or "",
        "大类净重": row.get("net_weight") or "",
        "长*高*宽": row.get("size_display") or "",
        "大类重量": row.get("weight") or "",
    }


def _invoice_record(detail, row) -> dict:
    return {
        "申报大类": row.get("material_category") or "",
        "英文名称": row.get("name_en") or "",
        "规格": row.get("specification") or "",
        "数量": row.get("quantity") or "",
        "单位": row.get("unit") or "",
        "人民币单价": row.get("price") or "",
        "人民币总价": row.get("amount") or "",
        "装箱数": row.get("package_count") or "",
        "发票号": detail.get("invoice_no") or "",
    }


def _write_pl(sheet, detail):
    categories = detail.get("categories") or []
    _repeat_placeholder_row(sheet, "物料大类", [_pl_record(row) for row in categories])
    _fill_remaining_tokens(sheet, _invoice_header(detail))


def _write_invoice_sheet(sheet, detail):
    categories = detail.get("categories") or []
    src = _repeat_placeholder_row(sheet, "申报大类", [_invoice_record(detail, row) for row in categories])
    _fill_remaining_tokens(sheet, _invoice_header(detail))
    rate = _num(detail.get("exchange_rate"))
    if rate is not None and sheet["L10"].value in (None, ""):
        sheet["L10"] = rate
    if not src:
        return
    for index, row in enumerate(categories):
        cell = sheet.cell(src + index, 9)
        if isinstance(cell.value, (int, float)):
            usd = _num(row.get("amount_usd"))
            cell.value = "" if usd is None else usd


def _write_details(sheet, detail, groups):
    title_style = sheet.cell(1, 1)
    header_style = [sheet.cell(2, col) for col in range(1, 13)]
    data_style = [sheet.cell(3, col) for col in range(1, 13)]
    total_style = sheet.cell(5, 2)
    styles = {
        "title": copy(title_style.font),
        "title_align": copy(title_style.alignment),
        "header": [(copy(cell.font), copy(cell.alignment), copy(cell.border), copy(cell.fill)) for cell in header_style],
        "data": [(copy(cell.font), copy(cell.alignment), copy(cell.border), copy(cell.fill), cell.number_format) for cell in data_style],
        "total_font": copy(total_style.font),
    }
    for merged in list(sheet.merged_cells.ranges):
        sheet.unmerge_cells(str(merged))
    for row in sheet.iter_rows(min_row=1, max_row=max(sheet.max_row, 1), max_col=15):
        for cell in row:
            cell.value = None
    rate = _num(detail.get("exchange_rate"))
    headers = ["编号", "中文名称", "英文名称", "型号", "单位", "数量", "FOB单价", "FOB总价", "Convert to USD", "询价名称/Inquiry Name", "询价型号\n/Inquiry Model", "询价数量\n/Inquiry Qty"]
    cursor = 1
    blocks = groups or []
    for index, group in enumerate(blocks):
        if index:
            cursor += 2
        title = sheet.cell(cursor, 1, group.get("label") or f"Specification # {index + 1}")
        title.font = styles["title"]
        title.alignment = styles["title_align"]
        sheet.merge_cells(start_row=cursor, start_column=1, end_row=cursor, end_column=12)
        cursor += 1
        for col, text in enumerate(headers, start=1):
            cell = sheet.cell(cursor, col, text)
            font, alignment, border, fill = styles["header"][col - 1]
            cell.font = font
            cell.alignment = alignment
            cell.border = border
            cell.fill = fill
        cursor += 1
        first = cursor
        for item in group.get("items") or []:
            values = [
                item.get("seq"),
                item.get("chinese_name") or "",
                item.get("english_name") or "",
                item.get("model") or "",
                item.get("unit") or "",
                _num(item.get("quantity")),
                _num(item.get("fob_unit_price")),
                None,
                None,
                item.get("inquiry_name") or "",
                item.get("inquiry_model") or "",
                _num(item.get("inquiry_qty")),
            ]
            for col, value in enumerate(values, start=1):
                cell = sheet.cell(cursor, col, value if value is not None else "")
                font, alignment, border, fill, number_format = styles["data"][col - 1]
                cell.font = font
                cell.alignment = alignment
                cell.border = border
                cell.fill = fill
                cell.number_format = number_format
            sheet.cell(cursor, 8, f"=F{cursor}*G{cursor}")
            if rate:
                sheet.cell(cursor, 9, f"=H{cursor}/{rate}")
            cursor += 1
        last = cursor - 1
        total = sheet.cell(cursor, 2, "TOTAL")
        total.font = styles["total_font"]
        if first <= last:
            sheet.cell(cursor, 6, f"=SUM(F{first}:F{last})")
            sheet.cell(cursor, 8, f"=SUM(H{first}:H{last})")
            sheet.cell(cursor, 9, f"=SUM(I{first}:I{last})")
        cursor += 1


def export_element_xlsx(detail: dict, template_path) -> bytes:
    """同一物料大类合并申报名称和申报要素，产品名称按物料明细展开；特殊申报一物料一行。"""
    book = load_workbook(template_path)
    sheet = book.active
    start = _find_row(sheet, "HS Code") or _find_row(sheet, "申报名称") or 2
    for merged in list(sheet.merged_cells.ranges):
        if merged.max_row >= start:
            sheet.unmerge_cells(str(merged))
    max_col = max(sheet.max_column or 1, 6)
    saved = []
    for col in range(1, max_col + 1):
        cell = sheet.cell(start, col)
        saved.append((copy(cell.font), copy(cell.border), copy(cell.fill), copy(cell.alignment), cell.number_format))
    if sheet.max_row >= start:
        sheet.delete_rows(start, sheet.max_row - start + 1)
    flat = []
    spans = []
    cursor = 0
    for item in detail.get("items") or []:
        lines = item.get("lines") or [{}]
        special = bool(item.get("special_declare"))
        block = lines or [{}]
        if not special and len(block) > 1:
            spans.append((cursor, len(block)))
        for line in block:
            flat.append([
                item.get("hs_code") or "",
                item.get("declaration_name") or "",
                line.get("product_name") or "",
                line.get("model") or item.get("specification") or "",
                item.get("declaration_elements") or "",
                line.get("origin_place") or "",
            ])
            cursor += 1
    wrap = Alignment(vertical="center", wrap_text=True)
    for index, values in enumerate(flat):
        row_no = start + index
        for col, value in enumerate(values, start=1):
            cell = sheet.cell(row_no, col, value)
            if col - 1 < len(saved):
                font, border, fill, alignment, number_format = saved[col - 1]
                cell.font = font
                cell.border = border
                cell.fill = fill
                cell.alignment = wrap if col == 5 else alignment
                cell.number_format = number_format
        sheet.row_dimensions[row_no].height = 22
    for begin, span in spans:
        first = start + begin
        last = first + span - 1
        for col in (1, 2, 5):
            sheet.merge_cells(start_row=first, start_column=col, end_row=last, end_column=col)
            sheet.cell(first, col).alignment = wrap
    output = BytesIO()
    book.save(output)
    return output.getvalue()


_ONES = (
    "ZERO", "ONE", "TWO", "THREE", "FOUR", "FIVE", "SIX", "SEVEN", "EIGHT", "NINE",
    "TEN", "ELEVEN", "TWELVE", "THIRTEEN", "FOURTEEN", "FIFTEEN", "SIXTEEN", "SEVENTEEN", "EIGHTEEN", "NINETEEN",
)
_TENS = ("", "", "TWENTY", "THIRTY", "FORTY", "FIFTY", "SIXTY", "SEVENTY", "EIGHTY", "NINETY")


def _english_under_thousand(number: int) -> str:
    parts = []
    if number >= 100:
        parts.append(f"{_ONES[number // 100]} HUNDRED")
        number %= 100
    if number >= 20:
        ten = _TENS[number // 10]
        rest = number % 10
        parts.append(f"{ten}-{_ONES[rest]}" if rest else ten)
    elif number:
        parts.append(_ONES[number])
    return " ".join(parts)


def _english_integer(number: int) -> str:
    if number <= 0:
        return "ZERO"
    parts = []
    for scale, name in ((1_000_000_000, "BILLION"), (1_000_000, "MILLION"), (1_000, "THOUSAND")):
        if number >= scale:
            parts.append(f"{_english_under_thousand(number // scale)} {name}")
            number %= scale
    if number:
        parts.append(_english_under_thousand(number))
    return " ".join(parts)


def usd_amount_words(amount) -> str:
    """美元金额写成发票上的英文大写，例如 USD ONE HUNDRED AND FIVE CENTS ONLY。"""
    number = _num(amount) or 0
    cents_total = int(round(number * 100))
    if cents_total < 0:
        cents_total = 0
    dollars, cents = divmod(cents_total, 100)
    return f"USD {_english_integer(dollars)} AND {_english_integer(cents)} CENTS ONLY"


def _materialize_sums(sheet, first: int, last: int, total_row: int):
    for cell in sheet[total_row]:
        formula = cell.value
        if not (isinstance(formula, str) and "SUM(" in formula.upper()):
            continue
        total = 0.0
        found = False
        for row in range(first, last + 1):
            number = _num(sheet.cell(row, cell.column).value)
            if number is None:
                continue
            total += number
            found = True
        if not found:
            continue
        total = round(total, 2)
        cell.value = int(total) if total == int(total) else total


def _write_amount_words(sheet, words: str):
    for row in sheet.iter_rows():
        for cell in row:
            if not isinstance(cell.value, str) or "TOTAL AMOUNT" not in cell.value.upper():
                continue
            for col in range(cell.column + 1, (sheet.max_column or cell.column) + 1):
                other = sheet.cell(cell.row, col)
                if isinstance(other.value, str) and "USD" in other.value.upper():
                    other.value = words
                    return
            sheet.cell(cell.row, 4, words)
            return


def _clearance_rate(detail) -> float | None:
    return _num(detail.get("exchange_rate"))


def _clearance_usd(item, rate):
    stored = _num(item.get("amount_usd"))
    if stored:
        return stored
    amount = _num(item.get("amount"))
    if amount is not None and rate:
        return round(amount / rate, 2)
    return 0


def export_clearance_xlsx(detail: dict, template_path) -> bytes:
    """清关发票按模板明细行展开：序号顺排，汇总写在明细下一行，美元合计写成英文大写。"""
    book = load_workbook(template_path)
    rate = _clearance_rate(detail)
    ci_records = []
    for item in detail.get("items") or []:
        ci_records.append({
            "中文名": item.get("name_cn") or "",
            "英文名": item.get("name_en") or "",
            "规格": item.get("specification") or "",
            "数量": item.get("quantity") or "",
            "单位": item.get("unit") or "",
            "人民币单价": item.get("price") or "",
            "人民币总价": item.get("amount") or "",
            "换算美元金额": _clearance_usd(item, rate),
            "PO单号": item.get("po_no") or "",
            "备注": item.get("remark") or "",
        })
    pl_records = []
    for item in detail.get("packing") or []:
        pl_records.append({
            "品名": item.get("category") or "",
            "中文名": item.get("category") or "",
            "规格": item.get("specification") or "",
            "数量": item.get("quantity") or "",
            "单位": item.get("unit") or "",
            "包件数": item.get("packages") if item.get("packages") not in (None, "") else 0,
            "包装件数": item.get("packages") if item.get("packages") not in (None, "") else 0,
            "毛重": item.get("gross_weight") or "",
            "净重": item.get("net_weight") or "",
            "尺寸": item.get("size_text") or "",
            "入仓单号": item.get("receipt_no") or "",
        })
    header = {
        "发票号": "" if detail.get("invoice_missing") else (detail.get("invoice_no") or ""),
        "制单日期": detail.get("invoice_date") or "",
        "船名/航号": detail.get("vessel_voyage") or "",
        "集装箱号": detail.get("container_no") or "",
        "提单号": detail.get("bill_no") or "",
        "出发港口": detail.get("port_of_shipment") or "",
        "出发港": detail.get("port_of_shipment") or "",
        "装箱号": detail.get("packing_list_no") or "",
        "汇率": detail.get("exchange_rate") or "",
    }
    ci = book["CI"] if "CI" in book.sheetnames else book.worksheets[0]
    pl = book["PL"] if "PL" in book.sheetnames else book.worksheets[-1]
    ci_src = _repeat_placeholder_row(ci, "中文名", ci_records)
    _fill_remaining_tokens(ci, header)
    if ci_src and ci_records:
        _materialize_sums(ci, ci_src, ci_src + len(ci_records) - 1, ci_src + len(ci_records))
        usd_total = 0.0
        for record in ci_records:
            usd_total += _num(record.get("换算美元金额")) or 0
        _write_amount_words(ci, usd_amount_words(round(usd_total, 2)))
    pl_src = _repeat_placeholder_row(pl, "品名", pl_records)
    if pl_src is None:
        pl_src = _repeat_placeholder_row(pl, "中文名", pl_records)
    _fill_remaining_tokens(pl, header)
    if pl_src and pl_records:
        _materialize_sums(pl, pl_src, pl_src + len(pl_records) - 1, pl_src + len(pl_records))
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def export_invoice_xlsx(detail: dict, specification_no=None, template_path=None) -> tuple[bytes, str]:
    groups = list(detail.get("detail_groups") or [])
    if specification_no not in (None, "", "all"):
        number = int(specification_no)
        groups = [group for group in groups if int(group.get("specification_no") or 0) == number]
    book = load_workbook(template_path or _invoice_template())
    _write_pl(book["PL"], detail)
    _write_invoice_sheet(book["Commercial Invoice"], detail)
    _write_details(book["明细表"], detail, groups)
    output = BytesIO()
    book.save(output)
    name = f"{detail.get('invoice_no') or '报关发票'}.xlsx"
    return output.getvalue(), name


def hs_template() -> tuple[bytes, str]:
    book = Workbook()
    sheet = book.active
    headers = ["HS Code", "申报名称", "产品名称", "申报规格", "报关要素", "描述", "状态"]
    for col, text in enumerate(headers, 1):
        _cell(sheet, 1, col, text, True)
    output = BytesIO()
    book.save(output)
    return output.getvalue(), "HS Code导入模板.xlsx"
