"""基础资料里的导出模板。一张单据可以有多份，每份记录模板绝对地址和导出名称。"""
import re
from pathlib import Path

from sqlalchemy import text

from ..config import ROOT
from ..database import ensure_columns, get_engine
from ..errors import BizError
from ..models import Company, ExportProfile, Party, now

TEMPLATE_DIR_NAME = "需求的模板"
PURCHASE_TEMPLATE = "{{合同编号}} {{物料名称1}} {{对方简称}} {{合同金额1}}.docx"
LEGACY_PURCHASE_TEMPLATE = "一般供应商模板.docx"
_TOKEN = re.compile(r"\{\{([^{}]+)\}\}")
_BAD = re.compile(r'[\\/:*?"<>|\r\n]+')
_EXT = re.compile(r"\.(docx|doc|dot|dotx|docm|xlsx|xls|xlsm|xlt|xltx|xltm|pdf|rtf|csv|txt)$", re.IGNORECASE)
ALLOWED_TEMPLATE_SUFFIXES = {
    ".doc", ".docx", ".dot", ".dotx", ".docm",
    ".xls", ".xlsx", ".xlsm", ".xlt", ".xltx", ".xltm",
    ".pdf", ".rtf", ".csv", ".txt",
}
OOXML_SUFFIXES = {".docx", ".dotx", ".docm", ".xlsx", ".xlsm", ".xltx", ".xltm"}
MEDIA_TYPES = {
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".dotx": "application/vnd.openxmlformats-officedocument.wordprocessingml.template",
    ".docm": "application/vnd.ms-word.document.macroEnabled.12",
    ".doc": "application/msword",
    ".dot": "application/msword",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".xlsm": "application/vnd.ms-excel.sheet.macroEnabled.12",
    ".xltx": "application/vnd.openxmlformats-officedocument.spreadsheetml.template",
    ".xltm": "application/vnd.ms-excel.template.macroEnabled.12",
    ".xls": "application/vnd.ms-excel",
    ".xlt": "application/vnd.ms-excel",
    ".pdf": "application/pdf",
    ".rtf": "application/rtf",
    ".csv": "text/csv",
    ".txt": "text/plain",
}
SUPPLIER_DOCS = {
    "purchase_contract", "purchase_order", "input_invoice", "payment", "purchase_in",
    "customs_invoice", "customs_clearance", "customs_element",
}


class Kind:
    def __init__(self, code: str, name: str, name_rule: str, placeholders: tuple[str, ...]):
        self.code = code
        self.name = name
        self.name_rule = name_rule
        self.placeholders = placeholders

    @property
    def party_type(self) -> str:
        return "supplier" if self.code in SUPPLIER_DOCS else "customer"


KINDS = (
    Kind("purchase_contract", "采购合同", PURCHASE_TEMPLATE, (
        "合同编号", "物料名称1", "对方简称", "对方名称", "我方名称", "我方简称", "签订日期", "标的名称", "合同金额1",
        "我方地址", "我方法人", "我方代理", "我方联系人", "我方联系电话", "我方开票电话", "我方银行账号", "我方开户支行", "我方税号",
        "对方地址", "对方法人", "对方代理", "对方联系人", "对方联系电话", "对方开票电话", "对方银行账号", "对方开户支行", "对方税号",
        "使用设备",
    )),
    Kind("sale_contract", "销售合同", "{{合同编号}}.docx", (
        "合同编号", "合同名称", "对方简称", "对方名称", "我方名称", "我方简称", "签订日期",
        "物料名称1", "标的名称", "合同金额1", "项目名称", "税率",
    )),
    Kind("purchase_order", "采购订单", "{{订单号}}.docx", ("订单号", "合同编号", "合同名称", "对方简称", "对方名称", "备注")),
    Kind("sale_order", "销售订单", "{{订单号}}.docx", ("订单号", "合同编号", "合同名称", "对方简称", "对方名称", "备注")),
    Kind("input_invoice", "进项发票", "{{发票号码}}.docx", ("发票号码", "单据编号", "对方简称", "客商名称", "对方名称", "发票日期", "合同编号")),
    Kind("output_invoice", "销项发票", "{{发票号码}}.docx", ("发票号码", "单据编号", "对方简称", "客商名称", "对方名称", "发票日期", "合同编号")),
    Kind("payment", "付款单", "{{单据编号}}.docx", ("单据编号", "单据日期", "对方简称", "对方名称", "合同编号", "合同金额", "账户")),
    Kind("receipt", "收款单", "{{单据编号}}.docx", ("单据编号", "单据日期", "对方简称", "对方名称", "合同编号", "合同金额", "账户")),
    Kind("purchase_in", "采购入库", "{{单据编号}}.docx", ("单据编号", "单据日期", "合同编号", "对方简称", "对方名称", "仓库", "使用设备")),
    Kind("sale_out", "销售出库", "{{单据编号}}.docx", ("单据编号", "单据日期", "合同编号", "对方简称", "对方名称", "仓库")),
    Kind("customs_invoice", "报关发票", "{{发票编号}}.docx", (
        "发票编号", "发票日期", "合同号", "总金额", "币种", "状态",
        "序号", "申报大类", "英文名称", "规格", "数量", "单位", "人民币单价", "人民币总价", "装箱数",
        "物料编号", "物料名称", "箱号", "集装箱号", "入仓单", "箱单号", "船名/航号",
    )),
    Kind("customs_clearance", "清关发票", "{{发票号}}.docx", (
        "发票号", "制单日期", "船名/航号", "船名", "航号", "集装箱号", "提单号", "出发港口", "出发港", "装箱号", "汇率", "人民币合计", "美元合计",
        "序号", "中文名", "英文名", "规格", "数量", "单位", "人民币单价", "人民币总价", "换算美元金额", "PO单号", "备注",
        "序列号", "品名", "包件数", "包装件数", "毛重", "净重", "尺寸", "入仓单号",
    )),
    Kind("customs_element", "报关要素", "{{报关发票编号}}-报关要素.docx", (
        "报关发票编号", "状态", "序号", "HS Code", "申报名称", "产品名称", "明细物料",
        "型号", "申报规格", "申报要素", "报关要素", "货源地", "数量",
    )),
)
_BY_CODE = {kind.code: kind for kind in KINDS}


def template_dir() -> Path:
    folder = (ROOT / TEMPLATE_DIR_NAME).resolve()
    folder.mkdir(parents=True, exist_ok=True)
    return folder


def _relax_unique():
    """旧表把一种单据限制成一行。改成同一单据可以挂多份模板。"""
    engine = get_engine()
    with engine.begin() as conn:
        row = conn.execute(text("SELECT sql FROM sqlite_master WHERE type = 'table' AND name = 'export_profile'")).fetchone()
    if row is None:
        return
    if "uq_export_doc" not in (row[0] or ""):
        ensure_columns("export_profile", {
            "template_path": "TEXT",
            "title": "VARCHAR(100)",
            "party_id": "INTEGER",
        })
        return
    with engine.begin() as conn:
        cols = {item[1] for item in conn.execute(text("PRAGMA table_info(export_profile)"))}
        path_sql = "template_path" if "template_path" in cols else "''"
        title_sql = "title" if "title" in cols else "''"
        party_sql = "party_id" if "party_id" in cols else "NULL"
        conn.execute(text("""
            CREATE TABLE export_profile_new (
                id INTEGER NOT NULL PRIMARY KEY,
                doc_code VARCHAR(40) NOT NULL,
                doc_name VARCHAR(40) NOT NULL,
                template_file VARCHAR(255) NOT NULL DEFAULT '',
                template_path TEXT NOT NULL DEFAULT '',
                title VARCHAR(100) NOT NULL DEFAULT '',
                party_id INTEGER,
                name_rule VARCHAR(255) NOT NULL DEFAULT '',
                sort_no INTEGER NOT NULL DEFAULT 0,
                company_id INTEGER NOT NULL,
                created_by VARCHAR(64) NOT NULL,
                created_at DATETIME NOT NULL,
                updated_by VARCHAR(64) NOT NULL,
                updated_at DATETIME NOT NULL,
                version INTEGER NOT NULL,
                deleted_at DATETIME,
                deleted_by VARCHAR(64),
                reserve1 VARCHAR(255),
                reserve2 VARCHAR(255),
                reserve3 VARCHAR(255)
            )
        """))
        conn.execute(text(f"""
            INSERT INTO export_profile_new (
                id, doc_code, doc_name, template_file, template_path, title, party_id, name_rule, sort_no,
                company_id, created_by, created_at, updated_by, updated_at, version,
                deleted_at, deleted_by, reserve1, reserve2, reserve3
            )
            SELECT
                id, doc_code, doc_name, template_file, {path_sql}, {title_sql}, {party_sql}, name_rule, sort_no,
                company_id, created_by, created_at, updated_by, updated_at, version,
                deleted_at, deleted_by, reserve1, reserve2, reserve3
            FROM export_profile
        """))
        conn.execute(text("DROP TABLE export_profile"))
        conn.execute(text("ALTER TABLE export_profile_new RENAME TO export_profile"))
        conn.execute(text("CREATE INDEX IF NOT EXISTS ix_export_profile_company_id ON export_profile (company_id)"))


def _default_file(kind: Kind) -> Path | None:
    if kind.code != "purchase_contract":
        return None
    folder = ROOT / TEMPLATE_DIR_NAME
    for name in (PURCHASE_TEMPLATE, LEGACY_PURCHASE_TEMPLATE):
        path = folder / name
        if path.is_file():
            return path.resolve()
    return None


def _fill_paths(db):
    folder = template_dir()
    rows = db.query(ExportProfile).filter(ExportProfile.deleted_at.is_(None)).all()
    for row in rows:
        if (row.template_path or "").strip():
            if not (row.title or "").strip():
                row.title = row.doc_name or ""
            continue
        name = (row.template_file or "").strip()
        candidate = None
        if name:
            raw = Path(name)
            candidate = raw if raw.is_absolute() else folder / name
        if candidate is None or not candidate.is_file():
            kind = _BY_CODE.get(row.doc_code)
            candidate = _default_file(kind) if kind and not name else candidate
        if candidate is not None and candidate.is_file():
            row.template_path = str(candidate.resolve())
            row.template_file = candidate.name
        if not (row.title or "").strip():
            row.title = row.doc_name or ""


def prepare_export_table():
    _relax_unique()


def ensure_export_profiles(db):
    companies = db.query(Company).filter(Company.deleted_at.is_(None)).all()
    for company in companies:
        existing = {
            row.doc_code
            for row in db.query(ExportProfile).filter(
                ExportProfile.company_id == company.id,
                ExportProfile.deleted_at.is_(None),
            ).all()
        }
        for index, kind in enumerate(KINDS):
            if kind.code in existing:
                continue
            path = _default_file(kind)
            row = ExportProfile(
                company_id=company.id,
                doc_code=kind.code,
                doc_name=kind.name,
                title=kind.name,
                template_file=path.name if path else "",
                template_path=str(path) if path else "",
                name_rule=kind.name_rule,
                sort_no=index,
                version=1,
                created_by="system",
                updated_by="system",
                created_at=now(),
                updated_at=now(),
            )
            db.add(row)
    db.flush()
    _fill_paths(db)


def _party_name(db, party_id) -> str:
    if not party_id:
        return ""
    party = db.get(Party, party_id)
    if party is None or party.deleted_at:
        return ""
    return party.short_name or party.name or ""


def profile_view(db, row: ExportProfile) -> dict:
    kind = _BY_CODE.get(row.doc_code)
    return {
        "id": row.id,
        "version": row.version,
        "doc_code": row.doc_code,
        "doc_name": row.doc_name,
        "title": row.title or row.doc_name or "",
        "template_path": row.template_path or "",
        "party_id": row.party_id,
        "party_name": _party_name(db, row.party_id),
        "name_rule": row.name_rule or "",
        "placeholders": list(kind.placeholders) if kind else [],
        "party_type": kind.party_type if kind else "supplier",
    }


def kind_views() -> list[dict]:
    return [
        {"code": kind.code, "name": kind.name, "placeholders": list(kind.placeholders), "party_type": kind.party_type, "name_rule": kind.name_rule}
        for kind in KINDS
    ]


def list_profiles(db, company_id: int, doc_code: str = "") -> dict:
    ensure_export_profiles(db)
    query = db.query(ExportProfile).filter(ExportProfile.company_id == company_id, ExportProfile.deleted_at.is_(None))
    if doc_code:
        query = query.filter(ExportProfile.doc_code == doc_code)
    rows = query.order_by(ExportProfile.sort_no, ExportProfile.id).all()
    return {"items": [profile_view(db, row) for row in rows], "kinds": kind_views()}


def _check_rule(rule: str):
    plain = _TOKEN.sub("", rule or "")
    if _BAD.search(plain):
        raise BizError('导出名称不能包含 \\ / : * ? " < > |')


def require_template_path(raw: str, *, for_save: bool = False) -> Path:
    text_path = (raw or "").strip().strip('"')
    if not text_path:
        raise BizError("请填写模板绝对地址")
    path = Path(text_path)
    if not path.is_absolute():
        raise BizError("请填写模板的绝对地址")
    path = path.resolve()
    if not path.is_file():
        raise BizError(f"找不到模板文件：{path}")
    suffix = path.suffix.lower()
    if suffix not in ALLOWED_TEMPLATE_SUFFIXES:
        raise BizError("模板可以是 doc、docx、xls、xlsx 等办公文件")
    return path


def align_export_name(filename: str, template: Path) -> str:
    suffix = template.suffix
    if not suffix:
        return filename
    current = Path(filename or "")
    if current.suffix.lower() == suffix.lower():
        return filename
    return f"{current.stem}{suffix}"


def media_type_for(filename: str) -> str:
    return MEDIA_TYPES.get(Path(filename or "").suffix.lower(), "application/octet-stream")


def _party_id(db, company_id: int, kind: Kind, raw) -> int | None:
    if raw in (None, "", 0, "0"):
        return None
    party = db.get(Party, int(raw))
    if party is None or party.deleted_at or party.company_id != company_id:
        raise BizError("客商不存在")
    if party.type != kind.party_type:
        raise BizError("适用客商和这张单据的方向不一致")
    return party.id


def _apply(db, user, company_id: int, row: ExportProfile, payload: dict):
    kind = _BY_CODE.get((payload.get("doc_code") or row.doc_code or "").strip())
    if kind is None:
        raise BizError("请选择单据")
    path = require_template_path(payload.get("template_path"), for_save=True)
    name_rule = (payload.get("name_rule") or "").strip()
    if not name_rule:
        raise BizError("请填写导出名称，用 {{字段}} 表示")
    _check_rule(name_rule)
    title = (payload.get("title") or "").strip() or path.stem
    row.doc_code = kind.code
    row.doc_name = kind.name
    row.title = title
    row.template_path = str(path)
    row.template_file = path.name
    row.party_id = _party_id(db, company_id, kind, payload.get("party_id"))
    row.name_rule = name_rule
    row.sort_no = KINDS.index(kind)


def save_profile(db, user, company_id: int, row_id: int | None, payload: dict) -> ExportProfile:
    from .util import bump, ensure_version, stamp_new

    if row_id:
        row = db.get(ExportProfile, row_id)
        if row is None or row.deleted_at or row.company_id != company_id:
            raise BizError("数据不存在")
        ensure_version(row, payload.get("version"))
        _apply(db, user, company_id, row, payload)
        bump(row, user.employee_id)
        return row
    row = ExportProfile(template_file="", template_path="", title="", name_rule="")
    stamp_new(row, user, company_id)
    _apply(db, user, company_id, row, payload)
    db.add(row)
    db.flush()
    return row


def remove_profile(db, user, company_id: int, row_id: int, version: int):
    from .util import ensure_version, soft_delete

    row = db.get(ExportProfile, row_id)
    if row is None or row.deleted_at or row.company_id != company_id:
        raise BizError("数据不存在")
    ensure_version(row, version)
    soft_delete(row, user.employee_id)


def save_upload(filename: str, content: bytes) -> str:
    name = Path(filename or "").name
    suffix = Path(name).suffix.lower()
    if suffix not in ALLOWED_TEMPLATE_SUFFIXES:
        raise BizError("请上传 doc、docx、xls、xlsx 等办公模板")
    if not name or name.startswith("~$") or name.startswith("."):
        raise BizError("模板文件名不正确")
    if len(content) > 20 * 1024 * 1024:
        raise BizError("模板文件不能超过 20MB")
    if not content:
        raise BizError("模板文件是空的")
    if suffix in OOXML_SUFFIXES and not content.startswith(b"PK"):
        raise BizError("这个文件内容和扩展名不一致")
    path = (template_dir() / name).resolve()
    if path.parent != template_dir():
        raise BizError("模板文件名不正确")
    path.write_bytes(content)
    return str(path)


def render_filename(rule: str, mapping: dict, fallback: str) -> str:
    source = (rule or "").strip() or fallback

    def repl(match):
        return str(mapping.get(match.group(1).strip(), "") or "")

    has_ext = bool(_EXT.search(source))
    text = _TOKEN.sub(repl, source)
    text = _BAD.sub("", text)
    text = re.sub(r"\s+", " ", text).strip().strip(".")
    if not text:
        text = Path(fallback).stem or "导出"
        has_ext = False
    if not has_ext:
        text += ".docx"
    return text


PURCHASE_FIELDS = (
    ("合同编号", "采购合同", "合同编号"),
    ("签订日期", "采购合同", "签订日期，写成 2026年9月26日"),
    ("标的名称", "采购合同", "标的名称；空着时用明细中文名称"),
    ("物料名称1", "采购合同", "与标的名称相同"),
    ("物料名称2", "采购合同", "与标的名称相同"),
    ("对方名称", "客商资料", "客商名称"),
    ("对方简称", "客商资料", "客商简称"),
    ("对方地址", "客商资料", "通讯地址"),
    ("对方法人", "客商资料", "法人代表"),
    ("对方代理", "客商资料", "委托代理人"),
    ("对方开票电话", "客商资料", "开票电话，没有则用手机"),
    ("对方银行账号", "客商资料", "银行账号"),
    ("对方开户支行", "客商资料", "开户银行或开户支行"),
    ("对方税号", "客商资料", "纳税人识别号"),
    ("我方名称", "公司资料", "公司名称"),
    ("我方简称", "公司资料", "公司简称"),
    ("我方地址", "公司资料", "通讯地址，没有则用开票地址"),
    ("我方法人", "公司资料", "法人代表"),
    ("我方代理", "公司资料", "委托代理人"),
    ("我方联系电话", "公司资料", "开票电话，没有则用手机或默认联系人"),
    ("我方银行账号", "公司资料", "银行账号"),
    ("我方开户支行", "公司资料", "开户银行或开户支行"),
    ("我方税号", "公司资料", "纳税人识别号"),
    ("合同金额1", "采购合同", "含税合同金额"),
    ("合同金额2", "采购合同", "含税合同金额"),
    ("合同金额3", "采购合同", "含税合同金额"),
    ("不含税金额", "采购合同", "不含税金额"),
    ("税率", "采购合同", "税率，例如 13%"),
    ("合同金额（大写）", "采购合同", "含税金额的人民币大写"),
    ("收货地址", "采购合同", "收货地址"),
    ("运输方式", "采购合同", "运输方式"),
    ("交货期限", "采购合同", "交货期限"),
    ("收货人", "采购合同", "收货人"),
    ("运费承担", "采购合同", "运费承担"),
    ("付款方式期限等条款", "采购合同", "付款条款"),
    ("货物验收标准和方法", "采购合同", "验收标准"),
    ("计划单号", "采购合同明细", "各行计划单号"),
    ("使用设备", "采购合同明细", "采购计划上的使用设备"),
    ("数量合计", "采购合同明细", "数量合计"),
    ("价税合计", "采购合同", "含税合同金额"),
    ("附件说明", "采购合同", "附件说明"),
    ("物料名称", "采购合同明细", "明细行中文名称"),
    ("英文名称", "采购合同明细", "明细行英文名称"),
    ("规格", "采购合同明细", "明细行规格"),
    ("数量", "采购合同明细", "明细行数量"),
    ("单位", "采购合同明细", "明细行单位"),
    ("含税单价", "采购合同明细", "明细行含税单价"),
    ("行价税合计", "采购合同明细", "明细行含税总价"),
)

CUSTOMS_FIELDS = (
    ("报关发票", "发票编号", "报关发票", "报关发票编号"),
    ("报关发票", "发票日期", "报关发票", "开票日期"),
    ("报关发票", "合同号", "报关发票", "合同号"),
    ("报关发票", "总金额", "报关发票", "人民币合计"),
    ("报关发票", "币种", "报关发票", "默认 CNY"),
    ("报关发票", "序号", "报关发票明细", "申报大类行号，表格按行展开"),
    ("报关发票", "申报大类", "报关发票明细", "合并后的物料大类/申报名称"),
    ("报关发票", "英文名称", "报关发票明细", "英文名称"),
    ("报关发票", "规格", "报关发票明细", "规格"),
    ("报关发票", "数量", "报关发票明细", "合并数量"),
    ("报关发票", "单位", "报关发票明细", "单位"),
    ("报关发票", "人民币单价", "报关发票明细", "离岸人民币单价"),
    ("报关发票", "人民币总价", "报关发票明细", "离岸人民币总价"),
    ("报关发票", "装箱数", "报关发票明细", "箱数"),
    ("报关发票", "物料编号", "来源物料", "原始物料编号，表格按行展开"),
    ("报关发票", "物料名称", "来源物料", "中文名称"),
    ("报关发票", "箱号", "来源物料", "箱号"),
    ("报关发票", "集装箱号", "来源物料", "集装箱号"),
    ("报关发票", "入仓单", "来源物料", "入仓单号"),
    ("报关发票", "箱单号", "来源物料", "箱单号"),
    ("报关发票", "船名/航号", "来源物料", "船名/航号"),
    ("清关发票", "发票号", "清关发票", "可为空，空着仍可编辑"),
    ("清关发票", "制单日期", "清关发票", "制单日期"),
    ("清关发票", "船名/航号", "清关发票", "合并字段"),
    ("清关发票", "船名", "清关发票", "船名"),
    ("清关发票", "航号", "清关发票", "航号"),
    ("清关发票", "集装箱号", "清关发票", "集装箱号"),
    ("清关发票", "提单号", "清关发票", "提单号"),
    ("清关发票", "出发港口", "清关发票", "商业发票 Port of shipment"),
    ("清关发票", "出发港", "清关发票", "与出发港口相同，对应箱单 From"),
    ("清关发票", "装箱号", "清关发票", "箱单号"),
    ("清关发票", "汇率", "清关发票", "当月美元汇率"),
    ("清关发票", "人民币合计", "清关发票", "人民币合计"),
    ("清关发票", "美元合计", "清关发票", "按汇率换算"),
    ("清关发票", "序号", "清关发票明细", "明细行号，表格按行展开"),
    ("清关发票", "中文名", "清关发票明细", "中文名称"),
    ("清关发票", "英文名", "清关发票明细", "英文名称"),
    ("清关发票", "规格", "清关发票明细", "规格"),
    ("清关发票", "数量", "清关发票明细", "数量"),
    ("清关发票", "单位", "清关发票明细", "单位"),
    ("清关发票", "人民币单价", "清关发票明细", "离岸人民币单价"),
    ("清关发票", "人民币总价", "清关发票明细", "离岸人民币总价"),
    ("清关发票", "换算美元金额", "清关发票明细", "行美元金额"),
    ("清关发票", "PO单号", "清关发票明细", "PO 单号"),
    ("清关发票", "备注", "清关发票明细", "备注"),
    ("清关发票", "序列号", "物料大类清单", "装箱清单行号，表格按行展开"),
    ("清关发票", "品名", "物料大类清单", "物料大类。箱单占位符也写作中文名"),
    ("清关发票", "包件数", "物料大类清单", "包装件数"),
    ("清关发票", "包装件数", "物料大类清单", "与包件数相同"),
    ("清关发票", "毛重", "物料大类清单", "毛重合计"),
    ("清关发票", "净重", "物料大类清单", "净重合计"),
    ("清关发票", "尺寸", "物料大类清单", "长×宽×高"),
    ("清关发票", "入仓单号", "物料大类清单", "来自报关物料"),
    ("报关要素", "报关发票编号", "报关要素", "关联的报关发票编号"),
    ("报关要素", "状态", "报关要素", "生成状态"),
    ("报关要素", "序号", "报关要素明细", "行号，同一申报名称下多行产品时重复"),
    ("报关要素", "HS Code", "报关要素明细", "HS 编码"),
    ("报关要素", "申报名称", "报关要素明细", "申报名称"),
    ("报关要素", "产品名称", "报关要素明细", "表单字段。模板占位符写作明细物料"),
    ("报关要素", "明细物料", "报关要素明细", "与产品名称相同，对应模板 {{明细物料}}"),
    ("报关要素", "型号", "报关要素明细", "表单字段。模板占位符写作申报规格"),
    ("报关要素", "申报规格", "报关要素明细", "与型号相同，对应模板 {{申报规格}}"),
    ("报关要素", "申报要素", "报关要素明细", "表单字段。模板占位符写作报关要素"),
    ("报关要素", "报关要素", "报关要素明细", "与申报要素相同，对应模板 {{报关要素}}"),
    ("报关要素", "货源地", "报关要素明细", "来自报关物料产地"),
    ("报关要素", "数量", "报关要素明细", "数量"),
)


# 每张表单一行字段。顺序与页面一致：区块、字段名称、字段标识、必填、说明。
_PARTY_CARD = (
    ("客商信息", "纳税人识别号", "tax_no", "只读", "选定客商后展示，来自客商资料"),
    ("客商信息", "开票名称", "invoice_title", "只读", "选定客商后展示"),
    ("客商信息", "开户银行", "bank_name", "只读", "选定客商后展示"),
    ("客商信息", "账户名称", "account_name", "只读", "选定客商后展示"),
    ("客商信息", "银行账号", "bank_account", "只读", "选定客商后展示"),
    ("客商信息", "开户支行", "bank_branch", "只读", "选定客商后展示"),
    ("客商信息", "开票电话", "invoice_phone", "只读", "选定客商后展示"),
    ("客商信息", "法人代表", "legal_name", "只读", "选定客商后展示"),
    ("客商信息", "委托代理人", "agent_name", "只读", "选定客商后展示"),
    ("客商信息", "联系人", "contact_name", "只读", "默认联系人"),
    ("客商信息", "联系电话", "contact_phone", "只读", "默认联系人电话"),
    ("客商信息", "开票地址", "invoice_address", "只读", "选定客商后展示"),
    ("客商信息", "通讯地址", "address", "只读", "选定客商后展示"),
)


def _contract_fields(purchase: bool):
    party = "供应商" if purchase else "客户"
    terms = "付款条款" if purchase else "条款"
    stage_flag = "到货阶段" if purchase else "发货阶段"
    rows = [
        ("基本信息", "合同编号", "contract_code", "否", "留空自动生成。采购前缀 CG-年份-流水，销售前缀 XS-年份-流水"),
        ("基本信息", "合同名称", "name", "是", ""),
        ("基本信息", "合同类型", "type_name", "否", "从合同类型中选择"),
        ("基本信息", party, "party_id", "是", "从客商主数据选择，保存名称和税号快照"),
        *_PARTY_CARD,
        ("基本信息", "业务员", "owner_id", "是", "按员工编号选择，页面显示姓名"),
        ("基本信息", "项目", "project_id", "否", "从项目台账选择，带出项目编号、名称、年度"),
        ("基本信息", "合同状态", "status", "是", "未签订 / 已签订 / 已终止，默认已签订"),
        ("基本信息", "归档", "archived", "只读", "当前版本恒为未归档"),
        ("基本信息", "计价方式", "pricing", "是", "固定总价 / 按实结算，默认固定总价"),
        ("基本信息", "税率", "tax_rate", "是", "默认 13，表示 13%"),
        ("基本信息", "签订日期", "sign_date", "否", ""),
        ("基本信息", "备注", "remark", "否", ""),
    ]
    if purchase:
        rows.extend([
            ("基本信息", "标的名称", "subject_name", "否", "写入合同首段和汇总行，留空则用物料名称"),
            ("基本信息", "收货地址", "ship_address", "否", ""),
            ("基本信息", "收货人", "receiver_name", "否", ""),
            ("基本信息", "运输方式", "transport_mode", "否", ""),
            ("基本信息", "交货期限", "delivery_term", "否", ""),
            ("基本信息", "运费承担", "freight_bearer", "否", ""),
            ("基本信息", "附件说明", "annex_note", "否", "例如车型、车架号，留空则不写入附件"),
        ])
    rows.extend([
        ("基本信息", "内容概述", "content", "否", ""),
        ("基本信息", terms, "terms", "否", "采购写入合同第四节：付款方式、期限及发票" if purchase else ""),
    ])
    if purchase:
        rows.append(("基本信息", "验收标准", "acceptance_text", "否", ""))
    rows.extend([
        ("物料明细", "物料编号", "material_code", "是", "同一合同内未删除行不能重复"),
        ("物料明细", "中文名称", "name_cn", "否", ""),
    ])
    if purchase:
        rows.extend([
            ("物料明细", "英文名称", "name_en", "否", ""),
            ("物料明细", "计划单号", "plan_no", "否", "从采购计划带入或手填"),
            ("物料明细", "使用设备", "category", "否", "从采购计划带入"),
        ])
    rows.extend([
        ("物料明细", "规格", "spec", "否", ""),
        ("物料明细", "单位", "unit", "否", ""),
        ("物料明细", "数量", "qty", "是", "必须大于 0。按实结算时签约数量不被出入库改写"),
    ])
    if not purchase:
        rows.append(("物料明细", "采购单价", "cost", "只读", "由采购合同生成销售合同时显示"))
    rows.extend([
        ("物料明细", "含税单价", "price", "是", "必须大于等于 0"),
        ("物料明细", "含税总价", "amount", "只读", "固定总价按签约数量计算，按实结算按已出入库数量计算"),
        ("物料明细", "合同金额", "amount", "只读", "有物料时为各行含税总价合计"),
        ("阶段付款", stage_flag, "arrival_stage_id" if purchase else "ship_stage_id", "否", "单选一个阶段，默认第一行"),
        ("阶段付款", "阶段名称", "name", "是", "采购可选预付款、进度款、到货款、验收款、尾款、质保金" if purchase else "销售可选预收款、进度款、验收款、尾款、质保金"),
        ("阶段付款", "比例", "ratio", "否", "只用于分摊，不强制合计 100%"),
        ("阶段付款", "计划金额", "plan_amount", "是", "各阶段合计必须等于合同金额"),
        ("阶段付款", "计划日期", "plan_date", "否", ""),
        ("阶段付款", "已收/已付", "settled_amount", "只读", "由收付款单回写"),
        ("阶段付款", "超收/超付", "over", "只读", "已收付超出计划金额的部分"),
        ("阶段付款", "备注", "remark", "否", ""),
    ])
    return tuple(rows)


_ORDER_FIELDS = (
    ("基本信息", "源合同", "contract_id", "是", "只能选已签订且尚未有未删除订单的合同"),
    ("基本信息", "订单号", "order_code", "只读", "保存时生成，首次为 order 加合同编号"),
    ("基本信息", "备注", "remark", "否", "空着时用合同内容概述"),
    ("物料明细", "物料编号", "material_code", "只读", "从合同物料带出"),
    ("物料明细", "中文名称", "name_cn", "只读", "从合同物料带出"),
    ("物料明细", "数量", "qty", "是", "必须大于 0"),
    ("物料明细", "含税单价", "price", "是", ""),
    ("物料明细", "到货 / 发货", "moved_text", "只读", "按该合同出入库累计"),
)

_INVOICE_HEAD = (
    ("基本信息", "客商", "party_id", "是", ""),
    ("基本信息", "合同", "contract_id", "是", "方向须与发票一致"),
    ("基本信息", "发票号码", "invoice_no", "是", "公司内唯一"),
    ("基本信息", "发票日期", "invoice_date", "是", ""),
    ("基本信息", "发票类型", "invoice_type", "是", "增值税专用发票 / 增值税普通发票 / 电子专票 / 电子普票，默认专票"),
)


def _invoice_fields(input_side: bool):
    party_name = "销售方名称" if input_side else "购买方名称"
    return _INVOICE_HEAD + (
        ("基本信息", party_name, "party_name", "否", "选客商后带出，可改"),
        ("基本信息", "税号", "party_tax_no", "否", "选客商后带出，可改"),
        ("物料明细", "物料编号", "material_code", "是", "同一发票内未删除行不能重复"),
        ("物料明细", "名称", "name_cn", "否", ""),
        ("物料明细", "数量", "qty", "是", "必须大于 0"),
        ("物料明细", "含税单价", "price", "是", ""),
        ("物料明细", "不含税金额", "amount_ex_tax", "否", "可手改，改后按不含税金额和税额重算价税合计"),
        ("物料明细", "税额", "tax", "否", "可手改"),
        ("物料明细", "税率", "tax_rate", "否", "默认带出合同税率"),
    )


def _settle_fields(payment: bool):
    verb = "付" if payment else "收"
    return (
        ("基本信息", "客商", "party_id", "是", ""),
        ("基本信息", "合同", "contract_id", "是", "只列出该客商下未结清的合同"),
        ("基本信息", "单号", "bill_code", "是", "一次勾选多个阶段时，在单号后加 -01、-02 拆成多张"),
        ("基本信息", "日期", "bill_date", "是", ""),
        ("基本信息", "账户", "account_name", "否", "新增时只显示启用的资金账户"),
        ("阶段明细", "阶段", "stage_name", "只读", "来自合同阶段"),
        ("阶段明细", "计划金额", "plan_amount", "只读", ""),
        ("阶段明细", "已结", "settled_amount", "只读", ""),
        ("阶段明细", "剩余", "remain", "只读", "计划金额减已结，小于 0 时按 0"),
        ("阶段明细", "本次", "amount", "是", f"本次{verb}款金额，必须大于 0 且不超过剩余"),
    )


def _wms_fields(inbound: bool):
    rows = [
        ("基本信息", "单据编号", "bill_code", "否", "留空按前缀-日期-流水自动生成"),
        ("基本信息", "单据日期", "bill_date", "是", ""),
        ("基本信息", "采购合同" if inbound else "销售合同", "contract_id", "是", "只选已签订合同"),
        ("基本信息", "仓库", "warehouse_id", "是", "只选启用仓库"),
        ("基本信息", "备注", "remark", "否", ""),
        ("基本信息", "供应商" if inbound else "客户", "party_name", "只读", "选定合同后带出"),
        ("物料明细", "物料编号", "material_code", "只读", "从合同物料带出，同一单内不重复"),
    ]
    if inbound:
        rows.append(("物料明细", "使用设备", "category", "只读", "从采购合同物料带出"))
    rows.extend([
        ("物料明细", "中文名称", "name_cn", "只读", ""),
        ("物料明细", "单位", "unit", "只读", ""),
        ("物料明细", "数量", "qty", "是", "必须大于 0；不入不出的行直接删除"),
        ("物料明细", "含税单价", "price", "是", ""),
    ])
    return tuple(rows)


_CUSTOMS_MATERIAL_FIELDS = (
    ("物料信息", "物料编号", "material_code", "是", "来自采购入库。批量导入时用它识别要填写的行"),
    ("物料信息", "物料名称", "material_name", "是", ""),
    ("物料信息", "英文名称", "material_name_en", "否", ""),
    ("物料信息", "物料大类", "material_category", "是", "在报关物料清单中手工填写或批量赋值，采购入库不会带入"),
    ("物料信息", "规格", "specification", "否", ""),
    ("物料信息", "单位", "unit", "否", ""),
    ("物料信息", "数量", "quantity", "否", "来自采购入库"),
    ("物料信息", "产地", "origin_country", "是", "可在清单中批量导入"),
    ("物料信息", "使用设备", "equipment", "是", ""),
    ("物流信息", "合同号", "contract_no", "是", "来自采购合同，只读"),
    ("物流信息", "箱单号", "packing_list_no", "是", ""),
    ("物流信息", "入仓单号", "receipt_no", "否", "与采购入库单号不是同一字段，采购入库不会自动填入。可手工填写或批量赋值；批量导入时用来区分同一物料编号的多条"),
    ("物流信息", "供应商", "supplier_name", "只读", "来自采购入库"),
    ("物流信息", "入仓日期", "receipt_date", "只读", "来自采购入库"),
    ("物流信息", "箱号", "box_no", "否", "可在清单中批量导入"),
    ("物流信息", "集装箱号", "container_no", "否", ""),
    ("物流信息", "提单号", "bill_no", "否", ""),
    ("物流信息", "船名/航号", "vessel_voyage", "否", ""),
    ("物流信息", "出发港", "port_of_shipment", "否", "可在清单中批量导入"),
    ("物流信息", "目的港", "port_of_destination", "否", ""),
    ("物流信息", "收货人", "consignee", "否", ""),
    ("重量尺寸", "毛重", "gross_weight", "是", "可在清单中批量导入"),
    ("重量尺寸", "净重", "net_weight", "是", "可在清单中批量导入"),
    ("重量尺寸", "长", "length", "否", "尺寸按长*宽*高填写，可批量导入"),
    ("重量尺寸", "宽", "width", "否", "尺寸的宽"),
    ("重量尺寸", "高", "height", "否", "尺寸的高"),
    ("重量尺寸", "体积", "volume", "是", "可在清单中批量导入"),
    ("重量尺寸", "包件件数", "package_count", "否", ""),
    ("价格与申报", "PO单号", "po_no", "否", ""),
    ("价格与申报", "含税单价", "price", "否", ""),
    ("价格与申报", "询价名称", "inquiry_name", "否", ""),
    ("价格与申报", "询价型号", "inquiry_model", "否", ""),
    ("价格与申报", "询价数量", "inquiry_qty", "否", ""),
    ("价格与申报", "HS Code", "hs_code", "否", ""),
    ("价格与申报", "特殊申报", "special_declare", "否", "勾选后生成报关发票时单独成行"),
    ("价格与申报", "付款条款", "payment_terms", "否", ""),
    ("价格与申报", "申报要素", "declaration_elements", "否", ""),
    ("价格与申报", "备注", "remark", "否", ""),
    ("清单", "状态", "status", "只读", ""),
    ("清单", "缺失字段", "missing_text", "只读", "必填未填时列出字段名称"),
    ("单据状态", "报关发票", "invoice_generated", "只读", "已生成 / 未生成"),
    ("单据状态", "清关发票", "clearance_generated", "只读", "已生成 / 未生成"),
    ("单据状态", "报关要素", "element_generated", "只读", "已生成 / 未生成"),
    ("单据状态", "报关完毕", "customs_done", "只读", "三张单据都已生成时为已报关"),
)

_CUSTOMS_INVOICE_FIELDS = (
    ("主页面", "发票编号", "invoice_no", "只读", "对应商业发票发票编号"),
    ("主页面", "发票日期1", "invoice_date_1", "否", "对应商业发票 Date。一键导出后仍可改"),
    ("主页面", "发票日期2", "invoice_date_2", "隐藏", "页面不显示，对应箱单日期，保存时不能清空"),
    ("主页面", "船名/航号", "vessel_voyage", "否", "来自报关物料清单"),
    ("主页面", "集装箱号", "container_no", "否", "多个用逗号分隔"),
    ("主页面", "提单号", "bl_no", "否", ""),
    ("主页面", "出发港", "port_of_shipment", "否", "默认 SHANGHAI, CHINA"),
    ("主页面", "目的港", "port_of_destination", "否", "默认 GEORGETOWN, GUYANA"),
    ("主页面", "报关要素", "customs_elements", "否", "自定义文本"),
    ("主页面", "报关要素关联", "element_status", "只读", "未关联 / 已关联"),
    ("主页面", "状态", "status", "只读", "一键导出或生成表单后记为已导出，主页面仍可改"),
    ("主页面", "总金额", "total_amount", "只读", "人民币合计"),
    ("主页面", "美元金额", "total_usd", "只读", "人民币合计除以汇率"),
    ("主页面", "汇率", "exchange_rate", "否", "开票当月第一个工作日汇率"),
    ("主页面", "箱单号", "packing_list_no", "否", ""),
    ("主页面", "合同编号", "contract_no", "否", ""),
    ("主页面", "PO NO.", "po_no", "否", ""),
    ("主页面", "付款条款", "payment_terms", "否", ""),
    ("主页面", "收货人", "consignee", "否", "默认 Bosai Minerals Group (Guyana) Inc."),
    ("物料大类详情", "序列号", "seq", "只读", "从 1 开始"),
    ("物料大类详情", "物料大类", "material_category", "只读", "按报关物料清单汇总"),
    ("物料大类详情", "规格", "specification", "只读", ""),
    ("物料大类详情", "数量", "quantity", "只读", ""),
    ("物料大类详情", "单位", "unit", "只读", ""),
    ("物料大类详情", "包件件数", "package_count", "只读", "有包件件数时相加，否则按箱号去重"),
    ("物料大类详情", "毛重", "gross_weight", "只读", "千克"),
    ("物料大类详情", "净重", "net_weight", "只读", "千克"),
    ("物料大类详情", "尺寸", "size_display", "只读", "长*宽*高，不计算体积"),
    ("物料大类详情", "重量", "weight", "只读", "按毛重合计"),
    ("物料明细表", "编号", "seq", "只读", "每个集装箱从 1 重新编号"),
    ("物料明细表", "中文名称", "chinese_name", "只读", ""),
    ("物料明细表", "英文名称", "english_name", "只读", ""),
    ("物料明细表", "型号", "model", "只读", "取物料规格"),
    ("物料明细表", "单位", "unit", "只读", ""),
    ("物料明细表", "数量", "quantity", "只读", ""),
    ("物料明细表", "FOB单价", "fob_unit_price", "只读", "物料含税单价"),
    ("物料明细表", "FOB总价", "fob_total_price", "只读", "单价乘数量"),
    ("物料明细表", "Convert to USD", "convert_to_usd", "只读", "FOB总价除以汇率"),
)

_CLEARANCE_FIELDS = (
    ("主表", "发票号", "invoice_no", "否", "可空。未填或已导出仍可编辑"),
    ("主表", "制单日期", "invoice_date", "否", "对应商业发票 Date"),
    ("主表", "状态", "status", "只读", "已导出只是标记，不锁编辑"),
    ("主表", "船名/航号", "vessel_voyage", "否", "对应 Ocean Vessel/Voy. No."),
    ("主表", "集装箱号", "container_no", "否", ""),
    ("主表", "提单号", "bill_no", "否", ""),
    ("主表", "出发港口", "port_of_shipment", "否", "默认 SHANGHAI, CHINA。箱单出发港用同一值"),
    ("主表", "装箱号", "packing_list_no", "否", "对应箱单 PACKING LIST NO."),
    ("主表", "汇率", "exchange_rate", "否", "取当月美元汇率，可改"),
    ("主表", "人民币合计", "total_amount", "只读", "明细人民币总价合计"),
    ("主表", "美元合计", "total_usd", "只读", "按汇率换算"),
    ("物料明细", "序号", "seq", "只读", ""),
    ("物料明细", "中文名", "name_cn", "否", ""),
    ("物料明细", "英文名", "name_en", "否", ""),
    ("物料明细", "规格", "specification", "否", ""),
    ("物料明细", "数量", "quantity", "否", ""),
    ("物料明细", "单位", "unit", "否", ""),
    ("物料明细", "人民币单价", "price", "否", "改单价或数量后重算总价和美元金额"),
    ("物料明细", "人民币总价", "amount", "只读", "单价乘数量"),
    ("物料明细", "换算美元金额", "amount_usd", "只读", "人民币总价除以汇率"),
    ("物料明细", "PO单号", "po_no", "否", ""),
    ("物料明细", "备注", "remark", "否", "来自报关物料备注，可改"),
    ("物料大类清单", "序列号", "seq", "只读", ""),
    ("物料大类清单", "品名", "category", "否", "物料大类"),
    ("物料大类清单", "规格", "specification", "否", ""),
    ("物料大类清单", "数量", "quantity", "否", ""),
    ("物料大类清单", "单位", "unit", "否", ""),
    ("物料大类清单", "包装件数", "packages", "否", ""),
    ("物料大类清单", "毛重", "gross_weight", "否", "千克"),
    ("物料大类清单", "净重", "net_weight", "否", "千克"),
    ("物料大类清单", "尺寸", "size_text", "否", "长×宽×高，不计算体积"),
    ("物料大类清单", "入仓单号", "receipt_no", "否", "来自报关物料"),
)

_ELEMENT_FIELDS = (
    ("基本信息", "关联报关发票", "invoice_id", "是", "已关联过的报关发票不再出现"),
    ("基本信息", "状态", "status", "只读", ""),
    ("明细", "序号", "seq", "只读", "同一申报名称下多行产品时合并显示"),
    ("明细", "HS CODE", "hs_code", "否", "可手改，也可按 HS Code 库匹配"),
    ("明细", "申报名称", "declaration_name", "否", "同一申报名称合并多行产品"),
    ("明细", "产品名称", "product_name", "否", "对应模板明细物料，来自报关物料名称"),
    ("明细", "型号", "model", "否", "模板列名。导出占位符仍是申报规格"),
    ("明细", "申报要素", "declaration_elements", "否", "模板列名。导出占位符仍是报关要素"),
    ("明细", "货源地", "origin_place", "否", "来自报关物料产地"),
    ("明细", "数量", "quantity", "只读", "从报关发票带出，用于和发票核对"),
    ("明细", "匹配", "match_source", "只读", "HS 匹配来源"),
)

_PARTY_FIELDS = (
    ("基本信息", "编码", "code", "是", "供应商页为供应商编码，客户页为客户编码；公司内唯一"),
    ("基本信息", "简称", "short_name", "否", ""),
    ("基本信息", "名称", "name", "是", ""),
    ("基本信息", "业务员", "salesman_id", "否", "填写员工编号或姓名"),
    ("基本信息", "状态", "status", "是", "启用 / 停用，默认启用。停用后不能被新单据选中"),
    ("基本信息", "默认税率", "default_tax_rate", "否", "默认 13。新增合同时带出，合同上仍可手改"),
    ("基本信息", "类型", "type", "只读", "由页面上的供应商 / 客户决定，保存后不可改"),
    ("联系方式", "联系人", "contacts.name", "否", "可加多行"),
    ("联系方式", "联系电话", "contacts.phone", "否", ""),
    ("联系方式", "默认", "contacts.is_default", "否", "同一客商最多一个默认联系人"),
    ("联系方式", "手机", "mobile", "否", ""),
    ("联系方式", "邮箱", "email", "否", ""),
    ("开票与银行", "纳税人识别号", "tax_no", "否", ""),
    ("开票与银行", "开票名称", "invoice_title", "否", ""),
    ("开票与银行", "开户银行", "bank_name", "否", ""),
    ("开票与银行", "账户名称", "account_name", "否", ""),
    ("开票与银行", "银行账号", "bank_account", "否", ""),
    ("开票与银行", "开户支行", "bank_branch", "否", ""),
    ("开票与银行", "开票电话", "invoice_phone", "否", ""),
    ("开票与银行", "开票地址", "invoice_address", "否", ""),
    ("开票与银行", "法人代表", "legal_name", "否", ""),
    ("开票与银行", "委托代理人", "agent_name", "否", ""),
    ("开票与银行", "通讯地址", "address", "否", ""),
    ("开票与银行", "备注", "remark", "否", ""),
)

_PLAN_FIELDS = (
    ("计划行", "计划月份", "plan_month", "否", "该行的来源月份，不按月拆成多行"),
    ("计划行", "采购序号", "purchase_seq", "否", "保持导入原值，不重新编号"),
    ("计划行", "物料编号", "material_code", "是", "公司内唯一，同一编号只有一行计划"),
    ("计划行", "采购计划单号", "plan_no", "否", ""),
    ("计划行", "使用设备", "category", "否", ""),
    ("计划行", "中文名称", "name_cn", "否", ""),
    ("计划行", "英文名称", "name_en", "否", ""),
    ("计划行", "规格型号", "spec", "否", ""),
    ("计划行", "单位", "unit", "否", ""),
    ("计划行", "采购数量", "qty", "是", ""),
    ("计划行", "用途", "purpose", "否", ""),
    ("计划行", "使用部门", "use_dept", "否", "有值才在列表显示"),
    ("计划行", "备注", "remark", "否", ""),
    ("计划行", "库存数量", "stock_qty", "否", "仅导入列，不改库存余额"),
    ("计划行", "在途数量", "transit_qty", "否", "仅导入列，不改库存余额"),
    ("计划行", "4个月消耗量", "consume_4m", "否", "仅导入列"),
    ("计划行", "已采购数量", "bought_qty", "只读", "全部采购合同中该物料编号的数量合计"),
    ("计划行", "待采购数量", "remain_qty", "只读", "采购数量减已采购数量"),
    ("计划行", "是否采购", "purchase_state", "只读", "按待采购数量判断"),
)

_PROJECT_FIELDS = (
    ("基本信息", "项目编号", "project_code", "是", "公司内唯一"),
    ("基本信息", "项目名称", "name", "是", ""),
    ("基本信息", "项目类型", "project_type", "否", ""),
    ("基本信息", "项目来源", "source", "否", ""),
    ("基本信息", "项目年度", "year", "否", ""),
    ("基本信息", "客户名称", "customer_name", "否", "手填文本，不关联客商"),
    ("基本信息", "状态", "status", "是", "进行中 / 已完成 / 已终止，默认进行中"),
    ("基本信息", "开始日期", "start_date", "否", ""),
    ("基本信息", "结束日期", "end_date", "否", ""),
    ("基本信息", "备注", "remark", "否", ""),
    ("基本信息", "关联合同", "contracts", "只读", "反查挂了本项目的合同编号，不在项目上保存"),
)


def _placeholder_block(code: str | None):
    if not code:
        return ()
    if code == "purchase_contract":
        return tuple((name, f"{source}。{note}") for name, source, note in PURCHASE_FIELDS)
    labels = {"customs_invoice": "报关发票", "customs_clearance": "清关发票", "customs_element": "报关要素"}
    if code in labels:
        doc = labels[code]
        return tuple((name, f"{source}。{note}") for item_doc, name, source, note in CUSTOMS_FIELDS if item_doc == doc)
    kind = _BY_CODE[code]
    return tuple((name, "可写在导出名称或 Word 模板里") for name in kind.placeholders)


# 工作表顺序：导出模板里的单据在前，客商、计划和项目附在后面。
FORM_SHEETS = (
    ("采购合同", _contract_fields(True), "purchase_contract"),
    ("销售合同", _contract_fields(False), "sale_contract"),
    ("采购订单", _ORDER_FIELDS, "purchase_order"),
    ("销售订单", _ORDER_FIELDS, "sale_order"),
    ("进项发票", _invoice_fields(True), "input_invoice"),
    ("销项发票", _invoice_fields(False), "output_invoice"),
    ("付款单", _settle_fields(True), "payment"),
    ("收款单", _settle_fields(False), "receipt"),
    ("采购入库", _wms_fields(True), "purchase_in"),
    ("销售出库", _wms_fields(False), "sale_out"),
    ("报关物料清单", _CUSTOMS_MATERIAL_FIELDS, None),
    ("报关发票", _CUSTOMS_INVOICE_FIELDS, "customs_invoice"),
    ("清关发票", _CLEARANCE_FIELDS, "customs_clearance"),
    ("报关要素", _ELEMENT_FIELDS, "customs_element"),
    ("客商管理", _PARTY_FIELDS, None),
    ("采购计划", _PLAN_FIELDS, None),
    ("项目", _PROJECT_FIELDS, None),
)


def build_field_workbook() -> bytes:
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    headers = ("序号", "区块", "字段名称", "字段标识", "必填", "说明")
    widths = (8, 16, 22, 24, 10, 62)
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF")
    thin = Border(
        left=Side(style="thin", color="D0D7E2"),
        right=Side(style="thin", color="D0D7E2"),
        top=Side(style="thin", color="D0D7E2"),
        bottom=Side(style="thin", color="D0D7E2"),
    )
    wrap = Alignment(vertical="center", wrap_text=True)
    book = Workbook()
    for index, (title, fields, code) in enumerate(FORM_SHEETS):
        sheet = book.active if index == 0 else book.create_sheet(title)
        sheet.title = title
        sheet.append(list(headers))
        seq = 0
        for section, name, key, required, note in fields:
            seq += 1
            sheet.append([seq, section, name, key, required, note])
        for name, note in _placeholder_block(code):
            seq += 1
            sheet.append([seq, "导出占位符", name, "{{" + name + "}}", "—", note])
        for cell in sheet[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center", horizontal="center")
        for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row, max_col=len(headers)):
            for cell in row:
                cell.alignment = wrap
                cell.border = thin
            row[0].alignment = Alignment(vertical="center", horizontal="center")
            row[4].alignment = Alignment(vertical="center", horizontal="center")
        for col, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(col)].width = width
        sheet.auto_filter.ref = f"A1:F{sheet.max_row}"
        sheet.freeze_panes = "A2"
        sheet.row_dimensions[1].height = 22
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.oddHeader.left.text = title
        sheet.oddFooter.right.text = "第 &P 页"
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def resolve_export(db, company_id: int, doc_code: str, template_id: int | None = None) -> tuple[Path, str]:
    """按基础资料里的导出模板取出绝对地址和导出名称。"""
    kind = _BY_CODE.get(doc_code)
    if kind is None:
        raise BizError("请选择单据")
    ensure_export_profiles(db)
    query = db.query(ExportProfile).filter(
        ExportProfile.company_id == company_id,
        ExportProfile.doc_code == kind.code,
        ExportProfile.deleted_at.is_(None),
    )
    if template_id:
        row = query.filter(ExportProfile.id == int(template_id)).first()
        if row is None:
            raise BizError("模板不存在")
    else:
        rows = [item for item in query.all() if (item.template_path or "").strip()]
        if not rows:
            raise BizError(f"请先在基础资料的导出模板里填写{kind.name}的模板绝对地址")
        if len(rows) > 1:
            raise BizError(f"这份{kind.name}有多份模板，请先选择")
        row = rows[0]
    path = require_template_path(row.template_path)
    return path, row.name_rule or ""


def purchase_export(db, company_id: int, template_id: int | None = None) -> tuple[Path, str]:
    return resolve_export(db, company_id, "purchase_contract", template_id)
