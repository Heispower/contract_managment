"""列表导出：按当前筛选出 Excel，勾了行只导出勾选行。"""
import csv
import json
from io import BytesIO, StringIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from starlette.datastructures import QueryParams

from ..errors import BizError

_SKIP = {"ids", "columns", "format", "title", "page", "pageSize", "scope"}
_FORMULA = ("=", "+", "-", "@")
_HIDE = {"id", "version", "party_id", "project_id", "contract_id", "order_id", "warehouse_id", "stage_id", "warehouse_out", "warehouse_in"}


class QueryRequest:
    def __init__(self, params: dict):
        clean = {}
        for key, value in params.items():
            if key in _SKIP or key.startswith("_") or value in (None, "", [], {}):
                continue
            if isinstance(value, list):
                clean[key] = ",".join(str(item) for item in value)
            else:
                clean[key] = str(value)
        clean["page"] = "1"
        clean["pageSize"] = "100000"
        self.query_params = QueryParams(clean)


def _named_filename(params, fallback):
    title = (params.get("title") or "").strip()
    if title:
        return title
    direction = params.get("direction") or ""
    names = {
        ("contract", "sale"): "销售合同",
        ("contract", "purchase"): "采购合同",
        ("order", "sale"): "销售订单",
        ("order", "purchase"): "采购订单",
        ("invoice", "output"): "销项发票",
        ("invoice", "input"): "进项发票",
        ("settle", "receipt"): "收款单",
        ("settle", "payment"): "付款单",
        ("wms", "inbound"): "采购入库",
        ("wms", "outbound"): "销售出库",
        ("pending", "purchase"): "待付款",
        ("pending", "sale"): "待收款",
        ("other", "income"): "其他收入",
        ("other", "expense"): "其他支出",
        ("party", "supplier"): "供应商",
        ("party", "customer"): "客户",
    }
    module = params.get("_module") or ""
    return names.get((module, direction)) or fallback


def _body(resp):
    raw = resp.body if hasattr(resp, "body") else resp
    if isinstance(raw, (bytes, bytearray)):
        payload = json.loads(raw)
    elif isinstance(raw, dict):
        payload = raw
    else:
        payload = json.loads(bytes(raw))
    return payload.get("data")


def _rows(data, params):
    if isinstance(data, list):
        return data
    if not isinstance(data, dict):
        return []
    if params.get("_module") == "supplier-detail" and params.get("scope") == "contracts":
        return data.get("contracts") or []
    if "items" in data:
        return data.get("items") or []
    if "suppliers" in data:
        return data.get("suppliers") or []
    return []


def row_key(row: dict) -> str:
    if row.get("id") not in (None, ""):
        return str(row["id"])
    return f"{row.get('order_id')}-{row.get('material_code')}"


def _filter_ids(rows, ids):
    if not ids:
        return rows
    if isinstance(ids, str):
        ids = [part for part in ids.split(",") if part]
    wanted = {str(item) for item in ids if str(item)}
    if not wanted:
        return rows
    return [row for row in rows if row_key(row) in wanted]


def _safe(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return "是" if value else "否"
    text = str(value)
    if text.startswith(_FORMULA):
        return "'" + text
    return text


def _columns(rows, requested):
    if requested:
        cols = []
        for item in requested:
            if isinstance(item, dict):
                key = item.get("key")
                if key:
                    cols.append({"key": key, "label": item.get("label") or key})
            elif item:
                cols.append({"key": str(item), "label": str(item)})
        return cols
    if not rows:
        return []
    return [{"key": key, "label": key} for key in rows[0].keys() if key not in _HIDE]


def build_xlsx(title: str, columns: list, rows: list) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = (title or "导出")[:31] or "导出"
    header_font = Font(name="Microsoft YaHei", bold=True, size=10, color="FFFFFF")
    body_font = Font(name="Microsoft YaHei", size=10)
    fill = PatternFill("solid", fgColor="1677FF")
    align = Alignment(vertical="center")
    sheet.append([col["label"] for col in columns])
    for cell in sheet[1]:
        cell.font = header_font
        cell.fill = fill
        cell.alignment = align
    for row in rows:
        values = []
        for col in columns:
            values.append(_safe(row.get(col["key"], "")))
        sheet.append(values)
        for cell in sheet[sheet.max_row]:
            cell.font = body_font
            cell.alignment = align
    for index, col in enumerate(columns, 1):
        width = max(8, min(36, len(col["label"]) * 2 + 4))
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def build_csv(columns: list, rows: list) -> bytes:
    buf = StringIO()
    writer = csv.writer(buf)
    writer.writerow([col["label"] for col in columns])
    for row in rows:
        writer.writerow([_safe(row.get(col["key"], "")) for col in columns])
    return ("\ufeff" + buf.getvalue()).encode("utf-8")


def _dict_list(path: str):
    from ..routers.misc_api import router

    for route in router.routes:
        methods = getattr(route, "methods", set()) or set()
        if "GET" not in methods:
            continue
        current = getattr(route, "path", "") or ""
        if current == path or current.endswith(path):
            return route.endpoint
    raise BizError("未知导出模块")


def collect(module: str, params: dict, user, db):
    name = (module or "").strip()
    extra = {
        "roles": ("/roles", "角色"),
        "depts": ("/depts", "部门"),
        "posts": ("/posts", "岗位"),
        "configs": (None, "安全设置"),
    }
    if name == "configs":
        from ..routers import misc_api
        resp = misc_api.list_configs(user=user, db=db)
        data = _body(resp)
        rows = data if isinstance(data, list) else []
        return rows, extra[name][1]
    if name in extra:
        fn = _dict_list(extra[name][0])
        resp = fn(QueryRequest(params), user=user, db=db)
        return _rows(_body(resp), params), extra[name][1]
    table = {
        "contract": ("list_contracts", "合同", "biz"),
        "contracts": ("list_contracts", "合同", "biz"),
        "order": ("list_orders", "订单", "biz"),
        "orders": ("list_orders", "订单", "biz"),
        "invoice": ("list_invoices", "发票", "biz"),
        "invoices": ("list_invoices", "发票", "biz"),
        "settle": ("list_settle", "收付款", "biz"),
        "settle-bills": ("list_settle", "收付款", "biz"),
        "wms": ("list_wms", "仓储", "biz"),
        "wms-bills": ("list_wms", "仓储", "biz"),
        "party": ("list_parties", "客商", "biz"),
        "parties": ("list_parties", "客商", "biz"),
        "project": ("list_projects", "项目", "misc"),
        "projects": ("list_projects", "项目", "misc"),
        "purchase-plan": ("list_plans", "采购计划", "misc"),
        "purchase-plans": ("list_plans", "采购计划", "misc"),
        "pending": ("pending", "待收付", "misc"),
        "goods": ("list_goods", "商品", "misc"),
        "stock": ("list_stocks", "库存余额", "misc"),
        "stocks": ("list_stocks", "库存余额", "misc"),
        "transfer": ("list_transfer", "调拨记录", "misc"),
        "transfer-bills": ("list_transfer", "调拨记录", "misc"),
        "check": ("list_checks", "盘点记录", "misc"),
        "check-bills": ("list_checks", "盘点记录", "misc"),
        "other": ("list_other", "其他收支", "misc"),
        "other-bills": ("list_other", "其他收支", "misc"),
        "users": ("list_users", "用户", "misc"),
        "login-logs": ("login_logs", "登录日志", "misc"),
        "op-logs": ("op_logs", "操作日志", "misc"),
        "reminders": ("reminders", "到期提醒", "misc"),
        "stats-customer": ("stats_customer", "客户统计", "misc"),
        "stats-supplier": ("stats_supplier", "供应商统计", "misc"),
        "supplier-detail": ("supplier_detail", "供应商往来统计", "misc"),
        "fund-accounts": ("list_accounts", "资金账户", "misc"),
        "contract-types": ("list_types", "合同类型", "misc"),
        "warehouses": ("list_warehouses", "仓库", "misc"),
        "goods-categories": ("list_cats", "商品分类", "misc"),
    }
    customs_table = {
        "warehouse-receipt": ("list_receipts", "收货入仓"),
        "customs-material": ("list_materials", "报关物料清单"),
        "customs-invoice": ("list_invoices", "报关发票"),
        "customs-clearance": ("list_clearances", "清关发票"),
        "customs-element": ("list_elements", "报关要素"),
        "hs-code": ("list_hs", "HS Code库"),
        "exchange-rate": ("list_rates", "汇率库"),
    }
    if name in customs_table:
        from . import customs as customs_svc
        fn_name, title = customs_table[name]
        rows = getattr(customs_svc, fn_name)(db, user.session_company_id, params)
        return rows, title
    spec = table.get(name)
    if spec is None:
        raise BizError("未知导出模块")
    from ..routers import biz_api, misc_api
    pack = biz_api if spec[2] == "biz" else misc_api
    fn = getattr(pack, spec[0])
    query = dict(params)
    if name in ("wms", "wms-bills") and not query.get("type"):
        query["type"] = "inbound"
    resp = fn(QueryRequest(query), user=user, db=db)
    data = _body(resp)
    return _rows(data, {**params, "_module": name}), spec[1]


def export_file(module: str, params: dict, user, db) -> tuple[bytes, str, str]:
    query = dict(params or {})
    query["_module"] = module
    rows, fallback = collect(module, query, user, db)
    rows = _filter_ids(rows, query.get("ids") or [])
    raw_cols = query.get("columns") or []
    if isinstance(raw_cols, str):
        try:
            raw_cols = json.loads(raw_cols)
        except json.JSONDecodeError:
            raw_cols = []
    columns = _columns(rows, raw_cols)
    if not columns:
        raise BizError("没有可导出的列")
    title = _named_filename(query, fallback)
    fmt = (query.get("format") or "xlsx").lower()
    if fmt == "csv":
        return build_csv(columns, rows), f"{title}.csv", "text/csv; charset=utf-8"
    return build_xlsx(title, columns, rows), f"{title}.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
