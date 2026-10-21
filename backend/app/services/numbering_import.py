"""合同编号、发票编号的 Excel 批量导入。

已有编号跳过，不覆盖。符合现行规则的编号占用顺序号；旧号原样登记。
以 -001 结尾的发票号归入同一批次。
"""
import hashlib
import re
from datetime import date, datetime, timedelta
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from ..errors import BizError
from ..models import (
    Contract,
    CustomsInvoice,
    NumberBatch,
    NumberBatchSlot,
    NumberContract,
    NumberInvoice,
    now,
)
from ..money import round2
from .numbering import (
    format_sub_no,
    matches_any_pattern,
    occupy_slot,
    parse_pool_number,
    party_a_name,
    rule_map,
    slot_blocked,
)
from .util import q, stamp_new

PLACEHOLDER = "（历史未填）"
CONTRACT_CODE_RE = re.compile(r"^[A-Za-z0-9][A-Za-z0-9\-]{3,63}$")
CONTRACT_EXTRACT_RE = re.compile(r"[A-Za-z]{2,}(?:-[A-Za-z0-9]+){1,8}")
INVOICE_CODE_RE = re.compile(r"^[A-Za-z0-9\u4e00-\u9fff（）()\-]{4,80}$")
YEAR_IN_NUMBER_RE = re.compile(r"(20\d{2})")
SHORT_YEAR_RE = re.compile(r"-(\d{2})(\d{3,})[A-Za-z]?(?:-\d{3})?$")
SUB_RE = re.compile(r"^(.+)-(\d{3})$")

CONTRACT_FIELDS = {
    "contract_no": ("合同编号", "合同号"),
    "business_type": ("业务类型",),
    "number_year": ("编号年份",),
    "product_name": ("产品/合同名称", "产品名称", "合同名称"),
    "party_b": ("乙方", "供应商名称", "供应商"),
    "party_a": ("甲方",),
    "amount": ("合同金额（CNY）", "合同金额", "合同总额"),
    "sign_date": ("签订时间", "签订日期", "合同日期"),
    "owner": ("归属人", "业务员"),
}
INVOICE_FIELDS = {
    "invoice_no": ("发票号", "发票编号"),
    "contract_no": ("关联合同", "合同编号", "合同号"),
    "business_type": ("业务类型",),
    "number_rule": ("编号规则",),
    "number_year": ("编号年份",),
    "content": ("所含内容", "内容"),
    "bl_date": ("提单日期",),
    "price_term": ("价格条款",),
    "currency": ("结算币种", "币种"),
    "amount_usd": ("美元金额", "USD金额"),
    "amount_cny": ("人民币金额", "CNY金额"),
    "rate": ("汇率",),
    "exchange_month": ("汇率月份",),
    "price_combined": ("发票价格", "发票金额"),
    "po_no": ("PO号", "PO"),
    "operator_name": ("登记人", "操作人"),
    "transport_type": ("运输方式",),
    "email_sent_date": ("发邮件时间", "邮件时间"),
    "remark": ("备注",),
    "bl_no": ("提单号",),
}


def template_bytes(kind: str) -> tuple[bytes, str]:
    book = openpyxl.Workbook()
    sheet = book.active
    if kind == "contract":
        sheet.title = "合同"
        headers = ["合同编号", "业务类型", "编号年份", "产品/合同名称", "乙方", "甲方", "合同金额（CNY）", "签订时间", "归属人"]
        notes = [
            "请在“合同”工作表从第 2 行填写，不要修改第 1 行表头。合同编号必填，按原样保存。",
            "业务类型填锰矿或铝矿。编号年份可空，空着时从编号里识别。签订时间格式 YYYY-MM-DD，可空。归属人填员工编号或姓名，可空。",
            "符合 BMI-GMI-YYYYXXX 或 BMGG-MAZ-YYXXX 的编号会占用顺序号；旧格式编号只登记。",
            "同一编号已存在则跳过，不覆盖。也可以直接上传表头含“合同号”的历史统计表。",
        ]
        filename = "合同编号导入模板.xlsx"
    else:
        sheet.title = "发票"
        headers = [
            "发票号", "关联合同", "业务类型", "编号规则", "编号年份", "所含内容", "提单日期",
            "价格条款", "结算币种", "美元金额", "人民币金额", "汇率", "汇率月份", "PO",
            "登记人", "运输方式", "发邮件时间", "备注", "提单号",
        ]
        notes = [
            "请在“发票”工作表从第 2 行填写，不要修改第 1 行表头。发票号必填，按原样保存。",
            "编号规则可填铝矿海运、铝矿空运或锰矿。运输方式填海运或空运。价格条款填 CIF 或 FOB。",
            "以 -001 这种三位分单号结尾时，归入该前缀的批次。已存在的发票号会跳过。",
            "也可以直接上传表头含“发票号”和“所含内容”的发票编号登记表。",
        ]
        filename = "发票编号导入模板.xlsx"
    sheet.append(headers)
    fill = PatternFill("solid", fgColor="E8F3FF")
    font = Font(bold=True, color="1D2129")
    for cell in sheet[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center")
    for column in sheet.columns:
        letter = column[0].column_letter
        sheet.column_dimensions[letter].width = min(max(14, len(str(column[0].value or "")) + 4), 36)
    sheet.freeze_panes = "A2"
    info = book.create_sheet("填写说明")
    for line in notes:
        info.append([line])
    info.column_dimensions["A"].width = 92
    buf = BytesIO()
    book.save(buf)
    return buf.getvalue(), filename


def parse_number_file(db, company_id, module, content: bytes, filename: str):
    kind = "contract" if module == "number-contract" else "invoice"
    sheets = _load_sheets(content, filename)
    rules = rule_map(db, company_id)
    if kind == "contract":
        items, used = _collect_contracts(sheets, rules)
        empty = "没有找到可导入的合同。请使用模板，或上传表头含「合同号」的统计表。"
    else:
        items, used = _collect_invoices(sheets, rules)
        empty = "没有找到可导入的发票。请使用模板，或上传表头含「发票号」的登记表。"
    ignored = [name for name, _grid in sheets if name not in used and name != "填写说明"]
    if not items:
        raise BizError(empty)
    rows = _preview_rows(db, company_id, kind, items)
    if not rows:
        raise BizError(empty)
    note = ""
    if ignored:
        note = "未识别表：" + "、".join(ignored[:8])
    return {
        "file_hash": hashlib.sha256(content).hexdigest(),
        "mode": "create",
        "rows": rows,
        "locks": [],
        "batch_error": note,
    }


def execute_number_file(db, user, company_id, module, rows):
    kind = "contract" if module == "number-contract" else "invoice"
    ready = []
    fail = 0
    errors = []
    for row in rows or []:
        if not row.get("ok"):
            fail += 1
            if row.get("message"):
                errors.append({"row_no": row.get("row_no"), "message": row.get("message")})
            continue
        data = row.get("data")
        if not isinstance(data, dict):
            fail += 1
            errors.append({"row_no": row.get("row_no"), "message": "导入数据不完整，请重新上传"})
            continue
        ready.append((row.get("row_no"), data))
    success = 0
    if kind == "contract":
        success, more_fail, more_errors = _write_contracts(db, user, company_id, ready)
    else:
        success, more_fail, more_errors = _write_invoices(db, user, company_id, ready)
    fail += more_fail
    errors.extend(more_errors)
    return {"success": success, "fail": fail, "errors": errors}


def _write_contracts(db, user, company_id, ready):
    success = 0
    fail = 0
    errors = []
    seen = set()
    for row_no, item in ready:
        number = item["contract_no"]
        if number in seen or _contract_taken(db, company_id, number):
            fail += 1
            errors.append({"row_no": row_no, "message": "编号已存在，已跳过"})
            continue
        try:
            with db.begin_nested():
                _insert_contract(db, user, company_id, item)
            seen.add(number)
            success += 1
        except BizError as exc:
            fail += 1
            errors.append({"row_no": row_no, "message": exc.message})
    return success, fail, errors


def _import_owner(db, company_id, item, fallback=""):
    text = (item.get("owner_id") or item.get("owner_text") or "").strip()
    if not text:
        return (fallback or "")[:64]
    from .numbering import resolve_owner
    owner_id, _name = resolve_owner(db, company_id, text)
    return owner_id


def _insert_contract(db, user, company_id, item):
    number = item["contract_no"]
    purchase = (
        db.query(Contract)
        .filter(
            Contract.company_id == company_id,
            Contract.contract_code == number,
            Contract.deleted_at.is_(None),
            Contract.direction == "purchase",
            Contract.nature == "main",
        )
        .first()
    )
    if purchase and q(db, NumberContract, company_id).filter(NumberContract.purchase_contract_id == purchase.id).first():
        raise BizError("该采购合同已有合同编号")
    pool = item.get("pool")
    seq = item.get("seq")
    if pool and seq:
        if slot_blocked(db, company_id, pool, seq):
            raise BizError(f"顺序号 {seq} 已被占用")
        occupy_slot(db, company_id, pool, seq, item["number_year"], number)
    row = NumberContract(
        contract_no=number,
        purchase_contract_id=purchase.id if purchase else None,
        business_type=item["business_type"],
        number_year=int(item["number_year"]),
        num_pool=pool,
        num_seq=seq,
        product_name=item["product_name"],
        party_b=item["party_b"],
        party_a=item["party_a"] or party_a_name(db, company_id),
        amount_cny=item["amount_cny"],
        sign_date=_as_date(item.get("sign_date")),
        owner_id=_import_owner(db, company_id, item, purchase.owner_id if purchase else ""),
    )
    stamp_new(row, user, company_id)
    db.add(row)
    db.flush()


def _write_invoices(db, user, company_id, ready):
    success = 0
    fail = 0
    errors = []
    seen = set()
    batches = {}
    ordered = [pair for pair in ready if not pair[1].get("is_sub")] + [pair for pair in ready if pair[1].get("is_sub")]
    for row_no, item in ordered:
        number = item["invoice_no"]
        prefix = item.get("prefix") or ""
        if number in seen or _invoice_taken(db, company_id, number, allow_prefix=prefix):
            fail += 1
            errors.append({"row_no": row_no, "message": "编号已存在，已跳过"})
            continue
        try:
            with db.begin_nested():
                batch = _ensure_batch(db, user, company_id, item, batches) if item.get("is_sub") or item.get("is_master") else None
                _insert_invoice(db, user, company_id, item, batch)
            seen.add(number)
            success += 1
        except BizError as exc:
            fail += 1
            errors.append({"row_no": row_no, "message": exc.message})
    for batch in batches.values():
        _refresh_batch(db, company_id, batch.id)
    return success, fail, errors


def _ensure_batch(db, user, company_id, item, batches):
    prefix = item.get("prefix") or item["invoice_no"]
    if prefix in batches:
        return batches[prefix]
    existing = q(db, NumberBatch, company_id).filter(NumberBatch.prefix == prefix).first()
    if existing:
        batches[prefix] = existing
        return existing
    if _contract_taken(db, company_id, prefix) or _invoice_taken(db, company_id, prefix):
        raise BizError(f"批次前缀 {prefix} 已被占用")
    batch = NumberBatch(
        prefix=prefix,
        number_rule=item["number_rule"],
        number_year=int(item["number_year"]),
        transport_type=item.get("transport_type") or "",
        business_type=item.get("business_type") or "",
    )
    stamp_new(batch, user, company_id)
    db.add(batch)
    db.flush()
    batches[prefix] = batch
    return batch


def _insert_invoice(db, user, company_id, item, batch):
    number = item["invoice_no"]
    pool = item.get("pool")
    seq = item.get("seq")
    if pool and seq:
        if slot_blocked(db, company_id, pool, seq):
            raise BizError(f"顺序号 {seq} 已被占用")
        occupy_slot(db, company_id, pool, seq, item["number_year"], number)
        if batch and item.get("is_master") and not batch.num_pool:
            batch.num_pool = pool
            batch.num_seq = seq
    if item.get("is_sub"):
        _occupy_sub(db, company_id, batch.id, item["sub_seq"], number)
    customs = (
        db.query(CustomsInvoice)
        .filter(CustomsInvoice.company_id == company_id, CustomsInvoice.invoice_no == number, CustomsInvoice.deleted_at.is_(None))
        .first()
    )
    if customs and q(db, NumberInvoice, company_id).filter(NumberInvoice.customs_invoice_id == customs.id).first():
        customs = None
    linked = q(db, NumberContract, company_id).filter(NumberContract.contract_no == (item.get("contract_no") or "")).first() if item.get("contract_no") else None
    row = NumberInvoice(
        invoice_no=number,
        number_contract_id=linked.id if linked else None,
        contract_no=(item.get("contract_no") or "")[:64],
        customs_invoice_id=customs.id if customs else None,
        business_type=item.get("business_type") or "",
        number_rule=item.get("number_rule") or "",
        number_year=int(item["number_year"]),
        num_pool=pool,
        num_seq=seq,
        content=item.get("content") or "",
        bl_date=_as_date(item.get("bl_date")),
        price_term=item.get("price_term") or "FOB",
        currency=item.get("currency") or "CNY",
        amount_usd=item.get("amount_usd") or "0.00",
        amount_cny=item.get("amount_cny") or "0.00",
        exchange_rate=item.get("exchange_rate") or "0",
        exchange_month=item.get("exchange_month") or "",
        rate_adjusted=bool(item.get("rate_adjusted")),
        po_no=(item.get("po_no") or "")[:100],
        operator_name=(item.get("operator_name") or getattr(user, "name", None) or user.employee_id or "")[:100],
        transport_type=item.get("transport_type") or "",
        email_sent_date=_as_date(item.get("email_sent_date")),
        remark=item.get("remark") or "",
        bl_no=(item.get("bl_no") or "")[:100],
        is_batch=bool(item.get("is_sub") or item.get("is_master")),
        is_batch_master=bool(item.get("is_master")),
        batch_id=batch.id if batch else None,
        batch_prefix=batch.prefix if batch else "",
        batch_sub_no=format_sub_no(item["sub_seq"]) if item.get("is_sub") else "",
        sub_seq=item.get("sub_seq"),
    )
    stamp_new(row, user, company_id)
    db.add(row)
    db.flush()


def _occupy_sub(db, company_id, batch_id, seq, number):
    seq = int(seq)
    have = {row.seq: row for row in db.query(NumberBatchSlot).filter(NumberBatchSlot.batch_id == batch_id).all()}
    if have.get(seq) and have[seq].status == "used":
        raise BizError("分单号已被占用")
    for index in range(1, seq):
        if index in have:
            continue
        slot = NumberBatchSlot(company_id=company_id, batch_id=batch_id, seq=index, status="free", number=None, updated_at=now())
        db.add(slot)
        have[index] = slot
    current = have.get(seq)
    if current:
        current.status = "used"
        current.number = number
        current.updated_at = now()
    else:
        db.add(NumberBatchSlot(company_id=company_id, batch_id=batch_id, seq=seq, status="used", number=number, updated_at=now()))
    db.flush()


def _refresh_batch(db, company_id, batch_id):
    children = (
        q(db, NumberInvoice, company_id)
        .filter(NumberInvoice.batch_id == batch_id, NumberInvoice.is_batch_master.is_(False))
        .all()
    )
    if not children:
        return
    master = (
        q(db, NumberInvoice, company_id)
        .filter(NumberInvoice.batch_id == batch_id, NumberInvoice.is_batch_master.is_(True))
        .first()
    )
    if master is None:
        return
    master.amount_usd = round2(sum((Decimal(str(row.amount_usd or 0)) for row in children), Decimal("0")))
    master.amount_cny = round2(sum((Decimal(str(row.amount_cny or 0)) for row in children), Decimal("0")))


def _contract_taken(db, company_id, number):
    return q(db, NumberContract, company_id).filter(NumberContract.contract_no == number).first() is not None


def _invoice_taken(db, company_id, number, allow_prefix=""):
    if q(db, NumberInvoice, company_id).filter(NumberInvoice.invoice_no == number).first():
        return True
    if _contract_taken(db, company_id, number):
        return True
    batch = q(db, NumberBatch, company_id).filter(NumberBatch.prefix == number).first()
    if batch and number != allow_prefix:
        return True
    return False


def _preview_rows(db, company_id, kind, items):
    rows = []
    seen = set()
    for item in items:
        if item.get("_error"):
            rows.append({"row_no": item["row_no"], "ok": False, "summary": item.get("summary") or "", "message": item["_error"]})
            continue
        number = item["contract_no"] if kind == "contract" else item["invoice_no"]
        notes = []
        ok = True
        if number in seen:
            ok = False
            notes.append("文件中编号重复")
        elif kind == "contract" and _contract_taken(db, company_id, number):
            ok = False
            notes.append("编号已存在，已跳过")
        elif kind == "invoice" and _invoice_taken(db, company_id, number, allow_prefix=item.get("prefix") or ""):
            ok = False
            notes.append("编号已存在，已跳过")
        else:
            if kind == "contract" and not item.get("party_a"):
                item["party_a"] = party_a_name(db, company_id)
            if item.get("pool"):
                notes.append("将占用现行顺序号")
            else:
                notes.append("历史编号，不占新号池")
            if item.get("filled_date"):
                notes.append(f"未填签订时间，按 {item.get('sign_date')} 保存")
            if item.get("zero_amount"):
                notes.append("未填金额，按 0 保存")
            if item.get("is_sub"):
                notes.append(f"归入批次 {item.get('prefix')}")
            if item.get("is_master"):
                notes.append("批次前缀")
            if kind == "contract" and item.get("owner_text"):
                try:
                    from .numbering import resolve_owner
                    owner_id, owner_name = resolve_owner(db, company_id, item["owner_text"])
                    item["owner_id"] = owner_id
                    notes.append(f"归属人 {owner_name}")
                except BizError as exc:
                    ok = False
                    notes = [exc.message]
        seen.add(number)
        if kind == "contract":
            summary = f"{number} {item.get('business_type') or ''} {item.get('product_name') or ''}".strip()
        else:
            summary = f"{number} {item.get('number_rule') or ''} {item.get('content') or ''}".strip()
        rows.append({
            "row_no": item["row_no"],
            "ok": ok,
            "summary": summary,
            "message": "；".join(notes),
            "data": item if ok else None,
        })
    return rows


def _collect_contracts(sheets, rules=None):
    items = []
    used = set()
    for name, grid in sheets:
        found = _find_header(grid, CONTRACT_FIELDS, "contract_no", ("product_name", "party_b", "amount"))
        if not found:
            continue
        used.add(name)
        header_idx, cols = found
        for excel_row, raw in _data_rows(grid, header_idx):
            item, error = _parse_contract_row(raw, cols, rules)
            if error == "blank":
                continue
            where = f"{name} 第{excel_row}行"
            if error:
                items.append({"row_no": excel_row, "where": where, "summary": where, "_error": error})
                continue
            item["row_no"] = excel_row
            item["where"] = where
            items.append(item)
    return items, used


def _collect_invoices(sheets, rules=None):
    items = []
    used = set()
    for name, grid in sheets:
        found = _find_header(grid, INVOICE_FIELDS, "invoice_no", ("content", "price_combined", "amount_cny", "amount_usd", "bl_date", "transport_type"))
        if not found:
            continue
        used.add(name)
        header_idx, cols = found
        for excel_row, raw in _data_rows(grid, header_idx):
            pieces = _split_invoice_numbers(_cell(raw, cols.get("invoice_no")))
            targets = [(raw, excel_row, f"{name} 第{excel_row}行")]
            if len(pieces) > 1:
                targets = []
                for index, number in enumerate(pieces, start=1):
                    copied = list(raw)
                    copied[cols["invoice_no"]] = number
                    targets.append((tuple(copied), excel_row, f"{name} 第{excel_row}行第{index}个编号"))
            for target, row_no, where in targets:
                item, error = _parse_invoice_row(target, cols, rules)
                if error == "blank":
                    continue
                if error:
                    items.append({"row_no": row_no, "where": where, "summary": where, "_error": error})
                    continue
                item["row_no"] = row_no
                item["where"] = where
                items.append(item)
    prefixes = {item["prefix"] for item in items if item.get("is_sub")}
    for item in items:
        if item.get("_error"):
            continue
        if not item.get("is_sub") and item.get("invoice_no") in prefixes:
            item["is_master"] = True
            item["prefix"] = item["invoice_no"]
    return items, used


def _parse_contract_row(raw, cols, rules=None):
    number_text = _text(_cell(raw, cols.get("contract_no"))).replace(" ", "")
    if number_text in ("", "合同号", "合同编号", "合计", "小计", "总计", "备注"):
        return None, "blank"
    number = _normalize_contract_no(number_text)
    if not number:
        return None, "合同号无法识别"
    pool, parsed = _match_pool(number, ("contract_mn", "contract_al"), rules)
    if pool is None and matches_any_pattern(number):
        return None, "该编号属于发票规则，不能作为合同导入"
    if pool == "contract_mn":
        business, year, seq = "锰矿", parsed[0], parsed[1]
    elif pool == "contract_al":
        business, year, seq = "铝矿", parsed[0], parsed[1]
    else:
        business = _norm_business(_cell(raw, cols.get("business_type"))) or _infer_contract_business(number)
        year = _norm_year(_cell(raw, cols.get("number_year"))) or _infer_year(number)
        seq = None
    product, _clipped = _clip(_text(_cell(raw, cols.get("product_name"))), 200)
    party_b, _clipped = _clip(_text(_cell(raw, cols.get("party_b"))), 200)
    party_a, _clipped = _clip(_text(_cell(raw, cols.get("party_a"))), 200)
    if not product:
        product = PLACEHOLDER
    if not party_b:
        party_b = PLACEHOLDER
    amount_raw = _cell(raw, cols.get("amount"))
    zero_amount = _is_empty(amount_raw)
    try:
        amount = "0.00" if zero_amount else _parse_money(amount_raw, "合同金额")
    except ValueError as exc:
        return None, str(exc)
    sign_raw = _cell(raw, cols.get("sign_date"))
    filled_date = _is_empty(sign_raw)
    try:
        sign_date = f"{year}-01-01" if filled_date else _parse_any_date(sign_raw, "签订时间")
    except ValueError as exc:
        return None, str(exc)
    return {
        "contract_no": number,
        "business_type": business,
        "number_year": year,
        "pool": pool,
        "seq": seq,
        "product_name": product,
        "party_b": party_b,
        "party_a": party_a,
        "amount_cny": amount,
        "sign_date": sign_date,
        "filled_date": filled_date,
        "zero_amount": zero_amount,
        "owner_text": _clip(_text(_cell(raw, cols.get("owner"))), 100)[0],
    }, None


def _parse_invoice_row(raw, cols, rules=None):
    number_text = _text(_cell(raw, cols.get("invoice_no"))).replace(" ", "")
    if number_text in ("", "发票号", "发票编号"):
        return None, "blank"
    if not INVOICE_CODE_RE.match(number_text) or not re.search(r"\d", number_text):
        return None, "发票号无法识别"
    sub = SUB_RE.match(number_text)
    prefix = None
    sub_seq = None
    if sub and len(sub.group(1)) >= 4:
        prefix = sub.group(1)
        sub_seq = int(sub.group(2))
        if sub_seq < 1:
            return None, "分单号不能为 0"
    pool, parsed = (None, None)
    if sub is None:
        pool, parsed = _match_pool(number_text, ("invoice_al_sea", "invoice_al_air", "invoice_mn"), rules)
        if pool is None and matches_any_pattern(number_text):
            return None, "该编号属于合同规则，不能作为发票导入"
    transport, transport_note = _infer_transport(number_text, _cell(raw, cols.get("transport_type")))
    explicit_rule = _norm_rule(_cell(raw, cols.get("number_rule")))
    explicit_business = _norm_business(_cell(raw, cols.get("business_type")))
    if pool:
        rule = {"invoice_al_sea": "铝矿海运", "invoice_al_air": "铝矿空运", "invoice_mn": "锰矿"}[pool]
        business = "锰矿" if rule == "锰矿" else "铝矿"
        year, seq = parsed
        transport = {"铝矿空运": "空运", "铝矿海运": "海运"}.get(rule, transport)
    else:
        business, rule = _legacy_invoice_rule(number_text, transport, explicit_business, explicit_rule)
        year = _norm_year(_cell(raw, cols.get("number_year"))) or _infer_year(number_text)
        seq = None
    content, _clipped = _clip(_text(_cell(raw, cols.get("content"))) or PLACEHOLDER, 500)
    try:
        amounts = _invoice_amounts(raw, cols)
    except ValueError as exc:
        return None, str(exc)
    bl_date, bl_note = _loose_date(_cell(raw, cols.get("bl_date")))
    email_date, email_note = _loose_date(_cell(raw, cols.get("email_sent_date")))
    month = amounts["exchange_month"]
    if not month:
        month = (bl_date or email_date or f"{year}-01-01")[:7]
    remark, _clipped = _clip(_text(_cell(raw, cols.get("remark"))), 500)
    notes = []
    if transport_note and transport_note not in ("海运", "空运"):
        notes.append(f"原运输方式：{transport_note}")
    if bl_note:
        notes.append(f"提单日期：{bl_note}")
    if email_note:
        notes.append(f"发邮件时间：{email_note}")
    if notes:
        extra = "；".join(notes)
        remark = extra if not remark else f"{remark}；{extra}"
        remark, _clipped = _clip(remark, 500)
    po, _clipped = _clip(_text(_cell(raw, cols.get("po_no"))), 100)
    operator, _clipped = _clip(_text(_cell(raw, cols.get("operator_name"))), 100)
    bl_no, _clipped = _clip(_text(_cell(raw, cols.get("bl_no"))), 100)
    price_term = _text(_cell(raw, cols.get("price_term"))).upper()
    if price_term not in ("CIF", "FOB"):
        price_term = amounts["term"] or "FOB"
    return {
        "invoice_no": number_text,
        "is_sub": prefix is not None,
        "is_master": False,
        "prefix": prefix,
        "sub_seq": sub_seq,
        "business_type": business,
        "number_rule": rule,
        "number_year": year,
        "pool": pool,
        "seq": seq,
        "content": content,
        "bl_date": bl_date,
        "price_term": price_term,
        "currency": amounts["currency"],
        "amount_usd": amounts["usd"],
        "amount_cny": amounts["cny"],
        "exchange_rate": amounts["rate"],
        "exchange_month": month,
        "rate_adjusted": amounts["adjusted"],
        "po_no": po,
        "operator_name": operator,
        "transport_type": transport,
        "email_sent_date": email_date,
        "remark": remark,
        "bl_no": bl_no,
        "contract_no": _normalize_contract_no(_text(_cell(raw, cols.get("contract_no")))),
        "zero_amount": amounts["zero"],
        "filled_date": False,
    }, None


def _load_sheets(content: bytes, filename: str):
    name = (filename or "").lower()
    if not name.endswith((".xlsx", ".xls")):
        raise BizError("文件格式不正确")
    if len(content) > 10 * 1024 * 1024:
        raise BizError("文件大小超过限制")
    if name.endswith(".xlsx"):
        if not content.startswith(b"PK"):
            raise BizError("文件格式不正确")
        try:
            book = openpyxl.load_workbook(BytesIO(content), data_only=False, read_only=True)
        except Exception as exc:
            raise BizError("无法读取 Excel，请确认文件是 xlsx 且未加密") from exc
        sheets = []
        try:
            for ws in book.worksheets:
                grid = []
                for index, row in enumerate(ws.iter_rows(values_only=True)):
                    if index >= 20000:
                        break
                    grid.append(tuple(row))
                sheets.append((ws.title, grid))
        finally:
            book.close()
        return sheets
    if not content.startswith(b"\xd0\xcf\x11\xe0"):
        raise BizError("文件格式不正确")
    try:
        import xlrd
        book = xlrd.open_workbook(file_contents=content)
    except Exception as exc:
        raise BizError("无法读取 Excel") from exc
    sheets = []
    for ws in book.sheets():
        grid = []
        for r in range(min(ws.nrows, 20000)):
            values = []
            for c in range(ws.ncols):
                cell = ws.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        values.append(xlrd.xldate_as_datetime(cell.value, book.datemode))
                    except Exception:
                        values.append(cell.value)
                else:
                    values.append("" if cell.ctype == xlrd.XL_CELL_ERROR else cell.value)
            grid.append(tuple(values))
        sheets.append((ws.name, grid))
    return sheets


def _find_header(grid, fields, key, required_any):
    for index, raw in enumerate(grid[:30]):
        headers = [_norm_header(item) for item in raw]
        if _col_index(headers, fields[key]) is None:
            continue
        if not any(_col_index(headers, fields[name]) is not None for name in required_any):
            continue
        return index, {name: _col_index(headers, aliases) for name, aliases in fields.items()}
    return None


def _data_rows(grid, header_idx):
    for offset, raw in enumerate(grid[header_idx + 1 :], start=header_idx + 2):
        yield offset, raw


def _match_pool(number, pools, rules=None):
    for pool in pools:
        for spec in (rules or {}).get(pool) or [None]:
            parsed = parse_pool_number(pool, number, spec)
            if parsed:
                return pool, parsed
    return None, None


def _infer_contract_business(number):
    upper = number.upper()
    if "MAZ" in upper or "BMGG" in upper:
        return "铝矿"
    return "锰矿"


def _legacy_invoice_rule(number, transport, explicit_business, explicit_rule):
    if explicit_rule in ("铝矿海运", "铝矿空运", "锰矿"):
        return ("锰矿" if explicit_rule == "锰矿" else "铝矿"), explicit_rule
    body = re.sub(r"-\d{3}$", "", number).upper()
    if "MAZ" not in body and "BMGG" not in body and "GMI" in body:
        return "锰矿", "锰矿"
    if transport == "空运" or body.endswith("A"):
        return "铝矿", "铝矿空运"
    if explicit_business == "锰矿":
        return "锰矿", "锰矿"
    return "铝矿", "铝矿海运"


def _infer_transport(number, raw):
    text = _text(raw)
    if text in ("海运", "空运"):
        return text, ""
    if "空运" in text:
        return "空运", text
    if "海运" in text:
        return "海运", text
    body = re.sub(r"-\d{3}$", "", number).upper()
    if body.endswith("A"):
        return "空运", text
    return "海运", text


def _infer_year(number):
    matched = YEAR_IN_NUMBER_RE.search(number or "")
    if matched:
        year = int(matched.group(1))
        if 2000 <= year <= 2099:
            return year
    matched = SHORT_YEAR_RE.search(number or "")
    if matched:
        return 2000 + int(matched.group(1))
    return date.today().year


def _norm_year(value):
    text = _text(value)
    if re.fullmatch(r"20\d{2}", text):
        return int(text)
    return None


def _norm_business(value):
    text = _text(value)
    if "铝" in text:
        return "铝矿"
    if "锰" in text:
        return "锰矿"
    return ""


def _norm_rule(value):
    text = _text(value)
    if text in ("铝矿海运", "铝矿空运", "锰矿"):
        return text
    if "空运" in text:
        return "铝矿空运"
    if "海运" in text:
        return "铝矿海运"
    return ""


def _normalize_contract_no(raw):
    text = "" if raw is None else str(raw).strip().replace(" ", "").replace("\u3000", "")
    text = text.replace("－", "-").replace("—", "-").replace("–", "-")
    if CONTRACT_CODE_RE.match(text) and re.search(r"\d", text):
        return text
    found = CONTRACT_EXTRACT_RE.search(text)
    if found and re.search(r"\d", found.group(0)):
        return found.group(0)[:64]
    return ""


def _invoice_amounts(raw, cols):
    usd = ""
    cny = ""
    term = ""
    usd_cell = _cell(raw, cols.get("amount_usd"))
    cny_cell = _cell(raw, cols.get("amount_cny"))
    if not _is_empty(usd_cell):
        usd = _parse_money(usd_cell, "美元金额")
    if not _is_empty(cny_cell):
        cny = _parse_money(cny_cell, "人民币金额")
    combined = _cell(raw, cols.get("price_combined"))
    if not usd and not cny and not _is_empty(combined):
        usd, cny, term = _parse_combined_price(combined)
    rate_cell = _cell(raw, cols.get("rate"))
    adjusted = False
    if not _is_empty(rate_cell):
        rate = _parse_rate(rate_cell)
        adjusted = True
    elif usd and cny and Decimal(usd) > 0 and Decimal(cny) > 0:
        rate = format((Decimal(cny) / Decimal(usd)).quantize(Decimal("0.000001"), rounding=ROUND_HALF_UP), "f")
        adjusted = True
    else:
        rate = "0"
    month_text = _text(_cell(raw, cols.get("exchange_month")))
    month = month_text if re.fullmatch(r"20\d{2}-\d{2}", month_text) else ""
    currency_cell = _text(_cell(raw, cols.get("currency"))).upper()
    if currency_cell in ("CNY", "USD", "人民币", "美元"):
        currency = "USD" if currency_cell in ("USD", "美元") else "CNY"
    elif usd and Decimal(usd) > 0 and (not cny or Decimal(cny) == 0):
        currency = "USD"
    else:
        currency = "CNY"
    usd = usd or "0.00"
    cny = cny or "0.00"
    return {
        "usd": usd,
        "cny": cny,
        "rate": rate,
        "exchange_month": month,
        "currency": currency,
        "term": term,
        "adjusted": adjusted,
        "zero": Decimal(usd) == 0 and Decimal(cny) == 0,
    }


def _parse_combined_price(value):
    if isinstance(value, bool):
        raise ValueError("金额必须是数字")
    if isinstance(value, (int, float)):
        return "", _parse_money(value, "金额"), ""
    text = _text(value)
    if not text:
        return "", "", ""
    upper = text.upper()
    term = "FOB" if "FOB" in upper else "CIF" if "CIF" in upper else ""
    usd_match = re.search(r"(?:USD|美元)\s*[:：]?\s*(\d[\d,，]*\.?\d*)", text, re.I)
    cny_match = re.search(r"(?:CNY|RMB|人民币|￥|¥)\s*[:：]?\s*(\d[\d,，]*\.?\d*)", text, re.I)
    usd = _parse_money(usd_match.group(1), "美元金额") if usd_match else ""
    cny = _parse_money(cny_match.group(1), "人民币金额") if cny_match else ""
    if not usd and not cny:
        cny = _parse_money(text, "金额")
    return usd, cny, term


def _parse_money(value, name):
    if isinstance(value, bool):
        raise ValueError(f"{name}必须是数字")
    if isinstance(value, (int, float, Decimal)):
        return format(round2(value), "f")
    text = _text(value).replace(",", "").replace("，", "").replace("￥", "").replace("¥", "").replace("元", "").replace(" ", "")
    matched = re.search(r"-?\d+(?:\.\d+)?", text)
    if not matched:
        raise ValueError(f"{name}必须是数字")
    return format(round2(matched.group(0)), "f")


def _parse_rate(value):
    text = _text(value).replace(",", "").replace("，", "")
    matched = re.search(r"\d+(?:\.\d+)?", text)
    if not matched:
        raise ValueError("汇率必须是数字")
    number = Decimal(matched.group(0))
    if number < 0:
        raise ValueError("汇率不能为负数")
    return format(number.quantize(Decimal("0.000001"), rounding=ROUND_HALF_UP), "f")


def _parse_any_date(value, name):
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        serial = float(value)
        if 20000 <= serial <= 80000:
            return (datetime(1899, 12, 30) + timedelta(days=int(serial))).strftime("%Y-%m-%d")
        raise ValueError(f"{name}不是有效日期")
    text = _text(value).replace(".", "-").replace("/", "-").replace("年", "-").replace("月", "-").replace("日", "")
    matched = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})", text)
    if not matched:
        raise ValueError(f"{name}必须为日期")
    try:
        return date(int(matched.group(1)), int(matched.group(2)), int(matched.group(3))).strftime("%Y-%m-%d")
    except ValueError as exc:
        raise ValueError(f"{name}不是有效日期") from exc


def _loose_date(value):
    if _is_empty(value):
        return "", ""
    try:
        return _parse_any_date(value, "日期"), ""
    except ValueError:
        return "", _text(value)


def _split_invoice_numbers(value):
    text = _text(value)
    if not text:
        return []
    return [part.strip() for part in re.split(r"[\r\n]+", text) if part.strip()]


def _as_date(value):
    if not value:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return datetime.strptime(str(value)[:10], "%Y-%m-%d").date()


def _cell(raw, index):
    if index is None or index >= len(raw):
        return None
    return raw[index]


def _is_empty(value):
    return value is None or (isinstance(value, str) and not value.strip())


def _text(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int) and not isinstance(value, bool):
        return str(value)
    return str(value).replace("\xa0", " ").replace("\u3000", " ").strip().replace("－", "-").replace("—", "-").replace("–", "-")


def _norm_header(value):
    return re.sub(r"\s+", "", _text(value))


def _col_index(headers, aliases):
    names = sorted(aliases, key=len, reverse=True)
    for index, header in enumerate(headers):
        if header and header in names:
            return index
    for index, header in enumerate(headers):
        if not header:
            continue
        folded = header.upper()
        for name in names:
            if len(name) < 4:
                continue
            if header.startswith(name) or (name.isascii() and folded.startswith(name.upper())):
                return index
    return None


def _clip(value, limit):
    text = "" if value is None else str(value).strip()
    if len(text) <= limit:
        return text, False
    return text[:limit], True
