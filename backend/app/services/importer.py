"""Excel 批量导入。按行保存点提交，失败行不带走其他行。"""
import hashlib
import json
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO

import openpyxl
from openpyxl.styles import Font

from ..errors import BizError
from ..models import (
    Contract,
    ContractOrder,
    ContractType,
    FundAccount,
    Goods,
    GoodsCategory,
    ImportIdem,
    Invoice,
    Party,
    PartyContact,
    Project,
    PurchasePlan,
    SettleBill,
)
from ..money import D, money_str, qty_str, round2, split_line
from . import calc, ops
from .util import (
    MSG_CODE,
    MSG_DIRECTION,
    MSG_LOCK,
    MSG_PERM,
    bump,
    can_overwrite,
    find_by_code,
    op_log,
    parse_date,
    q,
    soft_delete,
    stamp_new,
)

PARSE_FAIL = "文件解析失败，请检查模板格式或必填项"
MAX_SIZE = 10 * 1024 * 1024

MODULES = {"contract", "order", "invoice", "settle", "party", "purchase-plan", "goods", "warehouse-receipt", "hs-code", "customs-material-fill"}


def build_template(module: str, direction: str, party_type: str) -> tuple[bytes, str]:
    _check_module(module, direction, party_type)
    if module == "warehouse-receipt":
        from .customs_xlsx import receipt_template
        return receipt_template()
    if module == "customs-material-fill":
        from .customs_xlsx import material_fill_template
        return material_fill_template()
    if module == "hs-code":
        from .customs_xlsx import hs_template
        return hs_template()
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "导入"
    name, blocks = _template_blocks(module, direction, party_type)
    bold = Font(bold=True)
    row_no = 1
    for headers, _samples in blocks:
        for col, text in enumerate(headers, start=1):
            cell = sheet.cell(row_no, col, text)
            cell.font = bold
        row_no += 2
    for col in sheet.columns:
        sheet.column_dimensions[col[0].column_letter].width = 18
    buf = BytesIO()
    book.save(buf)
    return buf.getvalue(), name


def parse_upload(db, user, company_id, module, content: bytes, filename: str, mode: str, direction: str, party_type: str):
    mode = _normalize_mode(module, mode)
    _check_module(module, direction, party_type)
    if module in ("warehouse-receipt", "hs-code", "customs-material-fill"):
        from . import customs
        return customs.parse_import(db, user, company_id, module, content, filename)
    matrix = read_matrix(content, filename)
    file_hash = hashlib.sha256(content).hexdigest()
    rows, locks, batch_error = _parse_rows(db, user, company_id, module, mode, direction, party_type, matrix)
    if not rows:
        raise BizError(PARSE_FAIL)
    return {
        "file_hash": file_hash,
        "mode": mode,
        "rows": rows,
        "locks": locks,
        "batch_error": batch_error,
    }


def execute_upload(db, user, company_id, module, payload: dict):
    direction = payload.get("direction") or ""
    party_type = payload.get("party_type") or ""
    mode = _normalize_mode(module, payload.get("mode") or "")
    _check_module(module, direction, party_type)
    file_hash = (payload.get("file_hash") or "").strip()
    if not file_hash:
        raise BizError(PARSE_FAIL)
    mode_key = _mode_key(mode, direction, party_type)
    cached = (
        db.query(ImportIdem)
        .filter(
            ImportIdem.company_id == company_id,
            ImportIdem.module == module,
            ImportIdem.mode == mode_key,
            ImportIdem.file_hash == file_hash,
        )
        .first()
    )
    if cached:
        return json.loads(cached.result_json)
    rows = payload.get("rows") or []
    if not isinstance(rows, list) or not rows:
        raise BizError(PARSE_FAIL)
    if module in ("warehouse-receipt", "hs-code", "customs-material-fill"):
        from . import customs
        result = customs.execute_import(db, user, company_id, module, payload)
    elif module == "purchase-plan" and mode == "replace":
        result = _replace_plans(db, user, company_id, rows, payload.get("locks") or [])
    else:
        result = _run_each(db, user, company_id, module, mode, direction, party_type, rows)
    if result["success"] > 0:
        db.add(
            ImportIdem(
                company_id=company_id,
                module=module,
                mode=mode_key,
                file_hash=file_hash,
                result_json=json.dumps(result, ensure_ascii=False),
            )
        )
        op_log(db, user, company_id, "导入", module, "", {"mode": mode, "success": result["success"], "fail": result["fail"]})
    return result


def read_matrix(content: bytes, filename: str):
    name = (filename or "").lower()
    if not name.endswith((".xlsx", ".xls")):
        raise BizError("文件格式不正确")
    if len(content) > MAX_SIZE:
        raise BizError("文件大小超过限制")
    if name.endswith(".xlsx"):
        if not content.startswith(b"PK"):
            raise BizError("文件格式不正确")
        return _read_xlsx(content)
    if not content.startswith(b"\xd0\xcf\x11\xe0"):
        raise BizError("文件格式不正确")
    return _read_xls(content)


def cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        number = Decimal(str(value))
        if number == number.to_integral_value():
            return str(int(number))
        return format(number, "f").rstrip("0").rstrip(".")
    if isinstance(value, int):
        return str(value)
    if isinstance(value, Decimal):
        return format(value, "f")
    text = str(value).strip()
    while text and text[0] in "=+-@":
        text = text[1:]
    return text.strip()


def _read_xlsx(content: bytes):
    try:
        book = openpyxl.load_workbook(BytesIO(content), data_only=False, read_only=True)
        sheet = book.active
        rows = [tuple(row) for row in sheet.iter_rows(values_only=True)]
        book.close()
    except BizError:
        raise
    except Exception as exc:
        raise BizError(PARSE_FAIL) from exc
    return rows


def _read_xls(content: bytes):
    try:
        import xlrd
    except ImportError as exc:
        raise BizError("文件格式不正确") from exc
    try:
        book = xlrd.open_workbook(file_contents=content)
        sheet = book.sheet_by_index(0)
    except Exception as exc:
        raise BizError(PARSE_FAIL) from exc
    rows = []
    for r in range(sheet.nrows):
        values = []
        for c in range(sheet.ncols):
            cell = sheet.cell(r, c)
            if cell.ctype == xlrd.XL_CELL_DATE:
                try:
                    values.append(xlrd.xldate_as_datetime(cell.value, book.datemode))
                except Exception:
                    values.append(cell.value)
            elif cell.ctype == xlrd.XL_CELL_ERROR:
                values.append("")
            else:
                values.append(cell.value)
        rows.append(tuple(values))
    return rows


def _template_blocks(module, direction, party_type):
    if module == "contract":
        purchase = direction != "sale"
        title = "采购合同导入模板.xlsx" if purchase else "销售合同导入模板.xlsx"
        master = ["合同编号", "合同名称", "合同类型", "对方名称", "业务员员工编号", "税率", "签订日期", "标的名称"]
        detail = ["对应合同编号", "物料编号", "中文名称", "规格", "单位", "数量", "含税单价", "备注"]
        if purchase:
            detail.append("计划单号")
            master = [
                "合同编号", "合同名称", "合同类型", "对方名称", "业务员员工编号", "税率", "签订日期", "标的名称",
                "收货地址", "收货人", "运输方式", "交货期限", "运费承担", "付款条款", "验收标准",
            ]
        code = "CG-示例-001" if purchase else "XS-示例-001"
        if purchase:
            head = [code, "示例合同", "材料采购", "示例供应商", "E001", 13, "2026-01-01", "示例标的", "示例地址", "示例收货人", "汽运", "签订后30天", "供方", "按合同阶段付款", "按样品验收"]
        else:
            head = [code, "示例合同", "产品销售", "示例客户", "E001", 13, "2026-01-01", ""]
        line1 = [code, "M001", "热轧钢板", "10mm", "吨", 1, 100, ""]
        line2 = [code, "M002", "高强螺栓", "M16", "套", 2, 5, ""]
        if purchase:
            line1.append("")
            line2.append("")
        return title, [(master, [head]), (detail, [line1, line2])]
    if module == "order":
        purchase = direction != "sale"
        title = "采购订单导入模板.xlsx" if purchase else "销售订单导入模板.xlsx"
        code = "CG-2026-001" if purchase else "XS-2026-001"
        master = ["订单编号", "源合同编号", "单据日期", "备注"]
        detail = ["对应订单编号", "物料编号", "数量", "含税单价", "备注"]
        return title, [(master, [["", code, "2026-01-15", ""]]), (detail, [["", "M001", 1, 100, ""]])]
    if module == "invoice":
        incoming = direction != "output"
        title = "进项发票导入模板.xlsx" if incoming else "销项发票导入模板.xlsx"
        party = "S001" if incoming else "C001"
        code = "CG-2026-001" if incoming else "XS-2026-001"
        master = ["单据编号", "发票号码", "发票日期", "发票类型", "客商编码", "合同编号", "不含税金额", "税额", "价税合计", "备注"]
        detail = ["对应发票号码", "物料编号", "数量", "含税单价", "税率", "备注"]
        head = ["", "INV-示例-001", "2026-02-01", "增值税专用发票", party, code, "", "", "", ""]
        return title, [(master, [head]), (detail, [["INV-示例-001", "M001", 1, 100, 13, ""]])]
    if module == "settle":
        receipt = direction == "receipt"
        title = "收款单导入模板.xlsx" if receipt else "付款单导入模板.xlsx"
        code = "XS-2026-001" if receipt else "CG-2026-001"
        party = "C001" if receipt else "S001"
        headers = ["单据编号", "单据日期", "客商编码", "合同编号", "阶段序号", "阶段名称", "金额", "账户", "备注"]
        sample = ["", "2026-02-01", party, code, 1, "", 10, "基本户", ""]
        return title, [(headers, [sample])]
    if module == "party":
        supplier = party_type != "customer"
        title = "供应商导入模板.xlsx" if supplier else "客户导入模板.xlsx"
        prefix = "供应商" if supplier else "客户"
        headers = [
            f"{prefix}编码", f"{prefix}名称", f"{prefix}简称", "联系人", "联系电话", "手机", "邮箱", "业务员",
            "纳税人识别号", "默认税率", "开票名称", "开户银行", "账户名称", "银行账号", "开户支行",
            "开票电话", "开票地址", "通讯地址", "备注",
        ]
        if supplier:
            sample = [
                "TEST-GYS-001", "test供应商", "test简称", "test联系人", "13800000000", "", "test@test.com", "",
                "91500000TEST000001", 13, "test供应商", "test银行",
                "test供应商", "6222000000000000", "test支行",
                "13800000000", "test地址", "test地址", "test",
            ]
        else:
            sample = ["C900", "示例客户", "示例", "张三", "13800000000", "", "", "E001", "", 13, "示例客户", "", "示例客户", "", "", "", "", "", ""]
        return title, [(headers, [sample])]
    if module == "purchase-plan":
        headers = ["计划月份", "采购序号", "物料编号", "采购计划单号", "使用设备", "中文名称", "英文名称", "规格型号", "单位", "采购数量", "用途", "使用部门", "备注", "库存数量", "在途数量", "4个月消耗量"]
        sample = ["2026-01", "1", "M900", "PL-示例", "钢材", "示例钢板", "", "10mm", "吨", 10, "生产", "采购部", "", "", "", ""]
        return "采购计划导入模板.xlsx", [(headers, [sample])]
    headers = ["物料编号", "名称", "规格", "单位", "分类编码", "含税单价"]
    return "商品导入模板.xlsx", [(headers, [["M900", "示例商品", "10mm", "吨", "", 100]])]


def _parse_rows(db, user, company_id, module, mode, direction, party_type, matrix):
    if module == "party":
        return _parse_party(db, user, company_id, party_type, matrix), [], ""
    if module == "goods":
        return _parse_goods(db, user, company_id, matrix), [], ""
    if module == "purchase-plan":
        rows, batch_error = _parse_plan(db, user, company_id, mode, matrix)
        locks = [{"id": row.id, "version": row.version} for row in q(db, PurchasePlan, company_id).all()] if mode == "replace" else []
        return rows, locks, batch_error
    if module == "contract":
        return _parse_contract(db, user, company_id, direction, mode, matrix), [], ""
    if module == "order":
        return _parse_order(db, user, company_id, direction, matrix), [], ""
    if module == "invoice":
        return _parse_invoice(db, user, company_id, direction, matrix), [], ""
    return _parse_settle(db, user, company_id, direction, matrix), [], ""


def _run_each(db, user, company_id, module, mode, direction, party_type, rows):
    success = 0
    fail = 0
    errors = []
    party_ids = []
    for row in rows:
        if not row.get("ok"):
            fail += 1
            if row.get("message"):
                errors.append({"row_no": row.get("row_no"), "message": row.get("message")})
            continue
        try:
            with db.begin_nested():
                changed = _apply(db, user, company_id, module, mode, direction, party_type, row)
            success += 1
            if changed:
                party_ids.append(changed)
        except BizError as exc:
            fail += 1
            errors.append({"row_no": row.get("row_no"), "message": exc.message})
        except (InvalidOperation, ValueError):
            fail += 1
            errors.append({"row_no": row.get("row_no"), "message": PARSE_FAIL})
    return _result(success, fail, errors, party_ids)


def _apply(db, user, company_id, module, mode, direction, party_type, row):
    if module == "party":
        return _apply_party(db, user, company_id, party_type, row)
    if module == "goods":
        _apply_goods(db, user, company_id, row)
    elif module == "purchase-plan":
        _apply_plan(db, user, company_id, mode, row)
    elif module == "contract":
        _apply_contract(db, user, company_id, direction, mode, row)
    elif module == "order":
        _apply_order(db, user, company_id, direction, row)
    elif module == "invoice":
        _apply_invoice(db, user, company_id, direction, row)
    else:
        _apply_settle(db, user, company_id, direction, row)
    return None


_BAD_CONTACT = {"供应商类别", "客户类别"}


def _has_digit(value: str) -> bool:
    return any(ch.isdigit() for ch in (value or ""))


def _looks_like_employee(value: str) -> bool:
    text = (value or "").strip()
    return bool(text) and all(ch.isascii() and (ch.isalnum() or ch in "-_") for ch in text)


def _find_party_header(matrix):
    codes = {"编码", "客户编码", "供应商编码"}
    names = {"名称", "客户名称", "供应商名称"}
    end = min(len(matrix), 30)
    for i in range(end):
        texts = set(_texts(matrix[i]))
        if texts & codes and texts & names:
            return i
    return None


def _take_contact_token(value, names, phones, emails):
    text = (value or "").strip()
    if not text or text in _BAD_CONTACT:
        return
    if "@" in text:
        emails.append(text)
    elif _has_digit(text):
        phones.append(text)
    else:
        names.append(text)


def _party_contacts(raw, index):
    contact = _pick(raw, index, "联系人")
    phone = _pick(raw, index, "联系电话")
    mobile = _pick(raw, index, "手机")
    email = _pick(raw, index, "邮箱")
    if contact in _BAD_CONTACT:
        contact = ""
    names, phones, emails = [], [], []
    _take_contact_token(phone, names, phones, emails)
    _take_contact_token(mobile, names, phones, emails)
    _take_contact_token(email, names, phones, emails)
    if not contact and names:
        contact = names[0]
    if phone and _has_digit(phone) and "@" not in phone and phone.strip() not in _BAD_CONTACT:
        phone_out = phone.strip()
    else:
        phone_out = phones[0] if phones else ""
    rest = [item for item in phones if item != phone_out]
    if mobile and _has_digit(mobile) and "@" not in mobile:
        mobile_out = mobile.strip()
        if mobile_out == phone_out:
            mobile_out = rest[0] if rest else ""
    else:
        mobile_out = rest[0] if rest else ""
    return contact, phone_out, mobile_out, (emails[0] if emails else "")


def _parse_party(db, user, company_id, party_type, matrix):
    header_at = _find_party_header(matrix)
    if header_at is None:
        raise BizError(PARSE_FAIL)
    index = _header_index(matrix[header_at])
    rows = []
    seen = set()
    for offset, raw in enumerate(matrix[header_at + 1 :], start=header_at + 2):
        if _blank(raw):
            continue
        code = _pick(raw, index, "编码", "客户编码", "供应商编码")
        name = _pick(raw, index, "名称", "客户名称", "供应商名称")
        summary = f"{code} {name}".strip()
        kind_text = _pick(raw, index, "类型")
        kind = {"供应商": "supplier", "客户": "customer", "supplier": "supplier", "customer": "customer"}.get(kind_text, "")
        if not kind:
            kind = party_type
        contact_name, contact_phone, mobile, email = _party_contacts(raw, index)
        salesman = _pick(raw, index, "业务员")
        category = _pick(raw, index, "供应商类别", "供应商类别 / 业务员")
        if category in ("供应商类别", "客户类别"):
            category = ""
        data = {
            "code": code,
            "name": name,
            "short_name": _pick(raw, index, "简称", "客户简称", "供应商简称"),
            "type": kind,
            "tax_no": _pick(raw, index, "纳税人识别号", "税号"),
            "bank_name": _pick(raw, index, "开户银行"),
            "account_name": _pick(raw, index, "账户名称"),
            "bank_account": _pick(raw, index, "银行账号"),
            "bank_branch": _pick(raw, index, "开户支行"),
            "default_tax_rate": _pick(raw, index, "默认税率"),
            "invoice_title": _pick(raw, index, "开票名称"),
            "invoice_phone": _pick(raw, index, "开票电话"),
            "invoice_address": _pick(raw, index, "开票地址"),
            "address": _pick(raw, index, "通讯地址", "地址"),
            "remark": _pick(raw, index, "备注"),
            "email": email,
            "mobile": mobile,
            "contact_name": contact_name,
            "contact_phone": contact_phone,
            "supplier_category": category if kind != "customer" else "",
            "salesman_id": salesman,
        }
        message = ""
        ok = True
        row_id = None
        version = None
        if not code or not name:
            ok, message = False, "编码和名称不能为空"
        elif kind not in ("supplier", "customer"):
            ok, message = False, "类型不正确"
        elif kind != party_type:
            ok, message = False, "类型与当前列表不一致"
        elif code in seen:
            ok, message = False, "编码在文件中重复"
        else:
            current = find_by_code(db, Party, company_id, Party.code, code)
            if current is not None:
                row_id, version = current.id, current.version
                if current.type != kind:
                    ok, message = False, "类型保存后不可改"
                elif not can_overwrite(user, current.created_by):
                    ok, message = False, MSG_PERM
        if ok and data["salesman_id"] and _looks_like_employee(data["salesman_id"]):
            try:
                ops._owner_ok(db, company_id, data["salesman_id"])
            except BizError as exc:
                ok, message = False, exc.message
        if code:
            seen.add(code)
        rows.append(_preview(offset, ok, summary, message, row_id, version, data))
    return rows


def _apply_party(db, user, company_id, party_type, row):
    data = row.get("data") or {}
    current = _locked_master(db, user, Party, company_id, Party.code, data.get("code"), row)
    if current is None:
        current = Party(type=data.get("type") or party_type, status="启用")
        stamp_new(current, user, company_id)
        db.add(current)
        current.code = data.get("code") or ""
        current.default_tax_rate = D(data.get("default_tax_rate") or 0)
        renamed = None
    else:
        if current.type != (data.get("type") or current.type):
            raise BizError("类型保存后不可改")
        renamed = current.id if (current.name or "") != (data.get("name") or "") else None
        if current.deleted_at:
            current.deleted_at = None
            current.deleted_by = None
    if not (data.get("code") or "").strip() or not (data.get("name") or "").strip():
        raise BizError("编码和名称不能为空")
    current.code = data.get("code")
    current.name = data.get("name")
    current.short_name = data.get("short_name") or current.short_name or ""
    current.tax_no = data.get("tax_no") or ""
    current.bank_name = data.get("bank_name") or ""
    current.account_name = data.get("account_name") or ""
    current.bank_account = data.get("bank_account") or ""
    current.bank_branch = data.get("bank_branch") or ""
    current.invoice_title = data.get("invoice_title") or ""
    current.invoice_phone = data.get("invoice_phone") or ""
    current.invoice_address = data.get("invoice_address") or ""
    current.address = data.get("address") or ""
    current.email = data.get("email") or ""
    current.mobile = data.get("mobile") or ""
    current.remark = data.get("remark") or ""
    current.salesman_id = data.get("salesman_id") or ""
    if (data.get("type") or party_type) == "supplier":
        current.supplier_category = data.get("supplier_category") or ""
    if current.id:
        bump(current, user.employee_id)
    db.flush()
    _write_party_contact(db, user, company_id, current, data)
    return renamed


def _write_party_contact(db, user, company_id, party, data):
    name = (data.get("contact_name") or "").strip()
    phone = (data.get("contact_phone") or "").strip()
    if not name and not phone:
        return
    contact = (
        db.query(PartyContact)
        .filter(PartyContact.party_id == party.id, PartyContact.deleted_at.is_(None), PartyContact.is_default.is_(True))
        .first()
    )
    if contact is None:
        contact = PartyContact(party_id=party.id, is_default=True, name=name or "")
        stamp_new(contact, user, company_id)
        db.add(contact)
    contact.name = name or contact.name or ""
    contact.phone = phone
    contact.is_default = True


def _parse_goods(db, user, company_id, matrix):
    header_at = _find_row(matrix, ["物料编号", "名称"])
    if header_at is None:
        raise BizError(PARSE_FAIL)
    index = _header_index(matrix[header_at])
    rows = []
    seen = set()
    for offset, raw in enumerate(matrix[header_at + 1 :], start=header_at + 2):
        if _blank(raw):
            continue
        code = _pick(raw, index, "物料编号")
        name = _pick(raw, index, "名称")
        data = {
            "material_code": code,
            "name": name,
            "spec": _pick(raw, index, "规格"),
            "unit": _pick(raw, index, "单位"),
            "category_code": _pick(raw, index, "分类编码"),
            "price": _pick(raw, index, "含税单价"),
        }
        ok, message = True, ""
        row_id = version = None
        if not code or not name or not data["unit"]:
            ok, message = False, "物料编号、名称、单位不能为空"
        elif code in seen:
            ok, message = False, "物料编号在文件中重复"
        else:
            current = find_by_code(db, Goods, company_id, Goods.material_code, code)
            if current is not None:
                row_id, version = current.id, current.version
                if not can_overwrite(user, current.created_by):
                    ok, message = False, MSG_PERM
            if ok and data["category_code"]:
                found = q(db, GoodsCategory, company_id).filter(GoodsCategory.code == data["category_code"]).first()
                if found is None:
                    ok, message = False, "分类编码不存在"
        if code:
            seen.add(code)
        rows.append(_preview(offset, ok, f"{code} {name}".strip(), message, row_id, version, data))
    return rows


def _apply_goods(db, user, company_id, row):
    data = row.get("data") or {}
    code = (data.get("material_code") or "").strip()
    current = _locked_master(db, user, Goods, company_id, Goods.material_code, code, row)
    if data.get("category_code"):
        found = q(db, GoodsCategory, company_id).filter(GoodsCategory.code == data["category_code"]).first()
        if found is None:
            raise BizError("分类编码不存在")
    if not code or not (data.get("name") or "").strip() or not (data.get("unit") or "").strip():
        raise BizError("物料编号、名称、单位不能为空")
    if current is None:
        current = Goods(material_code=code)
        stamp_new(current, user, company_id)
        db.add(current)
    elif current.deleted_at:
        current.deleted_at = None
        current.deleted_by = None
    current.material_code = code
    current.name = data.get("name") or ""
    current.spec = data.get("spec") or ""
    current.unit = data.get("unit") or ""
    current.category_code = data.get("category_code") or ""
    current.price = round2(data.get("price") or 0)
    if current.id:
        bump(current, user.employee_id)
    db.flush()


def _fill_plan_codes(parsed):
    """来源表没有物料编号列时，用规格型号识别；同规格多行再拼采购序号。"""
    bases = []
    for _, data in parsed:
        if data.get("material_code"):
            bases.append("")
            continue
        bases.append((data.get("spec") or "").strip() or (data.get("name_cn") or "").strip())
    counts = {}
    for base in bases:
        if base:
            counts[base] = counts.get(base, 0) + 1
    for (_, data), base in zip(parsed, bases):
        if data.get("material_code") or not base:
            continue
        if counts[base] == 1:
            data["material_code"] = base
            continue
        seq = (data.get("purchase_seq") or "").strip()
        data["material_code"] = f"{base}-{seq}" if seq else base


def _parse_plan(db, user, company_id, mode, matrix):
    header_at = None
    for i, raw in enumerate(matrix[:12]):
        texts = set(_texts(raw))
        if "中文名称" in texts and "采购数量" in texts:
            header_at = i
            break
    if header_at is None:
        raise BizError(PARSE_FAIL)
    header_names = set(_texts(matrix[header_at]))
    has_code_col = bool(header_names & {"物料编号", "物料编码", "物料代码"})
    index = _header_index(matrix[header_at])
    carry = {}
    parsed = []
    for offset, raw in enumerate(matrix[header_at + 1 :], start=header_at + 2):
        if _blank(raw):
            continue
        if _is_total(raw, index):
            continue
        data = {
            "plan_month": _pick(raw, index, "计划月份", "月份"),
            "purchase_seq": _pick(raw, index, "采购序号", "序号"),
            "material_code": _pick(raw, index, "物料编号", "物料编码", "物料代码"),
            "plan_no": _pick(raw, index, "采购计划单号", "计划单号"),
            "category": _pick(raw, index, "使用设备", "类别"),
            "name_cn": _pick(raw, index, "中文名称"),
            "name_en": _pick(raw, index, "英文名称"),
            "spec": _pick(raw, index, "规格型号", "规格"),
            "unit": _pick(raw, index, "单位"),
            "qty": _pick(raw, index, "采购数量"),
            "purpose": _pick(raw, index, "用途"),
            "use_dept": _pick(raw, index, "使用部门"),
            "remark": _pick(raw, index, "备注"),
            "stock_qty": _pick(raw, index, "库存数量"),
            "transit_qty": _pick(raw, index, "在途数量"),
            "consume_4m": _pick(raw, index, "4个月消耗量", "4 个月消耗量"),
        }
        for key in ("plan_month", "plan_no", "category", "purpose", "use_dept"):
            if data[key]:
                carry[key] = data[key]
            elif carry.get(key):
                data[key] = carry[key]
        parsed.append((offset, data))
    if not has_code_col:
        _fill_plan_codes(parsed)
    counts = {}
    for _, data in parsed:
        if data["material_code"]:
            counts[data["material_code"]] = counts.get(data["material_code"], 0) + 1
    rows = []
    for offset, data in parsed:
        code = data["material_code"]
        ok, message = True, ""
        row_id = version = None
        if not code:
            ok, message = False, "物料编号不能为空"
        elif counts.get(code, 0) > 1:
            ok, message = False, "物料编号在文件中重复"
        elif not data["name_cn"] and not data["spec"]:
            ok, message = False, "没有中文名称且没有规格"
        elif data["qty"] == "":
            ok, message = False, "采购数量不能为空"
        else:
            try:
                qty = D(data["qty"])
            except (InvalidOperation, ValueError):
                ok, message = False, PARSE_FAIL
                qty = None
            current = find_by_code(db, PurchasePlan, company_id, PurchasePlan.material_code, code)
            alive = current is not None and not current.deleted_at
            if current is not None:
                row_id, version = current.id, current.version
            if ok and alive and mode == "append":
                ok, message = False, MSG_CODE
            elif ok and current is not None and not can_overwrite(user, current.created_by):
                ok, message = False, MSG_PERM
            elif ok and alive and mode in ("update", "replace") and qty is not None:
                bought = calc.purchased_qty(db, company_id, code)
                if bought > qty:
                    ok, message = False, f"物料 {code} 超出 {qty_str(bought - qty)}"
        rows.append(_preview(offset, ok, f"{code} {data['name_cn']}".strip(), message, row_id, version, data))
    batch_error = ""
    if mode == "replace":
        batch_error = _plan_batch_error(db, company_id, rows)
        if batch_error:
            for row in rows:
                row["ok"] = False
                if not row["message"]:
                    row["message"] = batch_error
    return rows, batch_error


def _plan_batch_error(db, company_id, rows):
    if any(not row["ok"] for row in rows):
        return "整批导入失败"
    new_qty = {}
    for row in rows:
        new_qty[row["data"]["material_code"]] = D(row["data"]["qty"])
    messages = []
    for plan in q(db, PurchasePlan, company_id).all():
        bought = calc.purchased_qty(db, company_id, plan.material_code)
        new = new_qty.get(plan.material_code, Decimal("0"))
        if bought > new:
            messages.append(f"物料 {plan.material_code} 超出 {qty_str(bought - new)}")
    return "；".join(messages)


def _replace_plans(db, user, company_id, rows, locks):
    if any(not row.get("ok") for row in rows):
        errors = [{"row_no": row.get("row_no"), "message": row.get("message") or "整批导入失败"} for row in rows]
        return _result(0, len(rows), errors, [])
    current = q(db, PurchasePlan, company_id).order_by(PurchasePlan.id).all()
    lock_map = {}
    for item in locks:
        try:
            lock_map[int(item.get("id"))] = int(item.get("version"))
        except (TypeError, ValueError):
            return _result(0, len(rows), [{"row_no": 0, "message": MSG_LOCK}], [])
    if {row.id for row in current} != set(lock_map) or any(int(row.version) != lock_map[row.id] for row in current):
        return _result(0, len(rows), [{"row_no": row.get("row_no"), "message": MSG_LOCK} for row in rows], [])
    codes = [row["data"]["material_code"] for row in rows]
    calc.lock_plans(db, company_id, [p.material_code for p in current] + codes)
    batch_error = _plan_batch_error(db, company_id, rows)
    if batch_error:
        return _result(0, len(rows), [{"row_no": row.get("row_no"), "message": batch_error} for row in rows], [])
    for plan in current:
        if not can_overwrite(user, plan.created_by):
            return _result(0, len(rows), [{"row_no": 0, "message": MSG_PERM}], [])
    new_codes = {row["data"]["material_code"] for row in rows}
    for plan in list(current):
        if plan.material_code not in new_codes:
            soft_delete(plan, user.employee_id)
    for row in rows:
        _apply_plan(db, user, company_id, "update", row)
    db.flush()
    return _result(len(rows), 0, [], [])


def _apply_plan(db, user, company_id, mode, row):
    data = row.get("data") or {}
    code = (data.get("material_code") or "").strip()
    if not code:
        raise BizError("物料编号不能为空")
    if not data.get("name_cn") and not data.get("spec"):
        raise BizError("没有中文名称且没有规格")
    if data.get("qty") in (None, ""):
        raise BizError("采购数量不能为空")
    qty = D(data.get("qty"))
    calc.lock_plans(db, company_id, [code])
    current = find_by_code(db, PurchasePlan, company_id, PurchasePlan.material_code, code)
    alive = current is not None and not current.deleted_at
    if alive and mode == "append":
        raise BizError(MSG_CODE)
    if current is not None:
        _check_lock(user, current, row)
        bought = calc.purchased_qty(db, company_id, code)
        if qty < bought:
            raise BizError(f"物料 {code} 超出 {qty_str(bought - qty)}")
    else:
        if row.get("id"):
            raise BizError(MSG_LOCK, 409)
        current = PurchasePlan(material_code=code)
        stamp_new(current, user, company_id)
        db.add(current)
    if current.deleted_at:
        current.deleted_at = None
        current.deleted_by = None
    current.plan_month = data.get("plan_month") or ""
    current.purchase_seq = data.get("purchase_seq") or ""
    current.material_code = code
    current.plan_no = data.get("plan_no") or ""
    current.category = data.get("category") or ""
    current.name_cn = data.get("name_cn") or ""
    current.name_en = data.get("name_en") or ""
    current.spec = data.get("spec") or ""
    current.unit = data.get("unit") or ""
    current.qty = qty
    current.purpose = data.get("purpose") or ""
    current.use_dept = data.get("use_dept") or ""
    current.remark = data.get("remark") or ""
    current.stock_qty = _opt_qty(data.get("stock_qty"))
    current.transit_qty = _opt_qty(data.get("transit_qty"))
    current.consume_4m = _opt_qty(data.get("consume_4m"))
    if current.id:
        bump(current, user.employee_id)
    db.flush()


def _parse_contract(db, user, company_id, direction, mode, matrix):
    master_at = _find_row(matrix, ["合同编号", "合同名称"])
    detail_at = _find_row(matrix, ["物料编号", "数量"])
    if master_at is None or detail_at is None:
        raise BizError(PARSE_FAIL)
    stage_at = _find_stage_header(matrix, detail_at)
    groups = {}
    order = []
    if master_at == detail_at:
        index = _header_index(matrix[master_at])
        end = stage_at if stage_at is not None else len(matrix)
        for offset, raw in enumerate(matrix[master_at + 1 : end], start=master_at + 2):
            if _blank(raw):
                continue
            code = _pick(raw, index, "合同编号", "对应合同编号")
            _touch_group(groups, order, code, offset)
            if not code:
                continue
            _fill_contract_header(groups[code], raw, index)
            _append_contract_item(groups[code], raw, index, direction)
    else:
        master_index = _header_index(matrix[master_at])
        detail_index = _header_index(matrix[detail_at])
        for offset, raw in enumerate(matrix[master_at + 1 : detail_at], start=master_at + 2):
            if _blank(raw):
                continue
            code = _pick(raw, master_index, "合同编号")
            _touch_group(groups, order, code, offset)
            if code:
                _fill_contract_header(groups[code], raw, master_index)
        end = stage_at if stage_at is not None else len(matrix)
        for offset, raw in enumerate(matrix[detail_at + 1 : end], start=detail_at + 2):
            if _blank(raw):
                continue
            code = _pick(raw, detail_index, "对应合同编号", "合同编号")
            _touch_group(groups, order, code, offset)
            if code:
                _append_contract_item(groups[code], raw, detail_index, direction)
    if stage_at is not None:
        _read_contract_stages(matrix, stage_at, groups, order)
    rows = []
    for code in order:
        group = groups[code]
        rows.append(_review_contract(db, user, company_id, direction, mode, code, group))
    return rows


def _plan_indexes(plans):
    by_code = {}
    by_spec = {}
    by_name = {}
    for plan in plans:
        by_code[plan.material_code] = plan
        spec = (plan.spec or "").strip()
        name = (plan.name_cn or "").strip()
        if spec:
            by_spec.setdefault(spec, []).append(plan)
        if name:
            by_name.setdefault(name, []).append(plan)
    return by_code, by_spec, by_name


def _match_plan(indexes, code, spec, name):
    by_code, by_spec, by_name = indexes
    code = (code or "").strip()
    spec = (spec or "").strip()
    name = (name or "").strip()
    if code and code in by_code:
        return by_code[code]
    for key in (spec, code):
        hits = by_spec.get(key) or []
        if key and len(hits) == 1:
            return hits[0]
    if not code and not spec and name:
        hits = by_name.get(name) or []
        if len(hits) == 1:
            return hits[0]
    return None


def _apply_plan_match(item, plan):
    item["material_code"] = plan.material_code
    item["name_cn"] = item.get("name_cn") or plan.name_cn or ""
    item["name_en"] = item.get("name_en") or plan.name_en or ""
    item["spec"] = item.get("spec") or plan.spec or ""
    item["unit"] = item.get("unit") or plan.unit or ""
    item["plan_no"] = item.get("plan_no") or plan.plan_no or ""
    item["category"] = item.get("category") or plan.category or ""
    seq = str(plan.purchase_seq or "")
    item["plan_row"] = int(seq) if seq.isdigit() else None
    item["matched_plan"] = True


def _review_contract(db, user, company_id, direction, mode, code, group):
    if direction == "purchase":
        indexes = _plan_indexes(q(db, PurchasePlan, company_id).all())
        for item in group["items"]:
            plan = _match_plan(indexes, item.get("material_code"), item.get("spec"), item.get("name_cn"))
            if plan is not None:
                _apply_plan_match(item, plan)
    data = group["header"]
    data["contract_code"] = code
    data["items"] = group["items"]
    data["stages"] = group.get("stages") or []
    summary = f"{code} {data.get('name') or ''}（{len(group['items'])} 行物料）".strip()
    ok, message = True, ""
    row_id = version = None
    if not code:
        ok, message = False, "合同编号不能为空"
    elif group.get("conflict"):
        ok, message = False, f"{group['conflict']}不一致"
    elif not data.get("name"):
        ok, message = False, "合同名称不能为空"
    elif not group["items"]:
        ok, message = False, "至少一行物料"
    else:
        seen = set()
        for item in group["items"]:
            if not item.get("material_code"):
                ok, message = False, "物料编号不能为空"
                break
            if item["material_code"] in seen:
                ok, message = False, f"物料 {item['material_code']} 重复"
                break
            seen.add(item["material_code"])
            if item.get("qty") in (None, "") or item.get("price") in (None, ""):
                ok, message = False, "数量和含税单价不能为空"
                break
        current = find_by_code(db, Contract, company_id, Contract.contract_code, code) if ok else None
        if ok and current is not None and current.deleted_at:
            ok, message = False, MSG_CODE
        elif ok and current is not None:
            row_id, version = current.id, current.version
            if current.direction != direction:
                ok, message = False, MSG_DIRECTION
            elif not can_overwrite(user, current.created_by):
                ok, message = False, MSG_PERM
            elif mode == "append":
                have = {item.material_code for item in calc.items_of(db, current.id)}
                dup = seen & have
                if dup:
                    ok, message = False, f"物料 {sorted(dup)[0]} 重复"
        if ok:
            party = _resolve_party(db, company_id, data, direction)
            if isinstance(party, str):
                ok, message = False, party
        if ok and data.get("type_name"):
            found = q(db, ContractType, company_id).filter(ContractType.name == data["type_name"], ContractType.status == "启用").first()
            if found is None:
                ok, message = False, "合同类型不存在或已停用"
        if ok and data.get("project_code"):
            project = find_by_code(db, Project, company_id, Project.project_code, data["project_code"])
            if project is None or project.deleted_at:
                ok, message = False, "项目不存在"
        if ok and direction == "purchase":
            qty_map = {item["material_code"]: item["qty"] for item in group["items"]}
            if mode == "append" and row_id:
                for item in calc.items_of(db, row_id):
                    qty_map[item.material_code] = item.qty
            try:
                calc.assert_plans(db, company_id, row_id, qty_map)
            except BizError as exc:
                ok, message = False, exc.message
    return _preview(group["row_no"], ok, summary, message, row_id, version, data)


def _text_or_current(data, key, current, attr=None):
    value = data.get(key)
    if value not in (None, ""):
        return value
    if current is None:
        return ""
    return getattr(current, attr or key) or ""


def _apply_contract(db, user, company_id, direction, mode, row):
    data = row.get("data") or {}
    code = (data.get("contract_code") or "").strip()
    current = find_by_code(db, Contract, company_id, Contract.contract_code, code)
    if current is not None and current.deleted_at:
        raise BizError(MSG_CODE)
    if current is not None:
        _check_lock(user, current, row)
        if current.direction != direction:
            raise BizError(MSG_DIRECTION)
    elif row.get("id"):
        raise BizError(MSG_LOCK, 409)
    party = _resolve_party(db, company_id, data, direction)
    if isinstance(party, str):
        raise BizError(party)
    items = list(data.get("items") or [])
    if current is not None and mode == "append":
        merged = []
        seen = set()
        for item in calc.items_of(db, current.id):
            seen.add(item.material_code)
            merged.append({
                "material_code": item.material_code,
                "name_cn": item.name_cn,
                "name_en": item.name_en,
                "spec": item.spec,
                "unit": item.unit,
                "qty": item.qty,
                "price": item.price,
                "plan_no": item.plan_no,
                "category": item.category,
                "remark": item.remark or "",
            })
        for item in items:
            if item.get("material_code") in seen:
                raise BizError(f"物料 {item.get('material_code')} 重复")
            seen.add(item.get("material_code"))
            merged.append(_complete_item(db, company_id, item))
        items = merged
    else:
        items = [_complete_item(db, company_id, item) for item in items]
    pricing = _pricing(data.get("pricing"))
    if pricing is None:
        pricing = current.pricing if current is not None else "fixed"
    payload = {
        "direction": direction,
        "nature": "main",
        "contract_code": code,
        "name": data.get("name") or (current.name if current else ""),
        "type_name": data.get("type_name") or (current.type_name if current else None),
        "party_id": party.id,
        "owner_id": data.get("owner_id") or (current.owner_id if current else ""),
        "status": data.get("status") or (current.status if current else "已签订"),
        "pricing": pricing,
        "tax_rate": data.get("tax_rate") if data.get("tax_rate") not in (None, "") else (current.tax_rate if current else 13),
        "sign_date": data.get("sign_date") or (current.sign_date if current else None),
        "start_date": data.get("start_date") or (current.start_date if current else None),
        "end_date": data.get("end_date") or (current.end_date if current else None),
        "content": _text_or_current(data, "content", current),
        "remark": _text_or_current(data, "remark", current),
        "terms": _text_or_current(data, "terms", current),
        "subject_name": _text_or_current(data, "subject_name", current),
        "ship_address": _text_or_current(data, "ship_address", current),
        "receiver_name": _text_or_current(data, "receiver_name", current),
        "transport_mode": _text_or_current(data, "transport_mode", current),
        "delivery_term": _text_or_current(data, "delivery_term", current),
        "freight_bearer": _text_or_current(data, "freight_bearer", current),
        "acceptance_text": _text_or_current(data, "acceptance_text", current),
        "annex_note": _text_or_current(data, "annex_note", current),
        "items": items,
        "keep_stages": current is not None,
        "version": row.get("version"),
    }
    project_code = (data.get("project_code") or "").strip()
    if project_code:
        project = find_by_code(db, Project, company_id, Project.project_code, project_code)
        if project is None or project.deleted_at:
            raise BizError("项目不存在")
        payload["project_id"] = project.id
    elif current is not None:
        payload["project_id"] = current.project_id
    stages = [stage for stage in (data.get("stages") or []) if (stage.get("name") or "").strip()]
    if stages and current is None:
        payload["stages"] = stages
        payload["realloc_stages"] = True
    ops.save_contract(db, user, company_id, payload, current, as_import=True)


def _parse_order(db, user, company_id, direction, matrix):
    master_at = _find_row(matrix, ["源合同编号"])
    detail_at = _find_row(matrix, ["物料编号", "数量"])
    if master_at is None or detail_at is None:
        raise BizError(PARSE_FAIL)
    groups = {}
    order = []
    if master_at == detail_at:
        index = _header_index(matrix[master_at])
        for offset, raw in enumerate(matrix[master_at + 1 :], start=master_at + 2):
            if _blank(raw):
                continue
            key = _order_key(raw, index)
            _touch_group(groups, order, key, offset)
            if not key:
                continue
            _fill_order_header(groups[key], raw, index)
            _append_order_item(groups[key], raw, index)
    else:
        master_index = _header_index(matrix[master_at])
        detail_index = _header_index(matrix[detail_at])
        for offset, raw in enumerate(matrix[master_at + 1 : detail_at], start=master_at + 2):
            if _blank(raw):
                continue
            key = _order_key(raw, master_index)
            _touch_group(groups, order, key, offset)
            if key:
                _fill_order_header(groups[key], raw, master_index)
        for offset, raw in enumerate(matrix[detail_at + 1 :], start=detail_at + 2):
            if _blank(raw):
                continue
            linked = _pick(raw, detail_index, "对应订单编号")
            if linked and ("order:" + linked) in groups:
                key = "order:" + linked
            else:
                key = _order_key(raw, detail_index)
            if not key and order:
                key = order[-1]
            _touch_group(groups, order, key, offset)
            if key:
                _append_order_item(groups[key], raw, detail_index)
    rows = []
    for key in order:
        group = groups[key]
        data = group["header"]
        data["items"] = group["items"]
        contract_code = data.get("contract_code") or ""
        summary = f"{data.get('order_code') or '自动编号'} {contract_code}（{len(group['items'])} 行物料）"
        ok, message = True, ""
        if not contract_code:
            ok, message = False, "源合同编号不能为空"
        elif not group["items"]:
            ok, message = False, "至少一行物料"
        else:
            contract = _alive_contract(db, company_id, contract_code)
            if contract is None:
                ok, message = False, "合同不存在"
            elif contract.direction != direction:
                ok, message = False, MSG_DIRECTION
            elif contract.status != "已签订":
                ok, message = False, "未签订与已终止的合同不能新建后续单据"
            elif q(db, ContractOrder, company_id).filter(ContractOrder.contract_id == contract.id).first():
                ok, message = False, "该合同已有订单"
            else:
                have = {item.material_code: item for item in calc.items_of(db, contract.id)}
                seen = set()
                amount = Decimal("0")
                for item in group["items"]:
                    code = item.get("material_code")
                    if not code or code not in have:
                        ok, message = False, f"物料 {code or ''} 不在合同中".strip()
                        break
                    if code in seen:
                        ok, message = False, f"物料 {code} 重复"
                        break
                    seen.add(code)
                    if item.get("qty") in (None, "") or item.get("price") in (None, ""):
                        ok, message = False, "数量和含税单价不能为空"
                        break
                    amount += round2(D(item["qty"]) * D(item["price"]))
                if ok and data.get("order_code"):
                    taken = find_by_code(db, ContractOrder, company_id, ContractOrder.order_code, data["order_code"])
                    if taken is not None:
                        ok, message = False, MSG_CODE
                if ok and amount > D(contract.amount):
                    ok, message = False, f"超出 {money_str(amount - D(contract.amount))}"
        rows.append(_preview(group["row_no"], ok, summary, message, None, None, data))
    return rows


def _apply_order(db, user, company_id, direction, row):
    data = row.get("data") or {}
    contract = _alive_contract(db, company_id, data.get("contract_code"))
    if contract is None:
        raise BizError("合同不存在")
    if contract.direction != direction:
        raise BizError(MSG_DIRECTION)
    have = {item.material_code: item for item in calc.items_of(db, contract.id)}
    items = []
    for raw in data.get("items") or []:
        src = have.get(raw.get("material_code"))
        if src is None:
            raise BizError(f"物料 {raw.get('material_code') or ''} 不在合同中")
        items.append({
            "material_code": src.material_code,
            "name_cn": raw.get("name_cn") or src.name_cn,
            "spec": raw.get("spec") or src.spec,
            "unit": raw.get("unit") or src.unit,
            "qty": raw.get("qty"),
            "price": raw.get("price"),
            "remark": raw.get("remark") or "",
        })
    ops.save_order(db, user, company_id, {
        "contract_id": contract.id,
        "direction": direction,
        "order_code": data.get("order_code") or "",
        "remark": data.get("remark") or "",
        "items": items,
    })


def _parse_invoice(db, user, company_id, direction, matrix):
    master_at = _find_row(matrix, ["发票号码"])
    detail_at = _find_row(matrix, ["物料编号", "数量"])
    if master_at is None or detail_at is None or "合同编号" not in set(_texts(matrix[master_at])):
        raise BizError(PARSE_FAIL)
    groups = {}
    order = []
    if master_at == detail_at:
        index = _header_index(matrix[master_at])
        for offset, raw in enumerate(matrix[master_at + 1 :], start=master_at + 2):
            if _blank(raw):
                continue
            number = _pick(raw, index, "发票号码", "对应发票号码")
            _touch_group(groups, order, number, offset)
            if not number:
                continue
            _fill_invoice_header(groups[number], raw, index)
            _append_invoice_item(groups[number], raw, index)
    else:
        master_index = _header_index(matrix[master_at])
        detail_index = _header_index(matrix[detail_at])
        for offset, raw in enumerate(matrix[master_at + 1 : detail_at], start=master_at + 2):
            if _blank(raw):
                continue
            number = _pick(raw, master_index, "发票号码")
            _touch_group(groups, order, number, offset)
            if number:
                _fill_invoice_header(groups[number], raw, master_index)
        for offset, raw in enumerate(matrix[detail_at + 1 :], start=detail_at + 2):
            if _blank(raw):
                continue
            number = _pick(raw, detail_index, "对应发票号码", "发票号码")
            _touch_group(groups, order, number, offset)
            if number:
                _append_invoice_item(groups[number], raw, detail_index)
    rows = []
    for number in order:
        group = groups[number]
        data = group["header"]
        data["invoice_no"] = number
        data["items"] = group["items"]
        summary = f"{number} {data.get('contract_code') or ''}（{len(group['items'])} 行明细）"
        ok, message = True, ""
        if not number:
            ok, message = False, "发票号码不能为空"
        elif group.get("conflict"):
            ok, message = False, "同一发票号码的头字段不一致"
        elif not data.get("contract_code"):
            ok, message = False, "合同编号不能为空"
        elif not group["items"]:
            ok, message = False, "至少一行明细才能保存"
        else:
            exists = db.query(Invoice).filter(Invoice.company_id == company_id, Invoice.invoice_no == number).first()
            if exists is not None:
                ok, message = False, "发票号码已存在"
            else:
                contract = _alive_contract(db, company_id, data.get("contract_code"))
                expect = "output" if contract and contract.direction == "sale" else "input"
                if contract is None:
                    ok, message = False, "合同不存在"
                elif direction != expect:
                    ok, message = False, MSG_DIRECTION
                elif contract.status != "已签订":
                    ok, message = False, "未签订与已终止的合同不能新建后续单据"
                else:
                    party_msg = _party_by_code(db, company_id, data.get("party_code") or _party_code(db, contract), direction)
                    if isinstance(party_msg, str):
                        ok, message = False, party_msg
                    have = {item.material_code for item in calc.items_of(db, contract.id)}
                    seen = set()
                    total = Decimal("0")
                    ex_tax = Decimal("0")
                    tax = Decimal("0")
                    for item in group["items"]:
                        code = item.get("material_code")
                        if not code or code not in have:
                            ok, message = False, f"物料 {code or ''} 不在合同中".strip()
                            break
                        if code in seen:
                            ok, message = False, f"物料 {code} 重复"
                            break
                        seen.add(code)
                        rate = item.get("tax_rate") or data.get("tax_rate") or contract.tax_rate
                        part = split_line(item.get("qty"), item.get("price"), rate)
                        total += part["total"]
                        ex_tax += part["ex_tax"]
                        tax += part["tax"]
                    if ok:
                        message = _amount_hint(data, ex_tax, tax, total)
        rows.append(_preview(group["row_no"], ok, summary, message, None, None, data))
    return rows


def _apply_invoice(db, user, company_id, direction, row):
    data = row.get("data") or {}
    number = (data.get("invoice_no") or "").strip()
    if db.query(Invoice).filter(Invoice.company_id == company_id, Invoice.invoice_no == number).first():
        raise BizError("发票号码已存在")
    contract = _alive_contract(db, company_id, data.get("contract_code"))
    if contract is None:
        raise BizError("合同不存在")
    party = _party_by_code(db, company_id, data.get("party_code") or _party_code(db, contract), direction)
    if isinstance(party, str):
        raise BizError(party)
    ops.save_invoice(db, user, company_id, {
        "contract_id": contract.id,
        "direction": direction,
        "party_id": party.id,
        "bill_code": data.get("bill_code") or "",
        "invoice_no": number,
        "invoice_date": data.get("invoice_date") or None,
        "invoice_type": data.get("invoice_type") or "增值税专用发票",
        "remark": data.get("remark") or "",
        "items": data.get("items") or [],
    })


def _parse_settle(db, user, company_id, direction, matrix):
    header_at = _find_row(matrix, ["合同编号", "金额"])
    if header_at is None:
        raise BizError(PARSE_FAIL)
    index = _header_index(matrix[header_at])
    rows = []
    for offset, raw in enumerate(matrix[header_at + 1 :], start=header_at + 2):
        if _blank(raw):
            continue
        data = {
            "bill_code": _pick(raw, index, "单据编号"),
            "bill_date": _pick(raw, index, "单据日期"),
            "party_code": _pick(raw, index, "客商编码"),
            "contract_code": _pick(raw, index, "合同编号"),
            "stage_seq": _pick(raw, index, "阶段序号"),
            "stage_name": _pick(raw, index, "阶段名称"),
            "amount": _pick(raw, index, "金额"),
            "account": _pick(raw, index, "账户"),
            "remark": _pick(raw, index, "备注"),
        }
        summary = f"{data['bill_code'] or '自动编号'} {data['contract_code']} 阶段{data['stage_seq'] or data['stage_name']}"
        ok, message = True, ""
        contract = _alive_contract(db, company_id, data["contract_code"]) if data["contract_code"] else None
        if not data["contract_code"] or data["amount"] == "":
            ok, message = False, "合同编号和金额不能为空"
        elif contract is None:
            ok, message = False, "合同不存在"
        else:
            expect = "receipt" if contract.direction == "sale" else "payment"
            if direction != expect:
                ok, message = False, MSG_DIRECTION
            elif contract.status != "已签订":
                ok, message = False, "未签订与已终止的合同不能新建后续单据"
            else:
                party = _party_by_code(db, company_id, data["party_code"] or _party_code(db, contract), direction)
                if isinstance(party, str):
                    ok, message = False, party
                elif party.id != contract.party_id:
                    ok, message = False, MSG_DIRECTION
                else:
                    stage, message = _find_stage(db, contract.id, data["stage_seq"], data["stage_name"])
                    if stage is None:
                        ok = False
                    elif data["bill_code"] and find_by_code(db, SettleBill, company_id, SettleBill.bill_code, data["bill_code"]):
                        ok, message = False, MSG_CODE
                    elif data["account"]:
                        found = q(db, FundAccount, company_id).filter(FundAccount.name == data["account"], FundAccount.status == "启用").first()
                        if found is None:
                            ok, message = False, "账户不存在或已停用"
        rows.append(_preview(offset, ok, summary, message, None, None, data))
    return rows


def _apply_settle(db, user, company_id, direction, row):
    data = row.get("data") or {}
    contract = _alive_contract(db, company_id, data.get("contract_code"))
    if contract is None:
        raise BizError("合同不存在")
    party = _party_by_code(db, company_id, data.get("party_code") or _party_code(db, contract), direction)
    if isinstance(party, str):
        raise BizError(party)
    stage, message = _find_stage(db, contract.id, data.get("stage_seq"), data.get("stage_name"))
    if stage is None:
        raise BizError(message or "阶段不存在")
    ops.save_settle_many(db, user, company_id, {
        "contract_id": contract.id,
        "direction": direction,
        "party_id": party.id,
        "bill_code": data.get("bill_code") or "",
        "bill_date": data.get("bill_date") or None,
        "account_name": data.get("account") or None,
        "remark": data.get("remark") or "",
        "lines": [{"stage_id": stage.id, "amount": data.get("amount"), "remark": data.get("remark") or ""}],
    })


def _find_stage_header(matrix, detail_at):
    for index, raw in enumerate(matrix):
        if index <= detail_at:
            continue
        texts = set(_texts(raw))
        if "阶段名称" in texts and "比例" in texts:
            return index
    return None


def _read_contract_stages(matrix, stage_at, groups, order):
    index = _header_index(matrix[stage_at])
    loose = []
    coded = {}
    for raw in matrix[stage_at + 1 :]:
        if _blank(raw):
            continue
        name = _pick(raw, index, "阶段名称")
        if not name or name.startswith("【"):
            continue
        stage = {
            "name": name,
            "ratio": _pick(raw, index, "比例") or "0",
            "plan_date": _pick(raw, index, "计划日期"),
            "remark": _pick(raw, index, "阶段备注", "备注"),
            "plan_amount": 0,
        }
        code = _pick(raw, index, "对应合同编号", "合同编号")
        if code:
            coded.setdefault(code, []).append(stage)
        else:
            loose.append(stage)
    if coded:
        for code, stages in coded.items():
            if code in groups:
                groups[code]["stages"] = stages
    elif len(order) == 1 and loose:
        groups[order[0]]["stages"] = loose


def _fill_contract_header(group, raw, index):
    fields = {
        "name": ("合同名称",),
        "type_name": ("合同类型",),
        "party_code": ("对方编码", "对方编号"),
        "party_name": ("对方名称",),
        "owner_id": ("业务员员工编号",),
        "status": ("合同状态",),
        "pricing": ("计价方式",),
        "tax_rate": ("税率",),
        "sign_date": ("签订日期",),
        "start_date": ("开始日期",),
        "end_date": ("结束日期",),
        "subject_name": ("标的名称",),
        "project_code": ("项目编号",),
        "ship_address": ("收货地址",),
        "receiver_name": ("收货人",),
        "transport_mode": ("运输方式",),
        "delivery_term": ("交货期限",),
        "freight_bearer": ("运费承担",),
        "terms": ("付款条款", "条款"),
        "acceptance_text": ("验收标准",),
        "annex_note": ("附件说明",),
        "content": ("内容概述",),
        "remark": ("备注",),
    }
    _fill_fields(group, raw, index, fields)


def _append_contract_item(group, raw, index, direction):
    code = _pick(raw, index, "物料编号")
    if not code:
        return
    item = {
        "material_code": code,
        "name_cn": _pick(raw, index, "中文名称"),
        "spec": _pick(raw, index, "规格", "规格型号"),
        "unit": _pick(raw, index, "单位"),
        "qty": _pick(raw, index, "数量"),
        "price": _pick(raw, index, "含税单价"),
        "remark": _pick(raw, index, "明细备注", "备注"),
        "plan_no": _pick(raw, index, "计划单号") if direction == "purchase" else "",
        "category": _pick(raw, index, "使用设备", "类别") if direction == "purchase" else "",
    }
    group["items"].append(item)


def _fill_order_header(group, raw, index):
    _fill_fields(group, raw, index, {
        "order_code": ("订单编号",),
        "contract_code": ("源合同编号",),
        "bill_date": ("单据日期",),
        "remark": ("备注",),
    })


def _append_order_item(group, raw, index):
    code = _pick(raw, index, "物料编号")
    if not code:
        return
    group["items"].append({
        "material_code": code,
        "qty": _pick(raw, index, "数量"),
        "price": _pick(raw, index, "含税单价"),
        "remark": _pick(raw, index, "明细备注", "备注"),
    })


def _order_key(raw, index):
    code = _pick(raw, index, "订单编号", "对应订单编号")
    contract = _pick(raw, index, "源合同编号")
    if code:
        return "order:" + code
    if contract:
        return "contract:" + contract
    return ""


def _fill_invoice_header(group, raw, index):
    _fill_fields(group, raw, index, {
        "bill_code": ("单据编号",),
        "invoice_date": ("发票日期",),
        "invoice_type": ("发票类型",),
        "party_code": ("客商编码",),
        "contract_code": ("合同编号",),
        "amount_ex_tax": ("不含税金额",),
        "tax": ("税额",),
        "total": ("价税合计",),
        "remark": ("备注",),
    }, conflict_keys={"invoice_date", "invoice_type", "party_code", "contract_code"})


def _append_invoice_item(group, raw, index):
    code = _pick(raw, index, "物料编号")
    if not code:
        return
    group["items"].append({
        "material_code": code,
        "qty": _pick(raw, index, "数量"),
        "price": _pick(raw, index, "含税单价"),
        "tax_rate": _pick(raw, index, "税率"),
        "remark": _pick(raw, index, "明细备注", "备注"),
    })


def _fill_fields(group, raw, index, fields, conflict_keys=None):
    conflict_keys = conflict_keys or set(fields)
    for key, names in fields.items():
        value = _pick(raw, index, *names)
        if not value:
            continue
        old = group["header"].get(key)
        if old and old != value and key in conflict_keys:
            group["conflict"] = names[0]
        elif not old:
            group["header"][key] = value


def _touch_group(groups, order, key, row_no):
    if key not in groups:
        groups[key] = {"header": {}, "items": [], "stages": [], "row_no": row_no, "conflict": ""}
        order.append(key)
    elif groups[key]["row_no"] is None:
        groups[key]["row_no"] = row_no


def _find_stage(db, contract_id, seq_text, name):
    stages = calc.stages_of(db, contract_id)
    if seq_text not in (None, ""):
        try:
            seq = int(D(seq_text))
        except (InvalidOperation, ValueError):
            return None, "阶段序号不正确"
        if seq < 1 or seq > len(stages):
            return None, "阶段不存在"
        return stages[seq - 1], ""
    if not name:
        return None, "阶段序号或阶段名称不能为空"
    found = [stage for stage in stages if stage.name == name]
    if len(found) != 1:
        return None, "阶段名称不唯一" if len(found) > 1 else "阶段不存在"
    return found[0], ""


def _amount_hint(data, ex_tax, tax, total):
    hints = []
    pairs = (("不含税金额", data.get("amount_ex_tax"), ex_tax), ("税额", data.get("tax"), tax), ("价税合计", data.get("total"), total))
    for label, raw, calc_value in pairs:
        if raw in (None, ""):
            continue
        try:
            if round2(raw) != round2(calc_value):
                hints.append(label)
        except (InvalidOperation, ValueError):
            hints.append(label)
    if not hints:
        return ""
    return "头金额与明细不一致，将以明细为准"


def _complete_item(db, company_id, item):
    goods = find_by_code(db, Goods, company_id, Goods.material_code, item.get("material_code"))
    if goods is not None and goods.deleted_at:
        goods = None
    return {
        "material_code": item.get("material_code"),
        "name_cn": item.get("name_cn") or (goods.name if goods else ""),
        "name_en": item.get("name_en") or "",
        "spec": item.get("spec") or (goods.spec if goods else ""),
        "unit": item.get("unit") or (goods.unit if goods else ""),
        "qty": item.get("qty"),
        "price": item.get("price"),
        "plan_no": item.get("plan_no") or None,
        "category": item.get("category") or None,
        "remark": item.get("remark") or "",
    }


def _resolve_party(db, company_id, data, direction):
    code = (data.get("party_code") or "").strip()
    name = (data.get("party_name") or "").strip()
    if code:
        party = _party_by_code(db, company_id, code, direction)
        if isinstance(party, str) or not name or party.name == name:
            return party
        return "对方名称和对方编码不一致"
    if name:
        return _party_by_name(db, company_id, name, direction)
    return "对方名称不能为空"


def _party_by_name(db, company_id, name, direction):
    expect = "supplier" if direction in ("purchase", "input", "payment") else "customer"
    rows = q(db, Party, company_id).filter(Party.name == name, Party.type == expect).all()
    if not rows:
        return "对方不存在"
    if len(rows) > 1:
        return "对方名称不唯一"
    if rows[0].status != "启用":
        return "客商已停用"
    return rows[0]


def _party_by_code(db, company_id, code, direction):
    expect = "supplier" if direction in ("purchase", "input", "payment") else "customer"
    if not code:
        return "客商不存在"
    party = find_by_code(db, Party, company_id, Party.code, code)
    if party is None or party.deleted_at:
        return "客商不存在"
    if party.type != expect:
        return MSG_DIRECTION
    if party.status != "启用":
        return "客商已停用"
    return party


def _party_code(db, contract):
    party = db.get(Party, contract.party_id)
    return party.code if party else ""


def _alive_contract(db, company_id, code):
    row = find_by_code(db, Contract, company_id, Contract.contract_code, (code or "").strip())
    if row is None or row.deleted_at:
        return None
    return row


def _locked_master(db, user, model, company_id, column, code, row):
    current = find_by_code(db, model, company_id, column, code)
    if current is None:
        if row.get("id"):
            raise BizError(MSG_LOCK, 409)
        return None
    _check_lock(user, current, row)
    return current


def _check_lock(user, current, row):
    if not can_overwrite(user, current.created_by):
        raise BizError(MSG_PERM, 403)
    version = row.get("version")
    if version is None or int(version) != int(current.version):
        raise BizError(MSG_LOCK, 409)


def _pricing(value):
    if value in (None, ""):
        return None
    mapped = {"固定总价": "fixed", "按实结算": "actual", "fixed": "fixed", "actual": "actual"}.get(value)
    if mapped is None:
        raise BizError("计价方式不正确")
    return mapped


def _opt_qty(value):
    if value in (None, ""):
        return None
    return D(value)


def _result(success, fail, errors, party_ids):
    return {
        "success": success,
        "fail": fail,
        "errors": errors,
        "need_name_sync": bool(party_ids),
        "party_ids": party_ids,
        "message": "导入完成",
    }


def _preview(row_no, ok, summary, message, row_id, version, data):
    return {
        "row_no": row_no,
        "ok": bool(ok),
        "summary": summary,
        "message": message or "",
        "id": row_id,
        "version": version,
        "data": data,
    }


def _find_row(matrix, required):
    end = min(len(matrix), 30)
    for i in range(end):
        texts = set(_texts(matrix[i]))
        if all(name in texts for name in required):
            return i
    return None


def _header_index(raw):
    index = {}
    for i, value in enumerate(raw):
        text = cell_text(value)
        if text and text not in index:
            index[text] = i
    return index


def _first_positive(raw, index, *names):
    """数量列可能叫「采购数量」或「数量」。空列或 0 不挡住后面的有效列。"""
    fallback = ""
    for name in names:
        text = _pick(raw, index, name)
        if not text:
            continue
        if not fallback:
            fallback = text
        try:
            if D(text) > 0:
                return text
        except BizError:
            return text
    return fallback or "0"


def _pick(raw, index, *names):
    for name in names:
        pos = index.get(name)
        if pos is None or pos >= len(raw):
            continue
        text = cell_text(raw[pos])
        if text != "":
            return text
    return ""


def _texts(raw):
    return [cell_text(value) for value in raw]


def _blank(raw):
    return not any(cell_text(value) for value in raw)


def _is_total(raw, index):
    for name in ("物料编号", "中文名称", "采购序号", "序号", "采购计划单号", "计划单号"):
        text = _pick(raw, index, name)
        if text in ("合计", "小计", "总计") or text.endswith("合计"):
            return True
    return False


def contract_item_template(direction: str) -> tuple[bytes, str]:
    purchase = direction != "sale"
    headers = ["物料编号", "中文名称", "英文名称", "规格", "单位", "数量", "含税单价", "备注"]
    if purchase:
        headers.extend(["计划单号", "使用设备"])
    title = "采购合同物料导入模板.xlsx" if purchase else "销售合同物料导入模板.xlsx"
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "物料明细"
    bold = Font(bold=True)
    for col, text in enumerate(headers, start=1):
        cell = sheet.cell(1, col, text)
        cell.font = bold
        sheet.column_dimensions[cell.column_letter].width = 16
    buf = BytesIO()
    book.save(buf)
    return buf.getvalue(), title


def parse_contract_items(db, company_id, content: bytes, filename: str, direction: str, existing_codes, contract_id):
    if direction not in ("sale", "purchase"):
        raise BizError(MSG_DIRECTION)
    if not content or len(content) > MAX_SIZE:
        raise BizError("文件大小超过限制")
    name = (filename or "").lower()
    if not (name.endswith(".xlsx") or name.endswith(".xls")):
        raise BizError("文件格式不正确")
    matrix = read_matrix(content, filename)
    header_at = _find_row(matrix, ["物料编号"])
    if header_at is None:
        for required in (["规格型号", "采购数量"], ["规格型号"], ["规格", "数量"], ["规格"]):
            header_at = _find_row(matrix, required)
            if header_at is not None:
                break
    if header_at is None:
        raise BizError(PARSE_FAIL)
    index = _header_index(matrix[header_at])
    purchase = direction == "purchase"
    indexes = _plan_indexes(q(db, PurchasePlan, company_id).all()) if purchase else None
    seen = set(existing_codes or [])
    items = []
    errors = []
    for raw in matrix[header_at + 1 :]:
        if _blank(raw) or _is_total(raw, index):
            continue
        code = _pick(raw, index, "物料编号", "物料编码", "物料代码")
        spec = _pick(raw, index, "规格", "规格型号")
        name = _pick(raw, index, "中文名称")
        plan = _match_plan(indexes, code, spec, name) if purchase else None
        label = code or spec or name
        if plan is not None:
            code = plan.material_code
        elif not code:
            code = spec or name
        if not code:
            errors.append("存在没有物料编号的行")
            continue
        if code in seen:
            errors.append(f"物料 {label or code} 重复")
            continue
        seen.add(code)
        qty_text = _first_positive(raw, index, "采购数量", "数量")
        price_text = _pick(raw, index, "含税单价", "单价(含税)", "单价") or "0"
        try:
            qty = D(qty_text)
            price = D(price_text)
        except BizError as exc:
            errors.append(f"物料 {label or code}：{exc.message}")
            continue
        if qty <= 0:
            errors.append(f"物料 {label or code} 数量必须大于 0")
            continue
        if price < 0:
            errors.append(f"物料 {label or code} 含税单价不能小于 0")
            continue
        if plan:
            bought = calc.purchased_qty(db, company_id, plan.material_code, exclude_contract_id=contract_id)
            remain = D(plan.qty) - bought
            if qty > remain:
                errors.append(f"物料 {label or code} 超出 {qty_str(qty - remain)}")
                continue
        item = {
            "material_code": code,
            "name_cn": name,
            "name_en": _pick(raw, index, "英文名称"),
            "spec": spec,
            "unit": _pick(raw, index, "单位"),
            "qty": qty_str(qty),
            "price": money_str(round2(price)),
            "remark": _pick(raw, index, "备注"),
            "plan_no": _pick(raw, index, "计划单号", "采购计划单号"),
            "category": _pick(raw, index, "使用设备", "类别"),
            "matched_plan": False,
        }
        if plan is not None:
            _apply_plan_match(item, plan)
        items.append(item)
    if errors:
        raise BizError("；".join(dict.fromkeys(errors)))
    if not items:
        raise BizError("没有可导入的物料")
    return items


def _check_module(module, direction, party_type):
    if module not in MODULES:
        raise BizError(PARSE_FAIL)
    if module in ("contract", "order") and direction not in ("sale", "purchase"):
        raise BizError(MSG_DIRECTION)
    if module == "invoice" and direction not in ("output", "input"):
        raise BizError(MSG_DIRECTION)
    if module == "settle" and direction not in ("receipt", "payment"):
        raise BizError(MSG_DIRECTION)
    if module == "party" and party_type not in ("supplier", "customer"):
        raise BizError("类型不正确")


def _normalize_mode(module, mode):
    mode = (mode or "").strip()
    if module == "contract":
        if mode not in ("overwrite", "append"):
            raise BizError("导入模式不正确")
        return mode
    if module == "purchase-plan":
        if mode not in ("update", "replace", "append"):
            raise BizError("导入模式不正确")
        return mode
    if module in ("party", "goods", "customs-material-fill"):
        return "update"
    return "create"


def _mode_key(mode, direction, party_type):
    return "|".join([mode, direction or "", party_type or ""])
