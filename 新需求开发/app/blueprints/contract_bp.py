from io import BytesIO

from flask import (
    Blueprint,
    abort,
    current_app,
    flash,
    jsonify,
    redirect,
    render_template,
    request,
    send_file,
    url_for,
)
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from app.auth import can_manual, can_mutate, can_write, current_user, login_required
from app.db import begin_immediate, get_db, log_op, safe_rollback
from app.importer import handle_import_request, import_page, template_bytes
from app.numbering import (
    CONTRACT_POOL,
    allocate_number,
    matches_any_pattern,
    number_exists,
    occupy_slot,
    parse_pool_number,
    preview_number,
    release_number,
    slot_blocked,
)
from app.utils import (
    DEFAULT_PARTY_A,
    build_pager,
    clean_text,
    default_number_year,
    excel_text,
    sort_sql,
    form_dict,
    like,
    now_str,
    page_args,
    parse_date,
    parse_decimal,
    year_options,
)

bp = Blueprint("contracts", __name__, url_prefix="/contracts")


CONTRACT_SORT = {
    "contract_no": "contract_no",
    "business_type": "business_type",
    "number_year": "number_year",
    "product_name": "product_name",
    "party_b": "party_b",
    "party_a": "party_a",
    "amount_cny": "CAST(amount_cny AS REAL)",
    "sign_date": "sign_date",
}
CONTRACT_RECENT = "number_year DESC, IFNULL(num_seq, 0) DESC, created_at DESC, id DESC"


def _filters():
    args = request.args
    where = ["deleted_at IS NULL"]
    params = []
    keyword = (args.get("keyword") or "").strip()
    if keyword:
        token = like(keyword)
        where.append(
            "(contract_no LIKE ? ESCAPE '\\' OR product_name LIKE ? ESCAPE '\\' OR party_b LIKE ? ESCAPE '\\' OR party_a LIKE ? ESCAPE '\\')"
        )
        params.extend([token, token, token, token])
    mapping = {
        "contract_no": "contract_no LIKE ? ESCAPE '\\'",
        "product_name": "product_name LIKE ? ESCAPE '\\'",
        "party_b": "party_b LIKE ? ESCAPE '\\'",
        "party_a": "party_a LIKE ? ESCAPE '\\'",
    }
    for key, sql in mapping.items():
        value = (args.get(key) or "").strip()
        if value:
            where.append(sql)
            params.append(like(value))
    if args.get("business_type") in ("锰矿", "铝矿"):
        where.append("business_type=?")
        params.append(args.get("business_type"))
    year = (args.get("number_year") or "").strip()
    if year.isdigit():
        where.append("number_year=?")
        params.append(int(year))
    sign_from = (args.get("sign_from") or "").strip()
    sign_to = (args.get("sign_to") or "").strip()
    if sign_from:
        where.append("sign_date>=?")
        params.append(sign_from)
    if sign_to:
        where.append("sign_date<=?")
        params.append(sign_to)
    return " AND ".join(where), params


def _blank_form():
    return {
        "business_type": "锰矿",
        "number_year": str(default_number_year()),
        "product_name": "",
        "party_b": "",
        "party_a": DEFAULT_PARTY_A,
        "amount_cny": "",
        "sign_date": "",
    }


def _validate(form):
    business_type = form.get("business_type") or ""
    if business_type not in CONTRACT_POOL:
        raise ValueError("请选择业务类型")
    year_text = (form.get("number_year") or "").strip()
    if not year_text.isdigit():
        raise ValueError("请选择编号年份")
    year = int(year_text)
    if year < 2000 or year > 2099:
        raise ValueError("编号年份需在 2000 到 2099 之间")
    return {
        "business_type": business_type,
        "number_year": year,
        "product_name": clean_text(form.get("product_name"), 200, "产品/合同名称", required=True),
        "party_b": clean_text(form.get("party_b"), 200, "乙方", required=True),
        "party_a": clean_text(form.get("party_a"), 200, "甲方", required=True),
        "amount_cny": parse_decimal(form.get("amount_cny"), 2, "合同金额", required=True),
        "sign_date": parse_date(form.get("sign_date"), "签订时间", required=True),
    }


def _form_from_row(row):
    return {
        "business_type": row["business_type"],
        "number_year": str(row["number_year"]),
        "product_name": row["product_name"],
        "party_b": row["party_b"],
        "party_a": row["party_a"],
        "amount_cny": row["amount_cny"],
        "sign_date": row["sign_date"],
    }


@bp.route("")
@login_required
def index():
    page, size = page_args()
    where, params = _filters()
    db = get_db()
    total = db.execute(f"SELECT COUNT(*) AS n FROM contracts WHERE {where}", params).fetchone()["n"]
    pager = build_pager(total, page, size)
    rows = db.execute(
        f"""SELECT * FROM contracts WHERE {where}
            ORDER BY {sort_sql(CONTRACT_SORT, CONTRACT_RECENT)} LIMIT ? OFFSET ?""",
        [*params, size, pager["offset"]],
    ).fetchall()
    return render_template(
        "contracts.html",
        rows=rows,
        pager=pager,
        filters=request.args,
        years=year_options(),
    )


@bp.route("/new", methods=["GET", "POST"])
@login_required
def new():
    user = current_user()
    if not can_write(user):
        abort(403)
    form = _blank_form()
    errors = []
    if request.method == "POST":
        form.update(form_dict())
        db = get_db()
        try:
            data = _validate(form)
            begin_immediate(db)
            pool = CONTRACT_POOL[data["business_type"]]
            got = allocate_number(db, pool, data["number_year"])
            ts = now_str()
            cur = db.execute(
                """INSERT INTO contracts
                   (contract_no, business_type, number_year, num_pool, num_seq, product_name,
                    party_b, party_a, amount_cny, sign_date, manual_edited, created_by,
                    updated_by, created_at, updated_at)
                   VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 0, ?, ?, ?, ?)""",
                (
                    got["number"],
                    data["business_type"],
                    data["number_year"],
                    got["pool"],
                    got["seq"],
                    data["product_name"],
                    data["party_b"],
                    data["party_a"],
                    data["amount_cny"],
                    data["sign_date"],
                    user["id"],
                    user["id"],
                    ts,
                    ts,
                ),
            )
            log_op(
                db,
                user,
                "contract_create",
                f"新增合同 {got['number']}" + ("，已复用空闲编号" if got["reused"] else ""),
                "contract",
                cur.lastrowid,
                {
                    "业务类型": data["business_type"],
                    "编号规则": "锰矿合同" if pool == "contract_mn" else "铝矿合同",
                    "编号年份": data["number_year"],
                    "生成编号": got["number"],
                    "操作人": user["display_name"],
                    "生成时间": ts,
                    "复用空闲编号": got["reused"],
                },
            )
            db.commit()
            message = f"已生成合同编号 {got['number']}"
            if got["reused"]:
                message += "，已复用空闲编号"
            flash(message, "success")
            return redirect(url_for("contracts.index"))
        except ValueError as exc:
            safe_rollback(db)
            errors = [str(exc)]
        except Exception:
            safe_rollback(db)
            current_app_logger().exception("create contract failed")
            errors = ["保存失败，请重试"]
    return render_template(
        "contract_form.html",
        form=form,
        errors=errors,
        editing=False,
        years=year_options(),
        record=None,
    )


@bp.route("/<int:cid>/edit", methods=["GET", "POST"])
@login_required
def edit(cid):
    user = current_user()
    db = get_db()
    row = _alive(db, cid)
    if not can_mutate(user, row):
        abort(403)
    form = _form_from_row(row)
    errors = []
    if request.method == "POST":
        posted = form_dict()
        form["product_name"] = posted.get("product_name", "")
        form["party_b"] = posted.get("party_b", "")
        form["party_a"] = posted.get("party_a", "")
        form["amount_cny"] = posted.get("amount_cny", "")
        form["sign_date"] = posted.get("sign_date", "")
        try:
            data = _validate({**form, "business_type": row["business_type"], "number_year": str(row["number_year"])})
            begin_immediate(db)
            ts = now_str()
            db.execute(
                """UPDATE contracts
                   SET product_name=?, party_b=?, party_a=?, amount_cny=?, sign_date=?,
                       updated_by=?, updated_at=?
                   WHERE id=? AND deleted_at IS NULL""",
                (
                    data["product_name"],
                    data["party_b"],
                    data["party_a"],
                    data["amount_cny"],
                    data["sign_date"],
                    user["id"],
                    ts,
                    cid,
                ),
            )
            log_op(db, user, "contract_update", f"编辑合同 {row['contract_no']}", "contract", cid)
            db.commit()
            flash("合同已保存", "success")
            return redirect(url_for("contracts.index"))
        except ValueError as exc:
            safe_rollback(db)
            errors = [str(exc)]
        except Exception:
            safe_rollback(db)
            current_app_logger().exception("edit contract failed")
            errors = ["保存失败，请重试"]
    return render_template(
        "contract_form.html",
        form=form,
        errors=errors,
        editing=True,
        years=year_options(),
        record=row,
    )


@bp.route("/<int:cid>/manual", methods=["GET", "POST"])
@login_required
def manual(cid):
    user = current_user()
    if not can_manual(user):
        abort(403)
    db = get_db()
    row = _alive(db, cid)
    errors = []
    history = db.execute(
        """SELECT * FROM number_edit_logs
           WHERE target_type='contract' AND target_id=?
           ORDER BY id DESC""",
        (cid,),
    ).fetchall()
    if request.method == "POST":
        new_no = (request.form.get("new_no") or "").strip()
        reason = (request.form.get("reason") or "").strip()
        try:
            begin_immediate(db)
            fresh = _alive(db, cid)
            _apply_manual(db, fresh, new_no, reason, user)
            db.commit()
            flash(f"合同编号已修改为 {new_no}", "success")
            return redirect(url_for("contracts.index"))
        except ValueError as exc:
            safe_rollback(db)
            errors = [str(exc)]
        except Exception:
            safe_rollback(db)
            current_app_logger().exception("manual contract failed")
            errors = ["保存失败，请重试"]
    return render_template(
        "manual_form.html",
        title="手动编辑合同编号",
        old_no=row["contract_no"],
        errors=errors,
        history=history,
        back_url=url_for("contracts.edit", cid=cid),
    )


@bp.route("/<int:cid>/delete", methods=["POST"])
@login_required
def delete(cid):
    user = current_user()
    db = get_db()
    try:
        begin_immediate(db)
        row = _alive(db, cid)
        if not can_mutate(user, row):
            safe_rollback(db)
            flash("只能删除自己登记的合同", "error")
            return redirect(url_for("contracts.index"))
        _soft_delete(db, row, user)
        db.commit()
        flash(f"已删除合同 {row['contract_no']}，编号已进入空闲池", "success")
    except Exception:
        safe_rollback(db)
        current_app_logger().exception("delete contract failed")
        flash("删除失败，请重试", "error")
    return redirect(url_for("contracts.index"))


@bp.route("/batch-delete", methods=["POST"])
@login_required
def batch_delete():
    user = current_user()
    ids = []
    for item in request.form.getlist("ids"):
        if str(item).isdigit():
            ids.append(int(item))
    if not ids:
        flash("请先选择合同", "error")
        return redirect(url_for("contracts.index"))
    db = get_db()
    done = 0
    skipped = 0
    try:
        begin_immediate(db)
        for cid in ids[:200]:
            row = db.execute(
                "SELECT * FROM contracts WHERE id=? AND deleted_at IS NULL",
                (cid,),
            ).fetchone()
            if not row or not can_mutate(user, row):
                skipped += 1
                continue
            _soft_delete(db, row, user)
            done += 1
        db.commit()
    except Exception:
        safe_rollback(db)
        current_app_logger().exception("batch delete contract failed")
        flash("删除失败，请重试", "error")
        return redirect(url_for("contracts.index"))
    message = f"已删除 {done} 条合同，编号已进入空闲池"
    if skipped:
        message += f"，跳过 {skipped} 条（无权限或已不存在）"
    flash(message, "success" if done else "error")
    return redirect(url_for("contracts.index"))


@bp.route("/export")
@login_required
def export_xlsx():
    where, params = _filters()
    db = get_db()
    rows = db.execute(
        f"SELECT * FROM contracts WHERE {where} ORDER BY {sort_sql(CONTRACT_SORT, CONTRACT_RECENT)} LIMIT 50000",
        params,
    ).fetchall()
    wb = Workbook()
    ws = wb.active
    ws.title = "合同登记"
    headers = ["序号", "合同编号", "业务类型", "编号年份", "产品/合同名称", "乙方", "甲方", "合同金额（CNY）", "签订时间"]
    ws.append(headers)
    for index, row in enumerate(rows, start=1):
        ws.append(
            [
                index,
                excel_text(row["contract_no"]),
                row["business_type"],
                row["number_year"],
                excel_text(row["product_name"]),
                excel_text(row["party_b"]),
                excel_text(row["party_a"]),
                float(row["amount_cny"]),
                row["sign_date"],
            ]
        )
    _style_sheet(ws, money_cols={8})
    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    from datetime import date

    filename = f"合同登记_{date.today().strftime('%Y%m%d')}.xlsx"
    return send_file(
        bio,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@bp.route("/import", methods=["GET", "POST"])
@login_required
def import_xlsx():
    user = current_user()
    if not can_write(user):
        abort(403)
    result = None
    if request.method == "POST":
        result, message, category = handle_import_request(
            "contract", get_db(), user, current_app.testing
        )
        flash(message, category)
    return render_template("import.html", result=result, **import_page("contract"))


@bp.route("/import-template")
@login_required
def import_template():
    user = current_user()
    if not can_write(user):
        abort(403)
    payload, filename = template_bytes("contract")
    return send_file(
        BytesIO(payload),
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@bp.route("/api/preview")
@login_required
def preview():
    business_type = request.args.get("business_type") or ""
    year = request.args.get("year") or ""
    if business_type not in CONTRACT_POOL or not str(year).isdigit():
        return jsonify({"ok": False, "message": "请选择业务类型和编号年份"})
    year = int(year)
    if year < 2000 or year > 2099:
        return jsonify({"ok": False, "message": "编号年份无效"})
    try:
        got = preview_number(get_db(), CONTRACT_POOL[business_type], year)
    except ValueError as exc:
        return jsonify({"ok": False, "message": str(exc)})
    return jsonify({"ok": True, "number": got["number"], "reused": got["reused"]})


def _alive(db, cid):
    row = db.execute(
        "SELECT * FROM contracts WHERE id=? AND deleted_at IS NULL",
        (cid,),
    ).fetchone()
    if not row:
        abort(404)
    return row


def _soft_delete(db, row, user):
    if row["num_pool"] and row["num_seq"] and not slot_blocked(
        db, row["num_pool"], row["num_seq"], self_contract_id=row["id"]
    ):
        release_number(db, row["num_pool"], row["num_seq"])
    ts = now_str()
    db.execute(
        "UPDATE contracts SET deleted_at=?, deleted_by=?, updated_at=? WHERE id=?",
        (ts, user["id"], ts, row["id"]),
    )
    log_op(
        db,
        user,
        "contract_delete",
        f"删除合同 {row['contract_no']}，编号回收",
        "contract",
        row["id"],
        {"编号": row["contract_no"], "编号池": row["num_pool"], "顺序号": row["num_seq"]},
    )


def _apply_manual(db, row, new_no, reason, user):
    if not new_no or len(new_no) > 64 or any(ch.isspace() for ch in new_no):
        raise ValueError("新编号不能为空、不能含空白，且不超过 64 个字符")
    if len(reason) < 2 or len(reason) > 200:
        raise ValueError("请填写修改原因（2 到 200 个字）")
    if new_no == row["contract_no"]:
        raise ValueError("新编号与原编号相同")
    if number_exists(db, new_no, except_contract_id=row["id"]):
        raise ValueError("编号已存在，请更换")
    pool = row["num_pool"] or CONTRACT_POOL[row["business_type"]]
    parsed = parse_pool_number(pool, new_no)
    if parsed is None and matches_any_pattern(new_no):
        raise ValueError("编号格式与当前业务类型的编号规则不一致")
    old_pool, old_seq = row["num_pool"], row["num_seq"]
    new_year = row["number_year"]
    new_pool = None
    new_seq = None
    if parsed:
        new_year, new_seq = parsed
        new_pool = pool
        if slot_blocked(db, pool, new_seq, self_contract_id=row["id"]):
            raise ValueError("该顺序号已被其他记录占用")
    ts = now_str()
    db.execute(
        """UPDATE contracts
           SET contract_no=?, number_year=?, num_pool=?, num_seq=?, manual_edited=1,
               updated_by=?, updated_at=?
           WHERE id=?""",
        (new_no, new_year, new_pool, new_seq, user["id"], ts, row["id"]),
    )
    if old_pool and old_seq and not (old_pool == new_pool and old_seq == new_seq):
        if not slot_blocked(db, old_pool, old_seq, self_contract_id=row["id"]):
            release_number(db, old_pool, old_seq)
    if parsed:
        occupy_slot(db, new_pool, new_seq, new_year, new_no)
    db.execute(
        """INSERT INTO number_edit_logs
           (target_type, target_id, old_no, new_no, reason, operator_id, operator_name, created_at)
           VALUES ('contract', ?, ?, ?, ?, ?, ?, ?)""",
        (row["id"], row["contract_no"], new_no, reason, user["id"], user["display_name"], ts),
    )
    log_op(
        db,
        user,
        "number_manual_edit",
        f"手动编辑合同编号 {row['contract_no']} -> {new_no}",
        "contract",
        row["id"],
        {"原编号": row["contract_no"], "新编号": new_no, "原因": reason, "操作人": user["display_name"], "操作时间": ts},
    )


def _style_sheet(ws, money_cols):
    fill = PatternFill("solid", fgColor="E8F3FF")
    font = Font(bold=True, color="1D2129")
    thin = Border(
        left=Side(style="thin", color="E5E6EB"),
        right=Side(style="thin", color="E5E6EB"),
        top=Side(style="thin", color="E5E6EB"),
        bottom=Side(style="thin", color="E5E6EB"),
    )
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = thin
            cell.alignment = Alignment(vertical="center")
            if cell.column in money_cols and isinstance(cell.value, float):
                cell.number_format = "#,##0.00"
    for column in ws.columns:
        letter = column[0].column_letter
        width = 14
        for cell in column:
            if cell.value:
                width = max(width, min(len(str(cell.value)) + 4, 40))
        ws.column_dimensions[letter].width = width
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22


def current_app_logger():
    import logging

    return logging.getLogger("invoiceapp")
