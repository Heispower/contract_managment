import logging
import re
from datetime import date, datetime, timedelta
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO

from flask import request, url_for
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.db import begin_immediate, log_op, safe_rollback
from app.utils import refresh_batch_amounts
from app.numbering import (
    format_sub_no,
    matches_any_pattern,
    occupy_slot,
    parse_pool_number,
    slot_blocked,
)
from app.utils import DEFAULT_PARTY_A, now_str, parse_decimal

SUB_RE = re.compile(r"^(.+)-(\d{3})$")
CONTRACT_CODE_RE = re.compile(r"^[A-Za-z0-9][A-Za-z0-9\-]{3,63}$")
CONTRACT_EXTRACT_RE = re.compile(r"[A-Za-z]{2,}(?:-[A-Za-z0-9]+){1,8}")
INVOICE_CODE_RE = re.compile(r"^[A-Za-z0-9\u4e00-\u9fff（）()\-]{4,80}$")
YEAR_IN_NUMBER_RE = re.compile(r"(20\d{2})")
SHORT_YEAR_RE = re.compile(r"-(\d{2})(\d{3,})[A-Za-z]?(?:-\d{3})?$")

CONTRACT_NO_ALIASES = ("合同编号", "合同号")
INVOICE_NO_ALIASES = ("发票号", "发票编号")
PLACEHOLDER = "（历史未填）"

CONTRACT_FIELDS = {
    "contract_no": CONTRACT_NO_ALIASES,
    "business_type": ("业务类型",),
    "number_year": ("编号年份",),
    "product_name": ("产品/合同名称", "产品名称", "合同名称"),
    "party_b": ("乙方", "供应商名称", "供应商"),
    "party_a": ("甲方",),
    "amount": ("合同金额（CNY）", "合同金额", "合同总额"),
    "sign_date": ("签订时间", "签订日期", "合同日期"),
}
INVOICE_FIELDS = {
    "invoice_no": INVOICE_NO_ALIASES,
    "contract_no": ("关联合同", "合同编号", "合同号"),
    "business_type": ("业务类型",),
    "number_rule": ("编号规则",),
    "number_year": ("编号年份",),
    "content": ("所含内容", "内容"),
    "bl_date": ("提单日期",),
    "price_term": ("价格条款",),
    "currency": ("结算币种", "币种"),
    "amount_usd": ("美元金额", "USD金额"),
    "amount_cny": ("人民币金额", "CNY金额"),
    "rate": ("汇率",),
    "exchange_month": ("汇率月份",),
    "price_combined": ("发票价格", "发票金额"),
    "po_no": ("PO号", "PO"),
    "operator_name": ("操作人",),
    "transport_type": ("运输方式",),
    "email_sent_date": ("发邮件时间", "邮件时间"),
    "remark": ("备注",),
    "bl_no": ("提单号",),
}

CONTRACT_TIPS = (
    "可下载模板填写，也可直接上传购销统计表。程序在每个工作表前 30 行查找“合同号”或“合同编号”，没有该表头的明细表不会导入。",
    "历史编号按原样保存。BS-GMI、BIC-GMI、KWI-GMI、BMGG-MAZ-供应商缩写这类旧号只登记，不占用 BMI-GMI / BMGG-MAZ 新号池。",
    "已经符合现行规则的编号会占用对应顺序号，中间空号仍可被下一笔新合同回收。",
    "未填签订时间时，按编号中的年份写成该年 1 月 1 日；未填金额时按 0 保存。同一编号已存在则跳过，不覆盖。",
)
INVOICE_TIPS = (
    "可下载模板填写，也可直接上传《发票编号登记表》（表头含发票号、所含内容）。",
    "以 -001、-002 这种三位分单号结尾的发票，会归到同一批次，前缀保持原编号，不重新生成。",
    "旧格式发票号原样保存，不占用铝矿海运、铝矿空运、锰矿发票的新号池。符合现行规则的编号会占用顺序号。",
    "价格只有一个数字时记入人民币，美元为 0。已存在的发票号会跳过。若要关联合同，请先导入合同。",
)


class ImportResult:
    def __init__(self):
        self.imported = 0
        self.skipped = 0
        self.occupied = 0
        self.legacy = 0
        self.filled_dates = 0
        self.zero_amounts = 0
        self.batch_count = 0
        self.clipped = 0
        self.errors = []
        self.hidden_errors = 0
        self.warnings = []
        self.hidden_warnings = 0
        self.ignored_sheets = []

    def add_error(self, where, message):
        if where != "文件":
            self.skipped += 1
        if len(self.errors) < 80:
            self.errors.append({"where": where, "message": message})
        else:
            self.hidden_errors += 1

    def add_warning(self, where, message):
        if len(self.warnings) < 40:
            self.warnings.append({"where": where, "message": message})
        else:
            self.hidden_warnings += 1

    def summary(self):
        if self.imported == 0 and self.skipped == 0 and self.errors:
            return self.errors[0]["message"]
        text = f"成功导入 {self.imported} 条，跳过 {self.skipped} 条"
        extra = []
        if self.occupied:
            extra.append(f"{self.occupied} 条已占用现行顺序号")
        if self.legacy:
            extra.append(f"{self.legacy} 条历史编号未占新号池")
        if self.filled_dates:
            extra.append(f"{self.filled_dates} 条未填签订时间，已按编号年份补为 1 月 1 日")
        if self.zero_amounts:
            extra.append(f"{self.zero_amounts} 条缺少金额，已按 0 保存")
        if self.batch_count:
            extra.append(f"归入 {self.batch_count} 个批次")
        if self.clipped:
            extra.append(f"{self.clipped} 条超长文本已截断")
        if extra:
            text += "（" + "，".join(extra) + "）"
        return text


def import_page(kind):
    if kind == "contract":
        return {
            "title": "批量导入合同",
            "intro": "把已有合同登记进系统。编号保持原样，已存在的编号自动跳过。",
            "tips": CONTRACT_TIPS,
            "back_url": url_for("contracts.index"),
            "template_url": url_for("contracts.import_template"),
        }
    return {
        "title": "批量导入发票",
        "intro": "把已有发票登记进系统。历史发票号保持原样，分单号会归入原批次。",
        "tips": INVOICE_TIPS,
        "back_url": url_for("invoices.index"),
        "template_url": url_for("invoices.import_template"),
    }


def perform_import(kind, db, user, storage):
    if storage is None or not (storage.filename or "").strip():
        raise ValueError("请选择要导入的 xlsx 文件")
    if not storage.filename.lower().endswith(".xlsx"):
        raise ValueError("只支持 .xlsx 文件")
    data = storage.read()
    if not data:
        raise ValueError("文件是空的")
    if kind == "contract":
        return import_contracts(db, user, data)
    return import_invoices(db, user, data)


def handle_import_request(kind, db, user, testing=False):
    try:
        result = perform_import(kind, db, user, request.files.get("file"))
    except ValueError as exc:
        return None, str(exc), "error"
    except Exception as exc:
        safe_rollback(db)
        logging.getLogger("invoiceapp").exception("import failed")
        if testing:
            return None, f"导入失败：{exc}", "error"
        return None, "导入失败，数据未写入。请检查文件后重试。", "error"
    category = "success" if result.imported else "error"
    return result, result.summary(), category


def template_bytes(kind):
    wb = Workbook()
    ws = wb.active
    if kind == "contract":
        ws.title = "合同"
        headers = ["合同编号", "业务类型", "编号年份", "产品/合同名称", "乙方", "甲方", "合同金额（CNY）", "签订时间"]
        notes = [
            "请在“合同”工作表从第 2 行填写，不要修改第 1 行表头。合同编号必填，按原样保存。",
            "业务类型填锰矿或铝矿。编号年份可空，空着时从编号里识别。签订时间格式 YYYY-MM-DD，可空。",
            "符合 BMI-GMI-YYYYXXX 或 BMGG-MAZ-YYXXX 的编号会占用顺序号；旧格式编号只登记。",
            "也可以不使用本模板，直接上传表头含“合同号”的历史统计表。",
        ]
        filename_hint = "合同"
    else:
        ws.title = "发票"
        headers = [
            "发票号", "关联合同", "业务类型", "编号规则", "编号年份", "所含内容", "提单日期",
            "价格条款", "结算币种", "美元金额", "人民币金额", "汇率", "汇率月份", "PO",
            "操作人", "运输方式", "发邮件时间", "备注", "提单号",
        ]
        notes = [
            "请在“发票”工作表从第 2 行填写，不要修改第 1 行表头。发票号必填，按原样保存。",
            "编号规则可填铝矿海运、铝矿空运或锰矿。运输方式填海运或空运。价格条款填 CIF 或 FOB，不填按 FOB。结算币种填 CNY 或 USD。",
            "以 -001 这种三位分单号结尾时，归入该前缀的批次。",
            "也可以直接上传《发票编号登记表》，表头含发票号、所含内容即可。",
        ]
        filename_hint = "发票"
    ws.append(headers)
    fill = PatternFill("solid", fgColor="E8F3FF")
    font = Font(bold=True, color="1D2129")
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center")
    for column in ws.columns:
        letter = column[0].column_letter
        width = min(max(14, len(str(column[0].value or "")) + 4), 36)
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    info = wb.create_sheet("填写说明")
    for line in notes:
        info.append([line])
    info.column_dimensions["A"].width = 92
    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio.getvalue(), f"{filename_hint}导入模板.xlsx"


def import_contracts(db, user, data):
    sheets = load_sheets(data)
    result = ImportResult()
    rows, used = collect_contracts(sheets, result)
    result.ignored_sheets = [name for name, _grid in sheets if name not in used and name != "填写说明"]
    if not rows and result.skipped == 0:
        result.add_error("文件", "没有找到可导入的合同。请使用模板，或上传表头含“合同号”的统计表。")
        return result
    if rows:
        _write_contracts(db, user, rows, result)
    return result


def import_invoices(db, user, data):
    sheets = load_sheets(data)
    result = ImportResult()
    rows, used = collect_invoices(sheets, result)
    result.ignored_sheets = [name for name, _grid in sheets if name not in used and name != "填写说明"]
    if not rows and result.skipped == 0:
        result.add_error("文件", "没有找到可导入的发票。请使用模板，或上传表头含“发票号”的登记表。")
        return result
    if rows:
        _write_invoices(db, user, rows, result)
    return result


def load_sheets(data):
    try:
        wb = load_workbook(BytesIO(data), data_only=True)
    except Exception as exc:
        raise ValueError("无法读取 Excel，请确认文件是 xlsx 且未加密") from exc
    sheets = []
    try:
        for ws in wb.worksheets:
            grid = []
            for index, row in enumerate(ws.iter_rows(values_only=True)):
                if index >= 20000:
                    break
                grid.append(tuple(row))
            sheets.append((ws.title, grid))
    finally:
        wb.close()
    return sheets


def collect_contracts(sheets, result):
    rows = []
    used = set()
    for name, grid in sheets:
        found = find_header(grid, CONTRACT_FIELDS, "contract_no", ("product_name", "party_b", "amount"))
        if not found:
            continue
        used.add(name)
        header_idx, cols = found
        for excel_row, raw in data_rows(grid, header_idx):
            item, error = parse_contract_row(raw, cols)
            if error == "blank":
                continue
            if error:
                result.add_error(f"{name} 第{excel_row}行", error)
                continue
            item["where"] = f"{name} 第{excel_row}行"
            rows.append(item)
    return rows, used


def collect_invoices(sheets, result):
    rows = []
    used = set()
    for name, grid in sheets:
        found = find_header(
            grid,
            INVOICE_FIELDS,
            "invoice_no",
            ("content", "price_combined", "amount_cny", "amount_usd", "bl_date", "transport_type"),
        )
        if not found:
            continue
        used.add(name)
        header_idx, cols = found
        for excel_row, raw in data_rows(grid, header_idx):
            pieces = split_invoice_numbers(cell(raw, cols.get("invoice_no")))
            if len(pieces) <= 1:
                targets = [(raw, f"{name} 第{excel_row}行")]
            else:
                targets = []
                for index, number in enumerate(pieces, start=1):
                    copied = list(raw)
                    copied[cols["invoice_no"]] = number
                    targets.append((tuple(copied), f"{name} 第{excel_row}行第{index}个编号"))
            for target, where in targets:
                item, error = parse_invoice_row(target, cols)
                if error == "blank":
                    continue
                if error:
                    result.add_error(where, error)
                    continue
                item["where"] = where
                rows.append(item)
    mark_batches(rows)
    return rows, used


def find_header(grid, fields, key, required_any):
    for index, raw in enumerate(grid[:30]):
        headers = [norm_header(cell) for cell in raw]
        if col_index(headers, fields[key]) is None:
            continue
        if not any(col_index(headers, fields[name]) is not None for name in required_any):
            continue
        return index, {name: col_index(headers, aliases) for name, aliases in fields.items()}
    return None


def data_rows(grid, header_idx):
    for offset, raw in enumerate(grid[header_idx + 1 :], start=header_idx + 2):
        yield offset, raw


def parse_contract_row(raw, cols):
    number_cell = cell(raw, cols.get("contract_no"))
    number_text = text_of(number_cell).replace(" ", "").replace("\u3000", "")
    if number_text in ("", "合同号", "合同编号", "合计", "小计", "总计", "备注"):
        if not number_text:
            return None, "blank"
        return None, "blank"
    number = normalize_contract_no(number_text)
    if not number:
        if not number_text:
            return None, "blank"
        return None, "合同号无法识别"
    pool, parsed = match_contract_pool(number)
    if pool is None and matches_any_pattern(number):
        return None, "该编号属于发票规则，不能作为合同导入"
    explicit = norm_business(cell(raw, cols.get("business_type")))
    if pool == "contract_mn":
        business = "锰矿"
        year, seq = parsed
    elif pool == "contract_al":
        business = "铝矿"
        year, seq = parsed
    else:
        business = explicit or infer_contract_business(number)
        year = norm_year(cell(raw, cols.get("number_year"))) or infer_year(number)
        seq = None
    product, clipped_product = clip(text_of(cell(raw, cols.get("product_name"))), 200)
    party_b, clipped_party = clip(text_of(cell(raw, cols.get("party_b"))), 200)
    party_a, clipped_a = clip(text_of(cell(raw, cols.get("party_a"))), 200)
    if not product:
        product = PLACEHOLDER
    if not party_b:
        party_b = PLACEHOLDER
    if not party_a:
        party_a = DEFAULT_PARTY_A
    amount_raw = cell(raw, cols.get("amount"))
    zero_amount = amount_raw is None or text_of(amount_raw) == ""
    try:
        amount = "0.00" if zero_amount else parse_money(amount_raw, "合同金额")
    except ValueError as exc:
        return None, str(exc)
    sign_raw = cell(raw, cols.get("sign_date"))
    filled_date = sign_raw is None or text_of(sign_raw) == ""
    try:
        sign_date = f"{year}-01-01" if filled_date else parse_any_date(sign_raw, "签订时间")
    except ValueError as exc:
        return None, str(exc)
    return {
        "contract_no": number,
        "business_type": business,
        "number_year": year,
        "pool": pool,
        "seq": seq,
        "product_name": product,
        "party_b": party_b,
        "party_a": party_a,
        "amount_cny": amount,
        "sign_date": sign_date,
        "filled_date": filled_date,
        "zero_amount": zero_amount,
        "clipped": clipped_product or clipped_party or clipped_a,
    }, None


def parse_invoice_row(raw, cols):
    number_cell = cell(raw, cols.get("invoice_no"))
    number_text = text_of(number_cell).replace(" ", "").replace("\u3000", "")
    content_text = text_of(cell(raw, cols.get("content")))
    if number_text in ("", "发票号", "发票编号"):
        return None, "blank"
    if not INVOICE_CODE_RE.match(number_text) or not re.search(r"\d", number_text):
        return None, "发票号无法识别"
    number = number_text
    sub = SUB_RE.match(number)
    prefix = None
    sub_seq = None
    if sub and len(sub.group(1)) >= 4:
        prefix = sub.group(1)
        sub_seq = int(sub.group(2))
        if sub_seq < 1:
            return None, "分单号不能为 0"
    pool, parsed = (None, None)
    if sub is None:
        pool, parsed = match_invoice_pool(number)
        if pool is None and matches_any_pattern(number):
            return None, "该编号属于合同规则，不能作为发票导入"
    transport, transport_note = infer_transport(number, cell(raw, cols.get("transport_type")))
    explicit_rule = norm_rule(cell(raw, cols.get("number_rule")))
    explicit_business = norm_business(cell(raw, cols.get("business_type")))
    if pool:
        rule = rule_of_pool(pool)
        business = "锰矿" if rule == "锰矿" else "铝矿"
        year, seq = parsed
        if rule == "铝矿空运":
            transport = "空运"
        elif rule == "铝矿海运":
            transport = "海运"
    else:
        business, rule = legacy_invoice_rule(number, transport, explicit_business, explicit_rule)
        year = norm_year(cell(raw, cols.get("number_year"))) or infer_year(number)
        seq = None
        if explicit_rule in ("铝矿海运", "铝矿空运", "锰矿"):
            rule = explicit_rule
            business = "锰矿" if rule == "锰矿" else "铝矿"
            if rule == "铝矿空运":
                transport = "空运"
            elif rule == "铝矿海运":
                transport = "海运"
    content, clipped_content = clip(content_text or PLACEHOLDER, 500)
    try:
        amounts = invoice_amounts(raw, cols)
    except ValueError as exc:
        return None, str(exc)
    bl_date, bl_note = loose_date(cell(raw, cols.get("bl_date")))
    email_date, email_note = loose_date(cell(raw, cols.get("email_sent_date")))
    if not amounts["exchange_month"]:
        if bl_date:
            amounts["exchange_month"] = bl_date[:7]
        elif email_date:
            amounts["exchange_month"] = email_date[:7]
        else:
            amounts["exchange_month"] = f"{year}-01"
    remark, clipped_remark = clip(text_of(cell(raw, cols.get("remark"))), 500)
    notes = []
    if transport_note and transport_note not in ("海运", "空运"):
        notes.append(f"原运输方式：{transport_note}")
    if bl_note:
        notes.append(f"提单日期：{bl_note}")
    if email_note:
        notes.append(f"发邮件时间：{email_note}")
    if notes:
        extra = "；".join(notes)
        remark = extra if not remark else f"{remark}；{extra}"
        remark, more = clip(remark, 500)
        clipped_remark = clipped_remark or more
    po, clipped_po = clip(text_of(cell(raw, cols.get("po_no"))), 200)
    bl_no, clipped_bl = clip(text_of(cell(raw, cols.get("bl_no"))), 100)
    operator, _clipped_op = clip(text_of(cell(raw, cols.get("operator_name"))), 32)
    price_term = text_of(cell(raw, cols.get("price_term"))).upper()
    if price_term not in ("CIF", "FOB"):
        price_term = amounts["term"] or "FOB"
    return {
        "invoice_no": number,
        "is_sub": prefix is not None,
        "is_master": False,
        "prefix": prefix,
        "sub_seq": sub_seq,
        "business_type": business,
        "number_rule": rule,
        "number_year": year,
        "pool": pool,
        "seq": seq,
        "content": content,
        "bl_date": bl_date,
        "price_term": price_term,
        "currency": amounts["currency"],
        "amount_usd": amounts["usd"],
        "amount_cny": amounts["cny"],
        "exchange_rate": amounts["rate"],
        "exchange_month": amounts["exchange_month"],
        "rate_adjusted": amounts["adjusted"],
        "po_no": po,
        "operator_name": operator,
        "transport_type": transport,
        "email_sent_date": email_date,
        "remark": remark,
        "bl_no": bl_no,
        "contract_no": normalize_contract_no(text_of(cell(raw, cols.get("contract_no")))),
        "zero_amount": amounts["zero"],
        "clipped": clipped_content or clipped_remark or clipped_po or clipped_bl,
    }, None


def mark_batches(rows):
    prefixes = {item["prefix"] for item in rows if item["is_sub"]}
    for item in rows:
        if not item["is_sub"] and item["invoice_no"] in prefixes:
            item["is_master"] = True
            item["prefix"] = item["invoice_no"]


def _write_contracts(db, user, rows, result):
    begin_immediate(db)
    try:
        ts = now_str()
        seen = set()
        for item in rows:
            number = item["contract_no"]
            where = item["where"]
            if number in seen:
                result.add_error(where, "文件中合同号重复")
                continue
            seen.add(number)
            conflict = find_conflict(db, number)
            if conflict:
                result.add_error(where, f"编号已存在于{conflict}")
                continue
            if item["pool"]:
                if slot_blocked(db, item["pool"], item["seq"]):
                    result.add_error(where, f"顺序号 {item['seq']} 已被占用")
                    continue
                occupy_slot(db, item["pool"], item["seq"], item["number_year"], number)
                result.occupied += 1
            else:
                result.legacy += 1
            if item["filled_date"]:
                result.filled_dates += 1
            if item["zero_amount"]:
                result.zero_amounts += 1
            if item["clipped"]:
                result.clipped += 1
            db.execute(
                """INSERT INTO contracts
                   (contract_no, business_type, number_year, num_pool, num_seq, product_name,
                    party_b, party_a, amount_cny, sign_date, manual_edited, created_by,
                    updated_by, created_at, updated_at)
                   VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 0, ?, ?, ?, ?)""",
                (
                    number,
                    item["business_type"],
                    item["number_year"],
                    item["pool"],
                    item["seq"],
                    item["product_name"],
                    item["party_b"],
                    item["party_a"],
                    item["amount_cny"],
                    item["sign_date"],
                    user["id"],
                    user["id"],
                    ts,
                    ts,
                ),
            )
            result.imported += 1
        if result.imported:
            log_op(
                db,
                user,
                "contract_import",
                f"批量导入合同 {result.imported} 条",
                "contract",
                None,
                {"成功": result.imported, "跳过": result.skipped, "占用顺序号": result.occupied},
            )
            db.commit()
        else:
            safe_rollback(db)
            result.occupied = 0
            result.legacy = 0
            result.filled_dates = 0
            result.zero_amounts = 0
            result.clipped = 0
    except Exception:
        safe_rollback(db)
        raise


def _write_invoices(db, user, rows, result):
    begin_immediate(db)
    try:
        ts = now_str()
        prefixes = {}
        for item in rows:
            if not (item["is_sub"] or item["is_master"]):
                continue
            meta = prefixes.setdefault(item["prefix"], {})
            if not meta or item["is_master"]:
                meta["number_rule"] = item["number_rule"]
                meta["number_year"] = item["number_year"]
                meta["transport_type"] = item["transport_type"]
                meta["business_type"] = item["business_type"]
        batch_ids = {}
        created_batches = set()
        batch_errors = {}
        for prefix, meta in prefixes.items():
            try:
                batch_id, created = ensure_batch(db, user, prefix, meta, ts)
                batch_ids[prefix] = batch_id
                if created:
                    created_batches.add(batch_id)
            except ValueError as exc:
                batch_errors[prefix] = str(exc)
        seen = set()
        used_batches = set()
        ordered = [item for item in rows if not item["is_sub"]] + [item for item in rows if item["is_sub"]]
        for item in ordered:
            where = item["where"]
            number = item["invoice_no"]
            prefix = item.get("prefix")
            if prefix and prefix in batch_errors:
                result.add_error(where, batch_errors[prefix])
                continue
            if number in seen:
                result.add_error(where, "文件中发票号重复")
                continue
            seen.add(number)
            allow_batch = batch_ids.get(number)
            conflict = find_conflict(db, number, allow_batch_id=allow_batch)
            if conflict:
                result.add_error(where, f"编号已存在于{conflict}")
                continue
            pool = item["pool"]
            seq = item["seq"]
            if pool:
                if slot_blocked(db, pool, seq):
                    result.add_error(where, f"顺序号 {seq} 已被占用")
                    continue
                occupy_slot(db, pool, seq, item["number_year"], number)
                result.occupied += 1
                if item["is_master"]:
                    db.execute(
                        """UPDATE batches SET num_pool=?, num_seq=?
                           WHERE id=? AND num_pool IS NULL""",
                        (pool, seq, batch_ids[prefix]),
                    )
            else:
                result.legacy += 1
            if item["zero_amount"]:
                result.zero_amounts += 1
            batch_id = None
            batch_prefix = None
            sub_no = None
            sub_seq = None
            is_batch = 0
            is_master = 0
            if item["is_sub"] or item["is_master"]:
                is_batch = 1
                batch_id = batch_ids[prefix]
                batch_prefix = prefix
                is_master = 1 if item["is_master"] else 0
                if item["is_sub"]:
                    sub_seq = item["sub_seq"]
                    sub_no = format_sub_no(sub_seq)
                    try:
                        occupy_sub(db, batch_id, sub_seq, number)
                    except ValueError as exc:
                        result.add_error(where, str(exc))
                        if pool:
                            result.occupied -= 1
                        else:
                            result.legacy -= 1
                        if item["zero_amount"]:
                            result.zero_amounts -= 1
                        continue
            contract_id = lookup_contract(db, item.get("contract_no") or "")
            if item.get("contract_no") and contract_id is None:
                result.add_warning(where, f"关联合同 {item['contract_no']} 不在系统中，发票已导入但未关联")
            operator = item["operator_name"] or user["display_name"]
            db.execute(
                """INSERT INTO invoices
                   (invoice_no, contract_id, business_type, number_rule, number_year, num_pool, num_seq,
                    content, bl_date, price_term, currency, amount_usd, amount_cny, exchange_rate,
                    exchange_month, rate_adjusted, po_no, operator_id, operator_name, transport_type,
                    email_sent_date, remark, bl_no, is_batch, is_batch_master, batch_id, batch_prefix,
                    batch_sub_no, sub_seq, manual_edited, created_by, updated_by, created_at, updated_at)
                   VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 0, ?, ?, ?, ?)""",
                (
                    number,
                    contract_id,
                    item["business_type"],
                    item["number_rule"],
                    item["number_year"],
                    pool,
                    seq,
                    item["content"],
                    item["bl_date"] or None,
                    item["price_term"],
                    item["currency"],
                    item["amount_usd"],
                    item["amount_cny"],
                    item["exchange_rate"],
                    item["exchange_month"],
                    item["rate_adjusted"],
                    item["po_no"] or None,
                    user["id"],
                    operator,
                    item["transport_type"],
                    item["email_sent_date"] or None,
                    item["remark"] or None,
                    item["bl_no"] or None,
                    is_batch,
                    is_master,
                    batch_id,
                    batch_prefix,
                    sub_no,
                    sub_seq,
                    user["id"],
                    user["id"],
                    ts,
                    ts,
                ),
            )
            if item.get("clipped"):
                result.clipped += 1
            if batch_id:
                used_batches.add(batch_id)
            result.imported += 1
        for batch_id in used_batches:
            refresh_batch_amounts(db, batch_id)
        for batch_id in created_batches - used_batches:
            db.execute("DELETE FROM batch_slots WHERE batch_id=?", (batch_id,))
            db.execute("DELETE FROM batches WHERE id=?", (batch_id,))
        result.batch_count = len(used_batches)
        if result.imported:
            log_op(
                db,
                user,
                "invoice_import",
                f"批量导入发票 {result.imported} 条",
                "invoice",
                None,
                {
                    "成功": result.imported,
                    "跳过": result.skipped,
                    "占用顺序号": result.occupied,
                    "批次": result.batch_count,
                },
            )
            db.commit()
        else:
            safe_rollback(db)
            result.occupied = 0
            result.legacy = 0
            result.zero_amounts = 0
            result.clipped = 0
            result.batch_count = 0
    except Exception:
        safe_rollback(db)
        raise


def ensure_batch(db, user, prefix, meta, ts):
    existing = db.execute(
        "SELECT id FROM batches WHERE prefix=? AND deleted_at IS NULL",
        (prefix,),
    ).fetchone()
    if existing:
        return existing["id"], False
    conflict = find_conflict(db, prefix)
    if conflict:
        raise ValueError(f"批次前缀已被{conflict}占用")
    cur = db.execute(
        """INSERT INTO batches
           (prefix, number_rule, number_year, transport_type, business_type,
            num_pool, num_seq, created_by, created_at)
           VALUES (?, ?, ?, ?, ?, NULL, NULL, ?, ?)""",
        (
            prefix,
            meta["number_rule"],
            meta["number_year"],
            meta["transport_type"],
            meta["business_type"],
            user["id"],
            ts,
        ),
    )
    return cur.lastrowid, True


def occupy_sub(db, batch_id, seq, number):
    ts = now_str()
    existing = db.execute(
        "SELECT seq, status FROM batch_slots WHERE batch_id=?",
        (batch_id,),
    ).fetchall()
    have = {row["seq"]: row["status"] for row in existing}
    if have.get(seq) == "used":
        raise ValueError("分单号已被占用")
    for index in range(1, int(seq)):
        if index not in have:
            db.execute(
                """INSERT INTO batch_slots (batch_id, seq, status, number, updated_at)
                   VALUES (?, ?, 'free', NULL, ?)""",
                (batch_id, index, ts),
            )
    if seq in have:
        db.execute(
            """UPDATE batch_slots
               SET status='used', number=?, updated_at=?
               WHERE batch_id=? AND seq=?""",
            (number, ts, batch_id, seq),
        )
    else:
        db.execute(
            """INSERT INTO batch_slots (batch_id, seq, status, number, updated_at)
               VALUES (?, ?, 'used', ?, ?)""",
            (batch_id, seq, number, ts),
        )


def lookup_contract(db, number):
    if not number:
        return None
    row = db.execute(
        "SELECT id FROM contracts WHERE contract_no=? AND deleted_at IS NULL",
        (number,),
    ).fetchone()
    return None if row is None else row["id"]


def find_conflict(db, number, allow_batch_id=None):
    if db.execute(
        "SELECT 1 FROM contracts WHERE contract_no=? AND deleted_at IS NULL",
        (number,),
    ).fetchone():
        return "合同"
    if db.execute(
        "SELECT 1 FROM invoices WHERE invoice_no=? AND deleted_at IS NULL",
        (number,),
    ).fetchone():
        return "发票"
    row = db.execute(
        "SELECT id FROM batches WHERE prefix=? AND deleted_at IS NULL",
        (number,),
    ).fetchone()
    if row and row["id"] != allow_batch_id:
        return "批次"
    return ""


def match_contract_pool(number):
    for pool in ("contract_mn", "contract_al"):
        parsed = parse_pool_number(pool, number)
        if parsed:
            return pool, parsed
    return None, None


def match_invoice_pool(number):
    for pool in ("invoice_al_sea", "invoice_al_air", "invoice_mn"):
        parsed = parse_pool_number(pool, number)
        if parsed:
            return pool, parsed
    return None, None


def rule_of_pool(pool):
    return {
        "invoice_al_sea": "铝矿海运",
        "invoice_al_air": "铝矿空运",
        "invoice_mn": "锰矿",
    }[pool]


def infer_contract_business(number):
    upper = number.upper()
    if "MAZ" in upper or "BMGG" in upper:
        return "铝矿"
    return "锰矿"


def legacy_invoice_rule(number, transport, explicit_business, explicit_rule):
    if explicit_rule in ("铝矿海运", "铝矿空运", "锰矿"):
        business = "锰矿" if explicit_rule == "锰矿" else "铝矿"
        return business, explicit_rule
    body = re.sub(r"-\d{3}$", "", number).upper()
    if "MAZ" not in body and "BMGG" not in body and "GMI" in body:
        return "锰矿", "锰矿"
    if transport == "空运" or body.endswith("A"):
        return "铝矿", "铝矿空运"
    if explicit_business == "锰矿":
        return "锰矿", "锰矿"
    return "铝矿", "铝矿海运"


def infer_transport(number, raw):
    text = text_of(raw)
    if text in ("海运", "空运"):
        return text, ""
    if "空运" in text:
        return "空运", text
    if "海运" in text:
        return "海运", text
    body = re.sub(r"-\d{3}$", "", number).upper()
    if body.endswith("A"):
        return "空运", text
    return "海运", text


def infer_year(number):
    matched = YEAR_IN_NUMBER_RE.search(number)
    if matched:
        year = int(matched.group(1))
        if 2000 <= year <= 2099:
            return year
    matched = SHORT_YEAR_RE.search(number)
    if matched:
        return 2000 + int(matched.group(1))
    return date.today().year


def norm_year(value):
    text = text_of(value)
    if re.fullmatch(r"20\d{2}", text):
        return int(text)
    return None


def norm_business(value):
    text = text_of(value)
    if "铝" in text:
        return "铝矿"
    if "锰" in text:
        return "锰矿"
    return ""


def norm_rule(value):
    text = text_of(value)
    if text in ("铝矿海运", "铝矿空运", "锰矿"):
        return text
    if "空运" in text:
        return "铝矿空运"
    if "海运" in text:
        return "铝矿海运"
    return ""


def normalize_contract_no(raw):
    text = "" if raw is None else str(raw).strip().replace(" ", "").replace("\u3000", "")
    text = text.replace("－", "-").replace("—", "-").replace("–", "-")
    if CONTRACT_CODE_RE.match(text) and re.search(r"\d", text):
        return text
    found = CONTRACT_EXTRACT_RE.search(text)
    if found and re.search(r"\d", found.group(0)):
        return found.group(0)[:64]
    return ""


def invoice_amounts(raw, cols):
    usd = ""
    cny = ""
    term = ""
    rate = ""
    month = ""
    adjusted = 1
    usd_cell = cell(raw, cols.get("amount_usd"))
    cny_cell = cell(raw, cols.get("amount_cny"))
    if not is_empty(usd_cell):
        usd = parse_money(usd_cell, "美元金额")
    if not is_empty(cny_cell):
        cny = parse_money(cny_cell, "人民币金额")
    combined = cell(raw, cols.get("price_combined"))
    if not usd and not cny and not is_empty(combined):
        usd, cny, term = parse_combined_price(combined)
    rate_cell = cell(raw, cols.get("rate"))
    if not is_empty(rate_cell):
        rate = parse_decimal(str(rate_cell).strip().replace(",", ""), 6, "汇率", required=True)
        if Decimal(rate) < 0:
            raise ValueError("汇率不能为负数")
        adjusted = 1
    elif usd and cny and Decimal(usd) > 0 and Decimal(cny) > 0:
        rate = format(
            (Decimal(cny) / Decimal(usd)).quantize(Decimal("0.000001"), rounding=ROUND_HALF_UP),
            "f",
        )
        adjusted = 1
    if not rate:
        rate = "0.000000"
    month_cell = cell(raw, cols.get("exchange_month"))
    month_text = text_of(month_cell)
    if re.fullmatch(r"20\d{2}-\d{2}", month_text):
        month = month_text
    currency_cell = text_of(cell(raw, cols.get("currency"))).upper()
    if currency_cell in ("CNY", "USD", "人民币", "美元"):
        currency = "USD" if currency_cell in ("USD", "美元") else "CNY"
    elif usd and Decimal(usd) > 0 and (not cny or Decimal(cny) == 0):
        currency = "USD"
    else:
        currency = "CNY"
    if not usd:
        usd = "0.00"
    if not cny:
        cny = "0.00"
    zero = Decimal(usd) == 0 and Decimal(cny) == 0
    return {
        "usd": usd,
        "cny": cny,
        "rate": rate,
        "exchange_month": month,
        "currency": currency,
        "term": term,
        "adjusted": adjusted,
        "zero": zero,
    }


def parse_combined_price(value):
    if isinstance(value, bool):
        raise ValueError("金额必须是数字")
    if isinstance(value, (int, float)):
        return "", parse_money(value, "金额"), ""
    text = text_of(value)
    if not text:
        return "", "", ""
    upper = text.upper()
    term = "FOB" if "FOB" in upper else "CIF" if "CIF" in upper else ""
    usd_match = re.search(r"(?:USD|美元)\s*[:：]?\s*(\d[\d,，]*\.?\d*)", text, re.I)
    cny_match = re.search(r"(?:CNY|RMB|人民币|￥|¥)\s*[:：]?\s*(\d[\d,，]*\.?\d*)", text, re.I)
    usd = parse_money(usd_match.group(1), "美元金额") if usd_match else ""
    cny = parse_money(cny_match.group(1), "人民币金额") if cny_match else ""
    if not usd and not cny:
        cny = parse_money(text, "金额")
    return usd, cny, term


def parse_money(value, name):
    if isinstance(value, bool):
        raise ValueError(f"{name}必须是数字")
    if isinstance(value, (int, float)):
        text = format(Decimal(str(value)), "f")
    else:
        text = text_of(value)
        text = text.replace(",", "").replace("，", "").replace("￥", "").replace("¥", "").replace("元", "").replace(" ", "")
        matched = re.search(r"\d+(?:\.\d+)?", text)
        if not matched:
            raise ValueError(f"{name}必须是数字")
        text = matched.group(0)
    return parse_decimal(text, 2, name, required=True)


def parse_any_date(value, name):
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        serial = float(value)
        if 20000 <= serial <= 80000:
            parsed = datetime(1899, 12, 30) + timedelta(days=int(serial))
            return parsed.strftime("%Y-%m-%d")
        raise ValueError(f"{name}不是有效日期")
    text = text_of(value).replace(".", "-").replace("/", "-")
    text = text.replace("年", "-").replace("月", "-").replace("日", "")
    matched = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})", text)
    if not matched:
        raise ValueError(f"{name}必须为日期")
    try:
        return date(int(matched.group(1)), int(matched.group(2)), int(matched.group(3))).strftime("%Y-%m-%d")
    except ValueError as exc:
        raise ValueError(f"{name}不是有效日期") from exc


def loose_date(value):
    if is_empty(value):
        return "", ""
    try:
        return parse_any_date(value, "日期"), ""
    except ValueError:
        return "", text_of(value)


def split_invoice_numbers(value):
    text = text_of(value)
    if not text:
        return []
    parts = re.split(r"[\r\n]+", text)
    return [part.strip() for part in parts if part.strip()]


def cell(raw, index):
    if index is None or index >= len(raw):
        return None
    return raw[index]


def is_empty(value):
    if value is None:
        return True
    if isinstance(value, str) and not value.strip():
        return True
    return False


def text_of(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int) and not isinstance(value, bool):
        return str(value)
    return (
        str(value)
        .replace("\xa0", " ")
        .replace("\u3000", " ")
        .strip()
        .replace("－", "-")
        .replace("—", "-")
        .replace("–", "-")
    )


def norm_header(value):
    return re.sub(r"\s+", "", text_of(value))


def col_index(headers, aliases):
    names = sorted(aliases, key=len, reverse=True)
    for index, header in enumerate(headers):
        if header and header in names:
            return index
    for index, header in enumerate(headers):
        if not header:
            continue
        folded = header.upper()
        for name in names:
            if len(name) < 4:
                continue
            if header.startswith(name) or (name.isascii() and folded.startswith(name.upper())):
                return index
    return None


def clip(value, limit):
    text = "" if value is None else str(value).strip()
    if len(text) <= limit:
        return text, False
    return text[:limit], True

