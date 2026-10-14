# 示例程序：从本行复制到「示例程序结束」，保存为 .py 后运行。															
# pip install requests openpyxl															
# 服务 http://127.0.0.1:8001    登录 admin / 123456															
# 编号和名称都带 test。已经存在就跳过，可以重复运行。															
# 采购计划 TEST-PL-001、示例合同 TEST-CG-001 的上传程序在工作表「上传示例合同」。															
import requests															
															
BASE = "http://127.0.0.1:8001"															
															
															
def call(method, path, token=None, **kw):															
    headers = {}															
    if token:															
        headers["Authorization"] = "Bearer " + token															
    resp = requests.request(method, BASE + path, headers=headers, timeout=60, **kw)															
    body = resp.json()															
    if body.get("code") != 0:															
        raise SystemExit(method + " " + path + " 失败: " + str(body.get("message")))															
    return body.get("data")															
															
															
def login():															
    data = call("POST", "/api/auth/login", json={"login_name": "admin", "password": "123456"})															
    user = data["user"]															
    print("登录成功", user["name"], user["employee_id"], "公司", user["company_id"])															
    return data["token"], user["employee_id"], user["company_id"]															
															
															
def items_of(data):															
    if isinstance(data, dict):															
        return data.get("items") or []															
    if isinstance(data, list):															
        return data															
    return []															
															
															
def find_one(token, path, params, **eq):															
    query = {"page": 1, "pageSize": 500}															
    query.update(params or {})															
    data = call("GET", path, token, params=query)															
    for row in items_of(data):															
        if all(str(row.get(key)) == str(value) for key, value in eq.items()):															
            return row															
    return None															
															
import io															
from openpyxl import Workbook															
															
															
def workbook_bytes(blocks):															
    book = Workbook()															
    sheet = book.active															
    sheet.title = "test"															
    row_no = 1															
    for headers, samples in blocks:															
        for col, text in enumerate(headers, start=1):															
            sheet.cell(row_no, col, text)															
        row_no += 1															
        for sample in samples:															
            for col, text in enumerate(sample, start=1):															
                sheet.cell(row_no, col, text)															
            row_no += 1															
        row_no += 1															
    buf = io.BytesIO()															
    book.save(buf)															
    return buf.getvalue()															
															
															
def upload(token, module, filename, content, mode, direction="", party_type=""):															
    parsed = call(															
        "POST", "/api/import/" + module + "/parse", token,															
        files={"file": (filename, content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},															
        data={"mode": mode, "direction": direction, "party_type": party_type},															
    )															
    bad = [row for row in (parsed.get("rows") or []) if not row.get("ok")]															
    if parsed.get("batch_error") or bad:															
        raise SystemExit("解析失败 " + str(parsed.get("batch_error") or bad))															
    result = call(															
        "POST", "/api/import/" + module + "/execute", token,															
        json={															
            "file_hash": parsed["file_hash"],															
            "mode": parsed.get("mode") or mode,															
            "direction": direction,															
            "party_type": party_type,															
            "rows": parsed["rows"],															
            "locks": parsed.get("locks") or [],															
        },															
    )															
    print("上传", filename, result.get("message"), "成功", result.get("success"), "失败", result.get("fail"))															
    if result.get("fail"):															
        raise SystemExit(result.get("errors"))															
    return result															
															
															
def ensure_party(token):															
    found = find_one(token, "/api/parties", {"keyword": "TEST-GYS-001", "type": "supplier"}, code="TEST-GYS-001")															
    if found:															
        print("已有 test 供应商", found["id"])															
        return found															
    call("POST", "/api/parties", token, json={															
        "type": "supplier", "code": "TEST-GYS-001", "name": "test供应商", "short_name": "test供",															
        "tax_no": "TEST91500000000000000X", "default_tax_rate": 13, "status": "启用",															
        "address": "test地址", "mobile": "13800000000", "bank_name": "test银行",															
        "account_name": "test供应商", "bank_account": "TEST6222000000000001",															
        "legal_name": "test法人", "agent_name": "test代理", "remark": "test客商",															
        "contacts": [{"name": "test联系人", "phone": "13800000000", "is_default": True}],															
    })															
    found = find_one(token, "/api/parties", {"keyword": "TEST-GYS-001", "type": "supplier"}, code="TEST-GYS-001")															
    print("已建 test 供应商", found["id"])															
    return found															
															
															
def ensure_project(token):															
    found = find_one(token, "/api/projects", {"keyword": "TEST-XM-001"}, project_code="TEST-XM-001")															
    if found:															
        return found															
    call("POST", "/api/projects", token, json={															
        "project_code": "TEST-XM-001", "name": "test项目", "year": "2026", "project_type": "test",															
        "source": "test", "customer_name": "test客户", "start_date": "2026-09-01",															
        "end_date": "2026-12-31", "status": "进行中", "remark": "test项目",															
    })															
    return find_one(token, "/api/projects", {"keyword": "TEST-XM-001"}, project_code="TEST-XM-001")															
															
															
def ensure_type(token):															
    found = find_one(token, "/api/contract-types", {"keyword": "TEST-HTLX"}, code="TEST-HTLX")															
    if found:															
        return found															
    call("POST", "/api/contract-types", token, json={															
        "code": "TEST-HTLX", "name": "test材料采购", "status": "启用", "remark": "test合同类型",															
    })															
    return find_one(token, "/api/contract-types", {"keyword": "TEST-HTLX"}, code="TEST-HTLX")															
															
															
def ensure_warehouse(token, code, name):															
    found = find_one(token, "/api/warehouses", {"keyword": code}, code=code)															
    if found:															
        return found															
    call("POST", "/api/warehouses", token, json={"code": code, "name": name, "status": "启用", "remark": "test仓库"})															
    return find_one(token, "/api/warehouses", {"keyword": code}, code=code)															
															
															
def ensure_account(token):															
    found = find_one(token, "/api/fund-accounts", {"keyword": "TEST-ZH-001"}, code="TEST-ZH-001")															
    if found:															
        return found															
    call("POST", "/api/fund-accounts", token, json={															
        "code": "TEST-ZH-001", "name": "test基本户", "account_no": "TEST9558800000000001", "status": "启用",															
    })															
    return find_one(token, "/api/fund-accounts", {"keyword": "TEST-ZH-001"}, code="TEST-ZH-001")															
															
															
def ensure_plans(token):															
    catalog = {															
        "TEST-M001": ["2026-09", "1", "TEST-M001", "TEST-PL-001", "test类别", "test钢板", "test plate", "TEST-10mm", "吨", 10, "test用途", "test部门", "test计划", "", "", ""],															
        "TEST-M002": ["2026-09", "2", "TEST-M002", "TEST-PL-001", "test类别", "test螺栓", "test bolt", "TEST-M16", "套", 20, "test用途", "test部门", "test计划", "", "", ""],															
    }															
    missing = [code for code in catalog if find_one(token, "/api/purchase-plans", {"keyword": code}, material_code=code) is None]															
    if not missing:															
        print("已有 test 采购计划 TEST-M001、TEST-M002")															
        return															
    headers = ["计划月份", "采购序号", "物料编号", "采购计划单号", "使用设备", "中文名称", "英文名称", "规格型号", "单位", "采购数量", "用途", "使用部门", "备注", "库存数量", "在途数量", "4个月消耗量"]															
    content = workbook_bytes([(headers, [catalog[code] for code in missing])])															
    upload(token, "purchase-plan", "test采购计划.xlsx", content, "append")															
															
															
def contract_file(employee_id):															
    master = ["合同编号", "合同名称", "合同类型", "对方编码", "业务员员工编号", "合同状态", "计价方式", "税率", "签订日期", "开始日期", "结束日期", "内容概述", "备注"]															
    head = ["TEST-CG-001", "test采购合同", "test材料采购", "TEST-GYS-001", employee_id, "已签订", "固定总价", 13, "2026-09-28", "2026-09-28", "2026-12-31", "test合同内容", "test"]															
    detail = ["对应合同编号", "物料编号", "中文名称", "规格", "单位", "数量", "含税单价", "备注", "计划单号"]															
    line1 = ["TEST-CG-001", "TEST-M001", "test钢板", "TEST-10mm", "吨", 2, 100, "test行", "TEST-PL-001"]															
    line2 = ["TEST-CG-001", "TEST-M002", "test螺栓", "TEST-M16", "套", 4, 5, "test行", "TEST-PL-001"]															
    return workbook_bytes([(master, [head]), (detail, [line1, line2])])															
															
															
def contract_body(employee_id, party_id, project_id):															
    return {															
        "direction": "purchase", "nature": "main", "contract_code": "TEST-CG-001", "name": "test采购合同",															
        "type_name": "test材料采购", "party_id": party_id, "project_id": project_id, "owner_id": employee_id,															
        "status": "已签订", "pricing": "fixed", "tax_rate": 13, "sign_date": "2026-09-28",															
        "start_date": "2026-09-28", "end_date": "2026-12-31", "content": "test合同内容", "remark": "test",															
        "subject_name": "test钢板、test螺栓",															
        "items": [															
            {"material_code": "TEST-M001", "name_cn": "test钢板", "name_en": "test plate", "spec": "TEST-10mm", "unit": "吨", "qty": 2, "price": 100, "plan_no": "TEST-PL-001", "remark": "test行"},															
            {"material_code": "TEST-M002", "name_cn": "test螺栓", "name_en": "test bolt", "spec": "TEST-M16", "unit": "套", "qty": 4, "price": 5, "plan_no": "TEST-PL-001", "remark": "test行"},															
        ],															
    }															
															
															
def find_contract(token):															
    return find_one(token, "/api/contracts", {"keyword": "TEST-CG-001", "direction": "purchase"}, contract_code="TEST-CG-001")															
															
															
def load_contract(token):															
    row = find_contract(token)															
    if row is None:															
        raise SystemExit("没有找到合同 TEST-CG-001")															
    return call("GET", "/api/contracts/" + str(row["id"]), token)															
															
															
def prepare_base(token):															
    party = ensure_party(token)															
    project = ensure_project(token)															
    ensure_type(token)															
    ensure_plans(token)															
    return party, project															
															
															
def ensure_contract_json(token, employee_id):															
    party, project = prepare_base(token)															
    if find_contract(token):															
        detail = load_contract(token)															
        print("已有 test 合同", detail.get("contract_code"), detail.get("id"))															
        return detail															
    created = call("POST", "/api/contracts", token, json=contract_body(employee_id, party["id"], project["id"]))															
    print("JSON 已建 test 合同", created.get("contract_code"), "含税", created.get("amount"), "不含税", created.get("amount_ex_tax"))															
    return created															
															
															
def ensure_contract_upload(token, employee_id):															
    prepare_base(token)															
    if find_contract(token):															
        detail = load_contract(token)															
        print("已有 test 合同", detail.get("contract_code"), detail.get("id"))															
        return detail															
    upload(token, "contract", "test采购合同.xlsx", contract_file(employee_id), "append", direction="purchase")															
    detail = load_contract(token)															
    print("已上传 test 合同", detail.get("contract_code"), "含税", detail.get("amount"), "不含税", detail.get("amount_ex_tax"))															
    return detail															
															
															
def ensure_order(token, contract):															
    rows = items_of(call("GET", "/api/orders", token, params={"contract_id": contract["id"], "direction": "purchase", "page": 1, "pageSize": 20}))															
    if rows:															
        print("已有 test 订单", rows[0].get("order_code"))															
        return rows[0]															
    created = call("POST", "/api/orders", token, json={															
        "contract_id": contract["id"], "direction": "purchase", "order_code": "orderTEST-CG-001", "remark": "test订单",															
        "items": [															
            {"material_code": "TEST-M001", "name_cn": "test钢板", "spec": "TEST-10mm", "unit": "吨", "qty": 2, "price": 100, "remark": "test行"},															
            {"material_code": "TEST-M002", "name_cn": "test螺栓", "spec": "TEST-M16", "unit": "套", "qty": 4, "price": 5, "remark": "test行"},															
        ],															
    })															
    print("已建 test 订单", created.get("order_code"), created.get("amount"))															
    return created															
															
															
def ensure_inbound(token, contract):															
    house = ensure_warehouse(token, "TEST-CK-001", "test仓库")															
    rows = items_of(call("GET", "/api/wms-bills", token, params={"type": "inbound", "contract_id": contract["id"], "page": 1, "pageSize": 50}))															
    for row in rows:															
        if row.get("bill_code") == "TEST-RK-001":															
            print("已有 test 入库", row["bill_code"])															
            return row															
    created = call("POST", "/api/wms-bills", token, json={															
        "type": "inbound", "contract_id": contract["id"], "warehouse_id": house["id"],															
        "bill_code": "TEST-RK-001", "bill_date": "2026-09-28", "remark": "test入库",															
        "items": [{"material_code": "TEST-M001", "name_cn": "test钢板", "spec": "TEST-10mm", "unit": "吨", "qty": 1, "price": 100, "remark": "test行"}],															
    })															
    print("已建 test 入库", created.get("bill_code") or created.get("id"))															
    return created															
															
															
def ensure_payment(token, contract):															
    account = ensure_account(token)															
    rows = items_of(call("GET", "/api/settle-bills", token, params={"contract_id": contract["id"], "page": 1, "pageSize": 50}))															
    for row in rows:															
        if str(row.get("bill_code") or "").startswith("TEST-FK"):															
            print("已有 test 付款", row.get("bill_code"))															
            return row															
    stage = (contract.get("stages") or [None])[0]															
    if not stage or float(stage.get("remain") or 0) <= 0:															
        print("test 合同第一阶段没有待付金额")															
        return None															
    created = call("POST", "/api/settle-bills/multi-stage", token, json={															
        "contract_id": contract["id"], "party_id": contract["party_id"], "direction": "payment",															
        "bill_code": "TEST-FK-001", "bill_date": "2026-09-28", "account_name": account["name"], "remark": "test付款",															
        "lines": [{"stage_id": stage["id"], "amount": stage["remain"], "remark": "test预付款"}],															
    })															
    print("已建 test 付款", created)															
    return created															
															
															
def ensure_invoice(token, contract):															
    found = find_one(token, "/api/invoices", {"keyword": "TEST-FP-001", "contract_id": contract["id"], "direction": "input"}, invoice_no="TEST-FP-001")															
    if found:															
        print("已有 test 发票", found["invoice_no"])															
        return found															
    created = call("POST", "/api/invoices", token, json={															
        "direction": "input", "contract_id": contract["id"], "party_id": contract["party_id"],															
        "invoice_no": "TEST-FP-001", "bill_code": "TEST-JX-001", "invoice_date": "2026-09-28",															
        "invoice_type": "增值税专用发票", "remark": "test发票",															
        "items": [{"material_code": "TEST-M001", "name_cn": "test钢板", "spec": "TEST-10mm", "unit": "吨", "qty": 1, "price": 100, "tax_rate": 13, "remark": "test行"}],															
    })															
    print("已建 test 发票", created.get("invoice_no"))															
    return created															
															
def main():															
    token, employee_id, company_id = login()															
    contract = ensure_contract_upload(token, employee_id)															
    ensure_order(token, contract)															
    ensure_inbound(token, contract)															
    ensure_payment(token, contract)															
    ensure_invoice(token, contract)															
    print("test 流程结束：计划 TEST-PL-001，合同 TEST-CG-001，订单 orderTEST-CG-001，入库 TEST-RK-001，付款 TEST-FK-001，发票 TEST-FP-001")															
    print("含税 2*100 + 4*5 = 220，不含税按 220/1.13 保留两位，应为 194.69")															
															
															
if __name__ == "__main__":															
    main()															
# 示例程序结束															
